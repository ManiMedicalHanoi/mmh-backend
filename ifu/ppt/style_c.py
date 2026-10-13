"""style_c.py  ·  MMH Style C (clean editorial) builder  ·  python-pptx
Write this file to /home/claude/style_c.py, then:  import sys; sys.path.insert(0, '/home/claude'); from style_c import *
All coordinates in inches, slide 13.333 x 7.5.
"""
import os, re, subprocess
from copy import deepcopy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.parts.slide import SlideLayoutPart
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palettes (roles, not raw colours)
PALETTES = {
 # ink/title = text · pri/sec = main colours · soft = light tint of pri · panel = card fill · line = hairlines
 # hi/hi_line = highlight fill + border (1 focus item) · acc = pop colour (badges, focus bar) · pos/neg = good/bad
 'navy':   dict(ink='3A4A5A', title='35506B', pri='35506B', sec='5B87B8', ter='9DBCDB', soft='EAF1F8', panel='F3F5F8', line='DDE3EA', muted='7F8C9A', hi='FFF6D9', hi_line='E9CF6E', acc='E9CF6E', pos='4E9A78', neg='C2605C', on_pri='FFFFFF'),
 'yellow': dict(ink='1E2328', title='1E2328', pri='FFD400', sec='5E6670', ter='B0B5BA', soft='FFF3B0', panel='F2F3F4', line='DDE0E3', muted='8A9199', hi='FFF3B0', hi_line='FFD400', acc='FFD400', pos='3E8E6B', neg='C2605C', on_pri='1E2328'),
 'green':  dict(ink='1B1B1F', title='1B1B1F', pri='0B6E4F', sec='55565C', ter='BDBBB2', soft='E2F0EA', panel='F1F0EC', line='BDBBB2', muted='55565C', hi='FFF4C2', hi_line='FFD23F', acc='FFD23F', pos='0B6E4F', neg='C2605C', on_pri='FFFFFF'),
 'indigo': dict(ink='1D1B26', title='302663', pri='302663', sec='D9482B', ter='C9C1E8', soft='ECE8F6', panel='FFFFFF', line='C9C1E8', muted='4A4658', hi='FBE7E2', hi_line='D9482B', acc='F0B429', pos='3E8E6B', neg='D9482B', on_pri='FFFFFF'),
 'duo':    dict(ink='222222', title='222222', pri='1F4E8C', sec='A3471F', ter='C98A1E', soft='E3ECF7', panel='F2F2F0', line='C9C9C9', muted='4D4D4D', hi='F6E6DE', hi_line='A3471F', acc='C98A1E', pos='1F4E8C', neg='A3471F', on_pri='FFFFFF'),
 'classic':dict(ink='1A1A1A', title='003047', pri='003047', sec='3A5CAA', ter='7D97CE', soft='F0F7FF', panel='F0F7FF', line='D9D9D9', muted='666666', hi='FFF8E1', hi_line='FFE100', acc='FFE100', pos='3E8E6B', neg='C2605C', on_pri='FFFFFF'),
 'teal':   dict(ink='1F2D2C', title='0F4C47', pri='0F766E', sec='2A9D8F', ter='99D5CF', soft='E0F2F1', panel='F4F7F7', line='D5E3E2', muted='5F7472', hi='FFF1E0', hi_line='F4A261', acc='F4A261', pos='2A9D8F', neg='D1495B', on_pri='FFFFFF'),
 'burgundy':dict(ink='2B1D22', title='5C1229', pri='7A1E3A', sec='B2475F', ter='E3B5C0', soft='F7E9ED', panel='F7F4F2', line='E2D8D4', muted='6E5F63', hi='FBF1D9', hi_line='C9A227', acc='C9A227', pos='3E8E6B', neg='B23A3A', on_pri='FFFFFF'),
 'charcoal':dict(ink='1E2124', title='1E2124', pri='2F3337', sec='5C6670', ter='A9B0B7', soft='ECEEF0', panel='F5F6F7', line='DADDE0', muted='6B7480', hi='FFEBDD', hi_line='F26B1D', acc='F26B1D', pos='3E8E6B', neg='C2402F', on_pri='FFFFFF'),
 'ocean':  dict(ink='1C2B36', title='0B3D5C', pri='0B5C8A', sec='2E8BC0', ter='9CCBE6', soft='E4F1F9', panel='F3F7FA', line='D3E2EC', muted='5E7385', hi='FFE9E3', hi_line='EF6F5E', acc='EF6F5E', pos='2E9D6F', neg='D64545', on_pri='FFFFFF'),
 'forest': dict(ink='22261F', title='1F3F36', pri='2F5D50', sec='8C6B4F', ter='CDBFA6', soft='E6EFEB', panel='F6F3EC', line='DDD5C6', muted='6B675E', hi='FFF3D6', hi_line='E0A526', acc='E0A526', pos='2F5D50', neg='B5523B', on_pri='FFFFFF'),
 'sky':    dict(ink='1F2A33', title='1F4E7A', pri='3B82C4', sec='34B39A', ter='A8D3EF', soft='E8F3FB', panel='F5FAFC', line='D6E6F0', muted='667A8A', hi='FFF4CC', hi_line='FFC93C', acc='FFC93C', pos='34B39A', neg='E05A5A', on_pri='FFFFFF'),
 'plum':   dict(ink='241C2B', title='3F2652', pri='5B3A70', sec='7E9C7A', ter='C9BED4', soft='EFE9F3', panel='F6F5F2', line='DED8E2', muted='6A6273', hi='EEF4E8', hi_line='7E9C7A', acc='D4A72C', pos='5E8C5A', neg='B34A4A', on_pri='FFFFFF'),
}
PALETTES['mmh'] = PALETTES['classic']          # old key kept for older scripts
PALETTE_NAMES = {'navy': 'Navy Data', 'yellow': 'Yellow · Graphite', 'green': 'Editorial Green', 'indigo': 'Storybook Indigo',
                 'duo': 'Dual Blue · Rust', 'classic': 'Deep Navy · Yellow', 'teal': 'Teal Clinical', 'burgundy': 'Burgundy · Gold',
                 'charcoal': 'Charcoal · Orange', 'ocean': 'Ocean · Coral', 'forest': 'Forest · Sand', 'sky': 'Sky · Mint', 'plum': 'Plum · Sage'}
PALETTE_ORDER = list(PALETTE_NAMES)            # number shown to the user = index + 1
def rgb(h): return RGBColor.from_string(h)
def I(x): return Inches(x)

# ---------------------------------------------------------------- image layouts (true picture placeholders)
_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
def add_image_layout(prs, base_layout, name, boxes, fill='EEF1F5', label_col='7F8C9A'):
    """Clone base_layout (keeps its banner/logo) into a new layout with picture placeholders.
    boxes = [(x, y, w, h, label, geom)] inches; geom = 'rect' | 'ellipse' | 'roundRect'."""
    master = base_layout.slide_master; pkg = prs.part.package
    el = deepcopy(base_layout._element); el.cSld.set('name', name); tree = el.cSld.spTree
    for sp in list(tree):
        ph = sp.find('.//' + qn('p:ph'))
        if ph is not None and ph.get('type') not in ('dt', 'ftr', 'sldNum'): tree.remove(sp)
    part = SlideLayoutPart(pkg.next_partname('/ppt/slideLayouts/slideLayout%d.xml'), base_layout.part.content_type, pkg, el)
    part.relate_to(master.part, RT.SLIDE_MASTER)
    remap = {}
    for rid, rel in base_layout.part.rels.items():
        if rel.is_external or rel.reltype == RT.SLIDE_MASTER: continue
        remap[rid] = part.relate_to(rel.target_part, rel.reltype)
    for node in el.iter():
        for att in (qn('r:embed'), qn('r:link'), qn('r:id')):
            if node.get(att) in remap: node.set(att, remap[node.get(att)])
    rid = master.part.relate_to(part, RT.SLIDE_LAYOUT)
    used = [int(x.get('id')) for x in prs.part._element.iter(qn('p:sldMasterId'))]
    for m in prs.slide_masters: used += [int(x.get('id')) for x in m._element.iter(qn('p:sldLayoutId'))]
    sid = etree.SubElement(master._element.get_or_add_sldLayoutIdLst(), qn('p:sldLayoutId'))
    sid.set('id', str(max(used + [2147483648]) + 1)); sid.set(qn('r:id'), rid)
    for i, b in enumerate(boxes):
        x, y, w, h, label = b[:5]; geom = b[5] if len(b) > 5 else 'rect'
        tree.append(etree.fromstring(
            f'<p:sp xmlns:p="{_P}" xmlns:a="{_A}"><p:nvSpPr><p:cNvPr id="{200+i}" name="{label}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph type="pic" idx="{20+i}" sz="quarter"/></p:nvPr></p:nvSpPr><p:spPr><a:xfrm><a:off x="{int(I(x))}" y="{int(I(y))}"/>'
            f'<a:ext cx="{int(I(w))}" cy="{int(I(h))}"/></a:xfrm><a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="{fill}"/></a:solidFill></p:spPr>'
            f'<p:txBody><a:bodyPr anchor="ctr"/><a:lstStyle><a:lvl1pPr algn="ctr"><a:defRPr sz="1400"><a:solidFill><a:srgbClr val="{label_col}"/></a:solidFill></a:defRPr></a:lvl1pPr></a:lstStyle>'
            f'<a:p><a:r><a:rPr lang="vi-VN"/><a:t>{label}</a:t></a:r></a:p></p:txBody></p:sp>'))
    return [l for l in master.slide_layouts if l.name == name][-1]

# ---------------------------------------------------------------- image style lock (same STYLE text for every prompt)
IMAGE_STYLES = {
 'iso3d': 'clean isometric 3D render in a minimal corporate style; white and very light grey matte surfaces with accents of {pri}, {sec} and a touch of {acc}; soft diffuse studio light, gentle contact shadows, same 30-degree isometric camera in every image, generous negative space, plain {bg} background',
 'flat':  'clean flat isometric infographic illustration in an editorial magazine style; crisp geometric shapes with thin {ink} outlines, subtle halftone dot texture, generous negative space; limited palette of {pri}, {acc}, {ink} and light warm grey on an off-white paper background; not neon, not pastel, no gradients',
 'paint': 'soft watercolor and ink illustration, loose brush edges, muted palette led by {pri} and {sec} with warm {acc} highlights, misty atmosphere, same paper texture and same light direction in every image',
 'story': 'warm storybook editorial illustration, flat shapes with fine grain texture, deep {pri} background tones with {sec} and {acc} accents, cosy lighting, the same character design and line weight in every image',
 'clinical': 'photorealistic clinical macro photograph, neutral light grey seamless background, soft even studio lighting, shallow depth of field, accurate anatomy and instruments, subtle {pri} colour accents only in props',
}
NO_TEXT = 'no text, no letters, no numbers, no logos, no brand names, no readable writing, no watermark'
def style_lock(style, pal, extra=''):
    p = {k: '#' + v for k, v in pal.items()}; p['bg'] = '#F7F8FA'
    return IMAGE_STYLES[style].format(**p) + '; strong contrast so it stays clear on a projector; ' + NO_TEXT + (('; ' + extra) if extra else '')

# ---------------------------------------------------------------- deck
class Deck:
    """mode: 'mmh' (MMH template, yellow banner + MANI logo) | 'blank' (white, own header) | 'custom' (user template)."""
    def __init__(self, mode='mmh', template=None, palette='navy', font='Aptos', title_font=None, sections=(),
                 image_style='iso3d', content_layout=None, safe_right=None, top=None):
        self.mode, self.font, self.tfont = mode, font, title_font or font
        self.P = PALETTES[palette] if isinstance(palette, str) else palette
        self.sections, self.image_style = list(sections), image_style
        self.prs = Presentation(template) if template else Presentation()
        self.prs.slide_width, self.prs.slide_height = I(13.333), I(7.5)
        L = self.prs.slide_layouts
        if mode == 'mmh':   self.layout = L[1]; self.safe_right = 10.9; self.top = 0.66
        elif mode == 'blank': self.layout = L[6]; self.safe_right = 12.4; self.top = 0.55
        else: self.layout = L[content_layout or 0]; self.safe_right = safe_right or 10.9; self.top = top or 0.66
        self.page, self.log, self._img_layouts, self.img_no, self.gallery = (1 if mode == 'mmh' else 0), [], {}, 0, []
        self.lock = style_lock(image_style, self.P)

    # -------- text & shapes
    def c(self, role): return rgb(self.P.get(role, role))
    def text(self, sl, x, y, w, h, paras, align='l', anchor='t', spc=0, font=None):
        tb = sl.shapes.add_textbox(I(x), I(y), I(w), I(h)); tf = tb.text_frame; tf.word_wrap = True
        tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
        for s_ in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, s_, 0)
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
            if spc: para.space_after = Pt(spc)
            for r_ in (p if isinstance(p, list) else [p]):
                t, sz, b, col = r_[:4]; it = r_[4] if len(r_) > 4 else False
                r = para.add_run(); r.text = t
                r.font.name, r.font.size, r.font.bold, r.font.italic, r.font.color.rgb = font or self.font, Pt(sz), b, it, self.c(col)
        return tb
    def box(self, sl, x, y, w, h, fill='panel', line=None, kind=MSO_SHAPE.ROUNDED_RECTANGLE, adj=None):
        s = sl.shapes.add_shape(kind, I(x), I(y), I(w), I(h)); s.shadow.inherit = False
        st = s._element.find(qn('p:style'))
        if st is not None: s._element.remove(st)          # no theme shadow / outline
        if fill: s.fill.solid(); s.fill.fore_color.rgb = self.c(fill)
        else: s.fill.background()
        if line: s.line.color.rgb = self.c(line); s.line.width = Pt(1)
        else: s.line.fill.background()
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = adj if adj is not None else (0.06 if min(w, h) > 1 else 0.14)
        return s
    def hline(self, sl, x1, y, x2, col='line', w=1):
        ln = sl.shapes.add_connector(1, I(x1), I(y), I(x2), I(y)); ln.line.color.rgb = self.c(col); ln.line.width = Pt(w)
    def badge(self, sl, x, y, n, d=0.38, fill='pri', col='on_pri', sz=13):
        self.box(sl, x, y, d, d, fill, kind=MSO_SHAPE.OVAL)
        self.text(sl, x, y, d, d, [(str(n), sz, True, col)], 'c', 'm')
    def chip(self, sl, x, y, t, fill='soft', col='title', sz=11, h=0.3):
        w = 0.25 + 0.085 * len(t) * sz / 11
        self.box(sl, x, y, w, h, fill, adj=0.5); self.text(sl, x, y, w, h, [(t, sz, True, col)], 'c', 'm'); return w
    def kpi(self, sl, x, y, w, h, value, label, sub='', dark=False, vsz=30):
        self.box(sl, x, y, w, h, 'pri' if dark else 'panel')
        c1, c2, c3 = ('on_pri', 'acc', 'on_pri') if dark else ('title', 'sec', 'muted')
        self.text(sl, x + .2, y + .14, w - .4, .62, [(value, vsz, True, c1)], anchor='m')
        self.text(sl, x + .2, y + .8, w - .4, .3, [(label, 13, True, c2)])
        if sub: self.text(sl, x + .2, y + 1.1, w - .4, max(h - 1.18, .3), [(sub, 11.5, False, c3)])
    def item(self, sl, x, y, w, n, head, body=None, fill='pri', hs=15, bs=13):
        self.badge(sl, x, y + .02, n, fill=fill)
        self.text(sl, x + .55, y, w - .55, .34, [(head, hs, True, 'title')])
        if body: self.text(sl, x + .55, y + .36, w - .55, .7, [(body, bs, False, 'ink')])
    def icon(self, sl, name, x, y, d=0.5, role='sec', circle=True):
        """Lucide icon tinted with a palette role. Needs ensure_icons(deck) once."""
        if circle: self.box(sl, x, y, d, d, 'soft', kind=MSO_SHAPE.OVAL)
        p = f'/home/claude/style_c_icons/{name}_{self.P[role]}.png'
        pad = d * .22 if circle else 0
        if os.path.exists(p): sl.shapes.add_picture(p, I(x + pad), I(y + pad), I(d - 2 * pad), I(d - 2 * pad))

    # -------- tables & charts (native, editable)
    def table(self, sl, x, y, col_w, rows, row_h=.36, head=12, body=13, styles=None, aligns=None):
        nr, nc = len(rows), len(rows[0])
        tb = sl.shapes.add_table(nr, nc, I(x), I(y), I(sum(col_w)), I(row_h * nr)).table
        tp = tb._tbl.tblPr; tp.set('firstRow', '0'); tp.set('bandRow', '0')
        st = tp.find(qn('a:tableStyleId'))
        if st is not None: st.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
        for j, w in enumerate(col_w): tb.columns[j].width = I(w)
        for i in range(nr):
            tb.rows[i].height = I(row_h); kind = 'head' if i == 0 else (styles or {}).get(i, 'body')
            for j in range(nc):
                c = tb.cell(i, j); c.margin_left = c.margin_right = I(.08); c.margin_top = c.margin_bottom = I(.02)
                c.vertical_anchor = MSO_ANCHOR.MIDDLE; p = c.text_frame.paragraphs[0]; p.text = ''
                r = p.add_run(); r.text = str(rows[i][j]); f = r.font; f.name = self.font
                f.size = Pt(head if kind == 'head' else body); f.bold = kind != 'body'
                f.color.rgb = self.c('title' if kind != 'body' else 'ink')
                fill_role = {'head': 'soft', 'focus': 'hi', 'group': 'panel'}.get(kind)
                if fill_role: c.fill.solid(); c.fill.fore_color.rgb = self.c(fill_role)
                else: c.fill.background()
                self._border(c, ('t', 'b') if kind == 'head' else ('b',), self.P['sec'] if kind == 'head' else self.P['line'])
                p.alignment = {'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT, 'c': PP_ALIGN.CENTER}[(aligns or ['l'] + ['r'] * (nc - 1))[j]]
        return tb
    def _border(self, cell, sides, color):
        tcPr = cell._tc.get_or_add_tcPr()
        for s_ in 'LRTB':
            for e in tcPr.findall(qn(f'a:ln{s_}')): tcPr.remove(e)
            ln = etree.SubElement(tcPr, qn(f'a:ln{s_}'))
            if s_.lower() in sides:
                ln.set('w', '12700'); c = etree.SubElement(etree.SubElement(ln, qn('a:solidFill')), qn('a:srgbClr')); c.set('val', color)
            else: ln.set('w', '0'); etree.SubElement(ln, qn('a:noFill'))
        for f in tcPr.findall(qn('a:solidFill')) + tcPr.findall(qn('a:noFill')): tcPr.remove(f); tcPr.append(f)
    def chart(self, sl, kind, x, y, w, h, cats, series, colors=None, fmt='0%', label_pos='out', legend=None, cat_sz=12, gap=60):
        """kind: 'col' | 'bar' | 'col_stack100' | 'bar_stack100' | 'donut' | 'line'. series = [(name, values)]."""
        K = {'col': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'col_stack100': XL_CHART_TYPE.COLUMN_STACKED_100,
             'bar_stack100': XL_CHART_TYPE.BAR_STACKED_100, 'donut': XL_CHART_TYPE.DOUGHNUT, 'line': XL_CHART_TYPE.LINE_MARKERS}[kind]
        cd = CategoryChartData(); cd.categories = cats
        for n, v in series: cd.add_series(n, v)
        ch = sl.shapes.add_chart(K, I(x), I(y), I(w), I(h), cd).chart
        ch.has_title = False; ch.font.name = self.font; ch.font.size = Pt(12); ch.font.color.rgb = self.c('ink')
        multi = len(series) > 1 or kind == 'donut'
        ch.has_legend = multi if legend is None else legend
        if ch.has_legend: ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False; ch.legend.font.size = Pt(12)
        colors = colors or ['pri', 'ter', 'sec', 'acc', 'muted']
        pl = ch.plots[0]
        if kind == 'donut':
            for i, _ in enumerate(cats):
                pt = pl.series[0].points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = self.c(colors[i % len(colors)])
        else:
            pl.gap_width = gap
            if 'stack' in kind: pl.overlap = 100
            elif multi: pl.overlap = -10
            for s_, col in zip(pl.series, colors):
                s_.format.fill.solid(); s_.format.fill.fore_color.rgb = self.c(col)
                if kind == 'line': s_.format.line.color.rgb = self.c(col)
            va = ch.value_axis; va.visible = False; va.has_major_gridlines = False
            ca = ch.category_axis; ca.format.line.color.rgb = self.c('line'); ca.tick_labels.font.size = Pt(cat_sz)
        for s_ in pl.series:
            dl = s_.data_labels; dl.show_value = True; dl.number_format = fmt; dl.number_format_is_linked = False
            dl.font.size = Pt(12); dl.font.bold = True; dl.font.color.rgb = self.c('on_pri' if label_pos == 'center' else 'title')
            if kind != 'donut': dl.position = XL_LABEL_POSITION.CENTER if label_pos == 'center' else XL_LABEL_POSITION.OUTSIDE_END
        return ch
    def highlight_point(self, ch, series_idx, point_idx, role='acc'):
        p = ch.plots[0].series[series_idx].points[point_idx]; p.format.fill.solid(); p.format.fill.fore_color.rgb = self.c(role)

    # -------- slides
    def _new(self, layout):
        self.page += 1
        sl = self.prs.slides.add_slide(layout)
        for sh in list(sl.shapes):
            if sh.is_placeholder and sh.placeholder_format.type != 18: sh._element.getparent().remove(sh._element)  # keep picture placeholders
        return sl
    def slide(self, kicker, title, section, pattern, source=None, images=(), chrome=True):
        """images = [(x, y, w, h, ratio, geom)] → picture placeholders (click the icon, or right-click > Change Picture).
        ratio must be a Gemini ratio: '16:9' '4:3' '3:2' '1:1' '3:4' '2:3' '9:16' '21:9'."""
        layout = self.layout; labels = []
        if images:
            boxes = []
            for (x, y, w, h, ratio, *g) in images:
                self.img_no += 1; lab = f'ẢNH {self.img_no:02d} · {ratio}'; labels.append((lab, ratio))
                boxes.append((x, y, w, h, lab, g[0] if g else 'rect'))
            layout = add_image_layout(self.prs, self.layout, f'IMG {self.page + 1:02d}', boxes, fill=self.P['panel'], label_col=self.P['muted'])
        sl = self._new(layout); sl._img_labels = labels; sl._page = self.page
        for (x_, y_, w_, h_, *_), (lab, _r) in zip(images, labels):  # visible frame behind the placeholder (preview / PDF)
            fr = self.box(sl, x_, y_, w_, h_, 'panel', kind=MSO_SHAPE.RECTANGLE)
            tb = self.text(sl, x_, y_ + h_ / 2 - .2, w_, .4, [(lab, 13, True, 'muted')], 'c', 'm')
            for e in (tb._element, fr._element): e.getparent().remove(e); sl.shapes._spTree.insert(2, e)
        self.log.append((self.page, pattern))
        if chrome: self._chrome(sl, kicker, title, section, source)
        return sl
    def _chrome(self, sl, kicker, title, section, source):
        top = self.top
        if self.mode == 'blank':
            self.hline(sl, .46, .32, 12.87, 'ink', 1.5)
            self.text(sl, .46, .1, 8, .2, [(kicker.upper(), 9, True, 'muted')])
            self.text(sl, 11.4, .1, 1.47, .2, [(f'TRANG {self.page:02d}', 9, True, 'muted')], 'r')
        else:
            self.text(sl, .46, .17, 6.0, .28, [(kicker.upper(), 11, True, 'sec')])
        self.text(sl, .46, top, self.safe_right - .46, .46, [(title, 25 if len(title) <= 62 else 23, True, 'title')], anchor='m', font=self.tfont)
        self.hline(sl, .46, top + .54, 12.87)
        if source: self.text(sl, .46, 6.66, 12.4, .25, [(source, 10, False, 'muted', True)])
        if self.sections:
            wt = (12.41 - .08 * (len(self.sections) - 1)) / len(self.sections)
            for i, s_ in enumerate(self.sections):
                x = .46 + i * (wt + .08)
                if s_ == section:
                    self.box(sl, x, 6.98, wt, .27, 'sec', adj=.3); self.text(sl, x + .1, 6.98, wt - .2, .27, [(s_, 10, True, 'on_pri')], anchor='m')
                else: self.text(sl, x + .1, 6.98, wt - .2, .27, [(s_, 10, False, 'muted')], anchor='m')
        if self.mode != 'blank':
            self.text(sl, 12.2, 6.98, .67, .27, [(f'{self.page:02d}', 10, False, 'muted')], 'r', 'm')

    def cover_mmh(self, kicker, line1, line2, subtitle):
        """MMH template slide 1: only replace text (title = 2 runs split by a line break)."""
        s0 = self.prs.slides[0]
        for sh in s0.shapes:
            if sh.name.startswith('タイトル'):
                runs = sh.text_frame.paragraphs[0].runs; runs[0].text = line1
                if len(runs) > 1: runs[1].text = line2
            elif sh.name.startswith('字幕'): sh.text_frame.paragraphs[0].runs[0].text = subtitle
            elif sh.name.startswith('テキスト'): sh.text_frame.paragraphs[0].runs[0].text = kicker
        self.log.append((1, 'C01'))
        return s0

    # -------- speaker notes (EVERY slide)
    def notes(self, sl, purpose, talk, reasoning=None, evidence=None, transition=None, images=()):
        """talk = list of spoken sentences. images = [subject prompt per placeholder on this slide, in order]."""
        parts = ['MỤC ĐÍCH', purpose, '', 'TRÌNH BÀY'] + [f'• {t}' for t in (talk if isinstance(talk, list) else [talk])]
        if reasoning: parts += ['', 'LẬP LUẬN', reasoning]
        if evidence: parts += ['', 'DẪN CHỨNG', evidence]
        if transition: parts += ['', 'CHUYỂN Ý', transition]
        for (lab, ratio), subj in zip(getattr(sl, '_img_labels', []), images):
            self.gallery.append((lab, getattr(sl, '_page', 1), subj))
            parts += ['', '────────', f'{lab}', 'Cách chèn: bấm biểu tượng ảnh trong khung, hoặc chuột phải vào khung > Change Picture.', '',
                      'PROMPT GEMINI (copy nguyên văn):', subj, '', 'STYLE (giữ nguyên cho mọi ảnh của deck):', self.lock, '', f'Aspect ratio {ratio}.']
        sl.notes_slide.notes_text_frame.text = '\n'.join(parts)

    def add_image_bible(self, sl):
        """Call after all slides are built: append 'BỘ ẢNH CỦA DECK' (style lock + list of every image) to the cover notes."""
        if not self.gallery: return
        lines = ['', '════════', 'BỘ ẢNH CỦA DECK (tạo cùng một phiên Gemini, dán STYLE vào mọi prompt)', '', 'STYLE:', self.lock, '']
        lines += [f'{lab} · slide {n} · {subj[:110]}' for lab, n, subj in self.gallery]
        tf = sl.notes_slide.notes_text_frame; tf.text = tf.text + '\n' + '\n'.join(lines)

    # -------- QA + save
    def check(self):
        issues = []
        for (p1, a), (p2, b) in zip(self.log, self.log[1:]):
            if a == b: issues.append(f'slide {p1}/{p2}: pattern {a} lặp ở 2 slide liền kề')
        from collections import Counter
        for k, n in Counter(p for _, p in self.log).items():
            if n > 2: issues.append(f'pattern {k} dùng {n} lần (tối đa 2)')
        for i, s_ in enumerate(self.prs.slides, 1):
            if self.mode == 'mmh' and i == 2: continue        # template sample slide, removed on save
            if not (s_.has_notes_slide and s_.notes_slide.notes_text_frame.text.strip()): issues.append(f'slide {i}: thiếu speaker note')
        return issues
    def save(self, path):
        if self.mode == 'mmh':   # template: slide 1 title, slide 2 sample content (delete), slide 3 thank-you (move last)
            lst = self.prs.slides._sldIdLst; ids = list(lst); lst.remove(ids[1]); th = list(lst)[1]; lst.remove(th); lst.append(th)
        self.prs.save(path); Presentation(path).save(path); return path

def ensure_icons(deck, names=('users', 'target', 'check', 'chart', 'book', 'shield', 'clock', 'file', 'alert', 'bulb', 'layers', 'arrow', 'box', 'globe', 'truck', 'wallet', 'award', 'quote', 'cal', 'eye'), roles=('sec', 'pri', 'on_pri')):
    """Render Lucide icons as PNG in palette colours (node + react-icons + sharp, global npm)."""
    lucide = dict(users='LuUsers', target='LuTarget', check='LuCircleCheck', chart='LuChartColumn', book='LuBookOpen', shield='LuShieldCheck',
                  clock='LuClock3', file='LuFileText', alert='LuTriangleAlert', bulb='LuLightbulb', layers='LuLayers', arrow='LuArrowRight',
                  box='LuPackage', globe='LuGlobe', truck='LuTruck', wallet='LuWallet', award='LuAward', quote='LuQuote', cal='LuCalendarDays',
                  eye='LuEye', scale='LuScale', building='LuBuilding2', search='LuSearch', gavel='LuGavel', tooth='LuSmile', flask='LuFlaskConical')
    out = '/home/claude/style_c_icons'; os.makedirs(out, exist_ok=True)
    want = {n: lucide[n] for n in names if n in lucide}; cols = {deck.P[r]: '#' + deck.P[r] for r in roles}
    js = """const React=require('react'),RDS=require('react-dom/server'),sharp=require('sharp'),lu=require('react-icons/lu');
const want=WANT, cols=COLS, out='OUT';
(async()=>{for(const [k,n] of Object.entries(want)){const C=lu[n]; if(!C) continue;
 for(const [ck,cv] of Object.entries(cols)){const svg=RDS.renderToStaticMarkup(React.createElement(C,{color:cv,size:256,strokeWidth:2}));
 await sharp(Buffer.from(svg)).resize(256,256).png().toFile(out+'/'+k+'_'+ck+'.png');}}})();"""
    import json
    js = js.replace('WANT', json.dumps(want)).replace('COLS', json.dumps(cols)).replace('OUT', out)
    root = '/home/claude/npmx/node_modules'
    subprocess.run(['node', '-e', js], env={**os.environ, 'NODE_PATH': root}, check=True)

def swatch_png(path='/home/claude/palettes.png', keys=None):
    """Numbered palette strips to show the user before Q3 (send with SendUserFile, display='render').
    The number on each row = PALETTE_ORDER index + 1, so the user can answer by typing the number."""
    from PIL import Image, ImageDraw, ImageFont
    keys = keys or PALETTE_ORDER
    roles = ['pri', 'sec', 'ter', 'acc', 'soft', 'panel', 'hi', 'ink']
    W, H, sw = 1400, 104, 150
    img = Image.new('RGB', (W, H * len(keys) + 20), 'white'); d = ImageDraw.Draw(img)
    try: f = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 24); fs = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 15)
    except Exception: f = fs = ImageFont.load_default()
    for i, k in enumerate(keys):
        y = 10 + i * H; n = PALETTE_ORDER.index(k) + 1 if k in PALETTE_ORDER else i + 1
        d.text((20, y + 32), f'{n} · {PALETTE_NAMES.get(k, k)}', fill='#222222', font=f)
        for j, r in enumerate(roles):
            x = 420 + j * (sw - 30); c = '#' + PALETTES[k][r]
            d.rounded_rectangle([x, y + 8, x + sw - 40, y + 66], 10, fill=c, outline='#DDDDDD')
            d.text((x, y + 72), c, fill='#555555', font=fs)
    img.save(path); return path
