"""PPT giới thiệu hệ thống dữ liệu nội bộ MMH · bản 2 (cuốn chiếu theo IFU, link IFU cuối mỗi slide)."""
import sys, os
sys.path.insert(0, '/home/claude')
from style_c import *
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE

SH = '/home/claude/ifu/shots/'
ICO = '/home/claude/img/icons/'
CACHE = '/home/claude/img/crop2'; os.makedirs(CACHE, exist_ok=True)
TPL = '/home/claude/build/MMH_PPT_TEMPLATE.pptx'

# ── Link IFU: đổi ở đây khi có link file trên Drive (mỗi phần 1 file) ──
IFU_FOLDER = 'https://drive.google.com/drive/folders/1g1lfTxLU51qTkxIl-bRqJwVyvN5r7j6q'
IFU_FILE = {'1': ('MMH_IFU_He_thong_du_lieu.pdf', IFU_FOLDER), '2': ('MMH_IFU_01_CRM.pdf', IFU_FOLDER),
            '3': ('MMH_IFU_02_Report_Hub.pdf', IFU_FOLDER), '4': ('MMH_IFU_03_KPI.pdf', IFU_FOLDER),
            '5': ('MMH_IFU_04_Training_Hub.pdf', IFU_FOLDER), '6': ('MMH_IFU_He_thong_du_lieu.pdf', IFU_FOLDER)}
PAGE = {'1.1': 8, '1.2': 9, '1.3': 9, '1.4': 10, '1.5': 11, '1.6': 12, '1.7': 15, '1.8': 16,
        '2.1': 6, '2.2': 7, '2.3': 9, '2.4': 12, '2.5': 14, '2.6': 17, '2.7': 21, '2.8': 23, '2.9': 26, '2.10': 29, '2.11': 31, '2.12': 33, '2.13': 38,
        '3.1': 6, '3.2': 9, '3.3': 12, '3.4': 15, '3.5': 18, '3.6': 21, '3.7': 24, '3.8': 26, '3.9': 29, '3.10': 31, '3.11': 33, '3.12': 38,
        '4.1': 6, '4.2': 7, '4.3': 8, '4.4': 10, '4.5': 10, '4.6': 11, '4.7': 11, '4.8': 14, '4.9': 16, '4.10': 17,
        '5.1': 6, '5.2': 7, '5.3': 9, '5.4': 11, '5.5': 13, '5.6': 15, '5.7': 17, '5.8': 19, '5.9': 23, '5.10': 28, '5.11': 31, '5.12': 32, '5.13': 33,
        '6.1': 134}

SEC = ['TỔNG QUAN', '1 · NỀN TẢNG', '2 · CRM', '3 · REPORT HUB', '4 · KPI', '5 · TRAINING HUB', 'TỔNG KẾT']
d = Deck('mmh', template=TPL, palette='classic', sections=SEC)
ensure_icons(d, names=('users', 'target', 'check', 'chart', 'book', 'shield', 'clock', 'file', 'alert', 'bulb', 'layers',
                       'arrow', 'box', 'globe', 'award', 'cal', 'eye', 'building', 'search', 'wallet', 'truck'), roles=('sec', 'pri', 'on_pri'))
LOG = []

# ───────────────────────── helpers
def modal_box(im):
    """Tìm khung hộp thoại trắng trên nền mờ (ảnh có lớp phủ); không thấy thì trả toàn ảnh."""
    import numpy as np
    g = np.asarray(im.convert('L').resize((im.width // 4, im.height // 4)), dtype=float)
    m = g >= 249
    rows = np.where(m.mean(1) > .18)[0]; cols = np.where(m.mean(0) > .18)[0]
    if len(rows) < 5 or len(cols) < 5: return None
    H, W = g.shape; y0, y1, x0, x1 = rows[0], rows[-1], cols[0], cols[-1]
    if (x1 - x0) * (y1 - y0) > .92 * W * H: return None
    pad = .012
    return (max(0, x0 / W - pad), max(0, y0 / H - pad), min(1, x1 / W + pad), min(1, y1 / H + pad))

def crop(path, w, h, box=None, anchor='top', maxw=2000, fit=True):
    """box = vùng cắt (tỉ lệ) hoặc 'auto' (khung hộp thoại). fit=True: giữ đủ vùng, thêm lề trắng cho đúng tỉ lệ khung."""
    im = Image.open(path).convert('RGB')
    if box == 'auto': box = modal_box(im)
    if box:
        W, H = im.size; im = im.crop((int(box[0] * W), int(box[1] * H), int(box[2] * W), int(box[3] * H)))
    W, H = im.size; r = w / h
    if fit:
        if abs(W / H - r) > .01:
            if W / H > r: nw, nh = W, int(W / r)
            else: nw, nh = int(H * r), H
            bg = Image.new('RGB', (nw, nh), (255, 255, 255))
            bg.paste(im, ((nw - W) // 2, 0 if anchor == 'top' and W / H > r else (nh - H) // 2)); im = bg
    else:
        if W / H > r:
            nw = int(H * r); x0 = (W - nw) // 2 if anchor != 'left' else 0; im = im.crop((x0, 0, x0 + nw, H))
        else:
            nh = int(W / r); y0 = 0 if anchor in ('top', 'left') else (H - nh) // 2; im = im.crop((0, y0, W, y0 + nh))
    if im.width > maxw: im = im.resize((maxw, int(maxw * im.height / im.width)), Image.LANCZOS)
    out = f'{CACHE}/{abs(hash((path, w, h, box, anchor, fit)))}.jpg'; im.save(out, quality=86); return out

AUTO = {'crm/team_switch', 'crm/trip_approve', 'crm/test_mode', 'rh/assign_mail', 'rh/assign_popup', 'rh/assign_track', 'th/quiz_question', 'th/quiz_result', 'crm/newacc_review'}
def shot(sl, sid, x, y, w, h, box=None, anchor='top'):
    p = SH + sid + '.jpg'
    if sid in AUTO: box = 'auto'
    d.box(sl, x - .03, y - .03, w + .06, h + .06, 'line', kind=MSO_SHAPE.RECTANGLE)
    sl.shapes.add_picture(crop(p, w, h, box, anchor), I(x), I(y), I(w), I(h))

def phone(sl, sid, x, y, h, box=None):
    im = Image.open(SH + sid + '.jpg'); r = im.width / im.height
    if box: r = r * (box[2] - box[0]) / (box[3] - box[1])
    w = h * r
    d.box(sl, x - .06, y - .06, w + .12, h + .12, 'ink', adj=.08)
    sl.shapes.add_picture(crop(SH + sid + '.jpg', w, h, box), I(x), I(y), I(w), I(h)); return w

def logo(sl, name, x, y, s=.4):
    sl.shapes.add_picture(ICO + name + '.png', I(x), I(y), I(s), I(s))

def footer(sl, sec, title, y=6.64, w=9.6, credit=True):
    """Dòng cuối slide: link tới đúng mục trong IFU."""
    part = sec.split('.')[0]; fname, url = IFU_FILE[part]; pg = PAGE.get(sec)
    tb = sl.shapes.add_textbox(I(.46), I(y), I(w), I(.27)); tf = tb.text_frame
    for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, a, 0)
    p = tf.paragraphs[0]
    r0 = p.add_run(); r0.text = 'Chi tiết  '; r0.font.size = Pt(10); r0.font.name = d.font; r0.font.color.rgb = d.c('muted')
    r = p.add_run(); r.text = f'IFU {sec} · {title}' + (f' · {fname} tr. {pg}' if pg else '') + '  ↗'
    r.font.size = Pt(10); r.font.bold = True; r.font.name = d.font; r.font.color.rgb = d.c('sec'); r.hyperlink.address = url
    if credit: d.text(sl, 10.2, 6.64, 2.67, .27, [('Ảnh thật · dữ liệu giả lập', 9, False, 'muted', True)], 'r')

def S(kicker, title, section, pattern, sec=None, sectitle=None):
    sl = d.slide(kicker, title, section, pattern)
    LOG.append((d.page, pattern, title))
    if sec: footer(sl, sec, sectitle or title)
    return sl

def steps(sl, x, y, w, items, gap=.62, sz=13.5, start=1):
    for i, t in enumerate(items):
        yy = y + i * gap
        d.badge(sl, x, yy, i + start, d=.34, sz=12)
        if isinstance(t, tuple):
            d.text(sl, x + .5, yy - .04, w - .5, .3, [(t[0], 14, True, 'title')])
            d.text(sl, x + .5, yy + .27, w - .5, gap - .3, [(t[1], 12.5, False, 'ink')])
        else:
            d.text(sl, x + .5, yy - .03, w - .5, gap, [(t, sz, False, 'ink')])

def bullets(sl, x, y, w, items, sz=13.5, gap=5):
    d.text(sl, x, y, w, .35 * len(items), [[('•  ', sz, True, 'sec'), (t, sz, False, 'ink')] for t in items], spc=gap)

def chev(sl, items, y=1.5, h=.7, x0=.46, W=12.41, last_dark=True, sub=None, subh=.9):
    n = len(items); w = (W + .1 * (n - 1)) / n
    for i, t in enumerate(items):
        x = x0 + i * (w - .1)
        fill = 'pri' if (last_dark and i == n - 1) else ('sec' if i == 0 else 'ter')
        d.box(sl, x, y, w, h, fill, kind=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        d.text(sl, x + .32, y, w - .55, h, [(t, 13, True, 'on_pri' if fill != 'ter' else 'title')], 'c', 'm')
        if sub: d.text(sl, x + .1, y + h + .15, w - .35, subh, [(sub[i], 12, False, 'ink')])

def note(sl, x, y, w, h, head, body, fill='hi', line='hi_line'):
    d.box(sl, x, y, w, h, fill, line=line)
    d.text(sl, x + .22, y + .1, w - .44, h - .2, [[(head + '  ', 13.5, True, 'title'), (body, 13, False, 'ink')]], anchor='m')

def divider(no, name, sub, points, img, sec, pattern='C04'):
    sl = d.slide(f'Phần {no}', f'Phần {no} · {name}', SEC[int(no)], pattern)
    LOG.append((d.page, pattern, name))
    d.text(sl, .46, 1.5, 5.4, .45, [(f'PHẦN {no}', 16, True, 'sec')])
    d.text(sl, .46, 1.95, 5.4, .85, [(name, 34, True, 'title')])
    d.text(sl, .46, 2.8, 5.4, .55, [(sub, 15, False, 'muted', True)])
    steps(sl, .46, 3.65, 5.4, points, gap=.56, sz=14)
    d.box(sl, 6.15, 1.38, 6.72, 5.12, 'pri')
    if img.endswith('_m'):
        pass
    shot(sl, img, 6.42, 1.65, 6.18, 4.58)
    footer(sl, sec, f'Chương {no} · {name}')
    return sl

def N(sl, purpose, talk, transition=None, reasoning=None):
    d.notes(sl, purpose, talk, reasoning=reasoning, transition=transition)

# ═════════════════════════ BÌA & MỤC LỤC
s1 = d.cover_mmh('GIỚI THIỆU HỆ THỐNG · DÀNH CHO TOÀN CÔNG TY', 'HỆ THỐNG DỮ LIỆU', 'NỘI BỘ MMH', 'Product Team · 10/2026')
LOG.append((1, 'C01', 'Bìa'))
footer(s1, '1.1', 'Sách hướng dẫn sử dụng đầy đủ', y=6.45, w=6.3, credit=False)
N(s1, 'Mở đầu, nêu mục tiêu và cách bài trình bày đi theo cuốn hướng dẫn IFU.',
  ['Chào mọi người. Hôm nay mình giới thiệu hệ thống dữ liệu nội bộ của MMH: CRM, Report Hub, KPI và Training Hub.',
   'Bài đi theo đúng thứ tự cuốn hướng dẫn sử dụng IFU: nền tảng chung trước, rồi lần lượt từng app.',
   'Cuối mỗi slide có một dòng link tới đúng mục trong IFU. Ai cần làm theo từng bước thì mở link đó.',
   'Ảnh trong bài là ảnh chụp thật từ app, dữ liệu giả lập.'], 'Mình xem mục lục.')

s = S('Tổng quan', 'Nội dung', 'TỔNG QUAN', 'C03', '1.1', 'Sách hướng dẫn đầy đủ')
rows = [('1', 'Nền tảng chung', 'Đăng nhập, quyền, lưu dữ liệu, nhắc hạn'), ('2', 'MMH CRM', 'Công cụ hằng ngày của đội Sales'),
        ('3', 'MMH Report Hub', 'Việc của các phòng ban, lịch chung'), ('4', 'KPI', 'Số liệu tự điền, cách tính, cách xem'),
        ('5', 'Training Hub', 'Học, kiểm tra, chứng chỉ'), ('', 'Tổng kết', '6 điều cần nhớ')]
for i, (n, t, ds) in enumerate(rows):
    y = 1.5 + i * .82
    d.text(s, .6, y, .8, .5, [(n, 24, True, 'sec')], anchor='m')
    d.text(s, 1.5, y, 4.0, .5, [(t, 20, True, 'title')], anchor='m')
    d.text(s, 5.6, y, 7.0, .5, [(ds, 14, False, 'muted')], anchor='m')
    d.hline(s, .6, y + .7, 12.8)
N(s, 'Cho thấy mạch bài: 5 phần theo đúng 5 chương IFU, rồi tổng kết.',
  ['Bài có năm phần, trùng với năm chương của IFU.', 'Phần một là nền tảng chung cho mọi app. Phần hai là CRM, phần ba Report Hub, phần bốn KPI, phần năm Training Hub.',
   'Cuối cùng là sáu điều cần nhớ.'], 'Bắt đầu với phần nền tảng chung.')

# ═════════════════════════ PHẦN 1 · NỀN TẢNG
s = divider('1', 'Nền tảng chung', 'Phần chung của mọi app', ['4 app, 1 hệ thống', 'Dữ liệu đi đâu', 'Đăng nhập bằng email', 'Ai thấy gì', 'Lưu không phải chờ'], 'rh/main', '1.1')
N(s, 'Mở phần 1.', ['Phần một nói về những gì giống nhau ở mọi app: đăng nhập, quyền xem, cách lưu dữ liệu, nhắc hạn.', 'Hiểu phần này một lần thì các app sau dùng rất nhanh.'], 'Trước hết, hệ thống gồm những gì.')

s = S('1 · Nền tảng', '4 app, 1 hệ thống', '1 · NỀN TẢNG', 'C27', '1.1', 'Hệ thống gồm những gì')
apps = [('MMH CRM', 'Đội Sales', 'crm/cal_week'), ('Report Hub', 'Các phòng ban', 'rh/main'), ('KPI', 'Mọi người · quản lý', 'crm/kpi_personal'), ('Training Hub', 'Học viên · trainer', 'th/home')]
for i, (t, u, img) in enumerate(apps):
    x = .46 + i * 3.15
    shot(s, img, x, 1.45, 2.95, 1.85)
    d.text(s, x, 3.5, 2.95, .4, [(t, 18, True, 'title')])
    d.text(s, x, 3.92, 2.95, .3, [(u, 13, False, 'muted')])
d.box(s, .46, 4.6, 12.41, 1.75, 'soft')
for i, (lg, t, ds) in enumerate([('googlesheets', 'Google Sheets', 'dữ liệu gốc'), ('googleappsscript', 'Apps Script', 'cầu nối'), ('googledrive', 'Drive', 'tài liệu, ảnh'), ('gmail', 'Gmail', 'email tự động'), ('github', 'GitHub Pages', 'nơi đặt app')]):
    x = .75 + i * 2.45
    logo(s, lg, x, 4.95, .5)
    d.text(s, x + .62, 4.9, 1.8, .3, [(t, 13.5, True, 'title')])
    d.text(s, x + .62, 5.2, 1.8, .3, [(ds, 12, False, 'muted')])
d.text(s, .75, 5.75, 11.9, .4, [('Mọi app mở bằng trình duyệt, không cài đặt; dữ liệu nằm trong Google Sheets của công ty.', 13.5, False, 'ink')])
N(s, 'Giới thiệu 4 app và nền tảng Google bên dưới.',
  ['Hàng trên là bốn app: CRM cho Sales, Report Hub cho các phòng ban, KPI cho mọi người, Training Hub cho đào tạo.',
   'Hàng dưới là nền tảng: dữ liệu gốc nằm trong Google Sheets của công ty, Apps Script là cầu nối, Drive giữ tài liệu và ảnh, Gmail gửi email tự động, app đặt trên GitHub Pages.',
   'Mọi app mở bằng trình duyệt, không cần cài.'], 'Vậy dữ liệu đi qua đâu?')

s = S('1 · Nền tảng', 'Dữ liệu đi đâu', '1 · NỀN TẢNG', 'C19', '1.2', 'Dữ liệu nằm ở đâu, đi qua đâu')
chev(s, ['Người dùng', 'App web', 'Apps Script', 'Google Sheets'], y=1.6, h=.8,
     sub=['Thao tác trên máy tính hoặc điện thoại', 'Hiện kết quả ngay, gửi lệnh ở phía sau', 'Kiểm tra phiên đăng nhập và quyền, rồi ghi', 'Dữ liệu gốc của công ty; Drive, Gmail'], subh=1.0)
note(s, .46, 4.1, 12.41, 1.0, 'Nhập một lần.', 'Sales báo kết quả trên CRM ⇒ số lượt tự sang KPI. Buổi đào tạo xong ⇒ KPI đào tạo tự đếm. Công tác được duyệt ⇒ tự hiện trên lịch chung.')
d.text(s, .46, 5.4, 12.41, .8, [('Mã backend quản lý tập trung trong kho mmh-backend: sửa xong tự cập nhật, đường link app không đổi, lỗi thì tự quay về bản trước.', 13, False, 'muted')])
N(s, 'Giải thích đường đi của dữ liệu và nguyên tắc nhập một lần.',
  ['Đọc từ trái sang phải: người dùng thao tác trên app; app hiện kết quả ngay và gửi lệnh ở phía sau; Apps Script kiểm tra đăng nhập, quyền rồi ghi vào Google Sheets.',
   'Ô vàng là nguyên tắc quan trọng nhất: nhập một lần. Dữ liệu nhập ở app này tự xuất hiện ở app khác.'], 'Tiếp theo là cách đăng nhập.')

s = S('1 · Nền tảng', 'Đăng nhập bằng email', '1 · NỀN TẢNG', 'C30', '1.3', 'Đăng nhập bằng email công ty')
shot(s, 'rh/login_email', .46, 1.45, 3.5, 3.0, 'auto')
shot(s, 'rh/login_code', 4.15, 1.45, 3.5, 3.0, 'auto')
steps(s, 8.0, 1.5, 4.87, [('Nhập email công ty', '@manimedicalhanoi.com hoặc @mani.inc'), ('Lấy mã 6 số', 'trong email vừa nhận'), ('Nhập mã', 'app tự nhận tên, nhóm, quyền')], gap=.95)
d.table(s, .46, 4.75, [3.6, 3.6], [['Quy định', 'Giá trị'], ['Mã hết hạn', '10 phút'], ['Nhập sai', '5 lần'], ['Dùng được', '30 ngày / máy']], row_h=.4, head=12, body=12.5, aligns=['l', 'l'])
note(s, 8.0, 4.75, 4.87, 1.6, 'Một lần, mọi app.', 'Đăng nhập Report Hub rồi thì mở CRM, Product Data không phải đăng nhập lại.', fill='soft', line=None)
N(s, 'Hướng dẫn đăng nhập.',
  ['Không cần mật khẩu. Nhập email công ty, hệ thống gửi mã 6 số, nhập mã là vào.', 'Mã hết hạn sau 10 phút, sai 5 lần phải xin mã mới, một lần đăng nhập dùng 30 ngày trên máy đó.',
   'Đăng nhập một app thì các app khác trên cùng trình duyệt cũng nhận luôn.', 'Riêng Training Hub có đăng nhập bằng mã PIN, mình nói ở phần năm.'], 'Đăng nhập xong, mỗi người thấy gì?')

s = S('1 · Nền tảng', 'Ai thấy gì', '1 · NỀN TẢNG', 'C33', '1.4', 'Quyền xem và quyền sửa')
d.table(s, .46, 1.45, [2.6, 4.9, 4.91], [['Vai trò', 'Thấy', 'Làm thêm'], ['PIC / Sales', 'Việc, lịch, KPI của mình', 'Sửa việc của mình'],
        ['Team Leader', 'Cả nhóm', 'Duyệt một số mục của nhóm'], ['HOD', 'Cả phòng ban', 'Theo phân quyền'],
        ['Director', 'Toàn công ty, Total KPI', 'Duyệt công tác'], ['Admin', 'Toàn bộ', 'Phân quyền, chế độ thử, email cập nhật']],
        row_h=.62, head=13, body=13.5, styles={4: 'focus'}, aligns=['l', 'l', 'l'])
note(s, .46, 5.4, 12.41, .85, 'Mặc định:', 'chuyến công tác của người khác bị ẩn; buổi đào tạo chỉ hiện với trainer và người được mời.', fill='soft', line=None)
N(s, 'Giải thích quyền xem theo vai trò.',
  ['PIC chỉ thấy việc của mình; Team Leader thấy cả nhóm; HOD thấy phòng ban; Director thấy toàn công ty và duyệt công tác; Admin phân quyền.',
   'Admin có thể bật thêm hoặc tắt bớt từng quyền cho từng người.'], 'Tiếp theo là cách app lưu dữ liệu.')

s = S('1 · Nền tảng', 'Lưu không phải chờ', '1 · NỀN TẢNG', 'C18', '1.5', 'Lưu không phải chờ')
steps(s, .46, 1.5, 6.3, [('Bấm Lưu', 'khung đóng ngay, thay đổi hiện trên màn hình'), ('App gửi ở phía sau', 'lần lượt từng lệnh, ảnh gửi sau chữ'),
                          ('Mạng chập chờn', 'tự thử lại, không tạo bản ghi trùng'), ('Lỗi thật', 'chip đỏ trên thanh trên cùng: Thử lại / Bỏ')], gap=1.0)
shot(s, 'rh/outbox', 7.1, 1.5, 5.77, 3.83, 'auto')
note(s, .46, 5.6, 12.41, .7, 'Lưu ý:', 'còn chip đỏ thì không xoá dữ liệu trình duyệt; đóng tab vẫn giữ, lần mở sau app gửi tiếp.', fill='hi')
N(s, 'Giải thích cơ chế lưu không phải chờ.',
  ['Bấm Lưu là xong, khung đóng ngay. App tự gửi lên Sheet ở phía sau, mạng yếu thì tự thử lại.', 'Chỉ khi có lỗi thật, ví dụ dòng đã bị xoá trên Sheet, app mới hiện chip đỏ để mình chọn Thử lại hoặc Bỏ.'],
  'Cuối phần một: các nhắc hạn tự động.')

s = S('1 · Nền tảng', 'App tự nhắc hạn', '1 · NỀN TẢNG', 'C40', '1.6', 'Nhắc hạn, thông báo và hướng dẫn')
L = [('Thông báo cập nhật', '7 ngày sau mỗi bản mới'), ('Hạn chứng từ kế toán', 'ngày 14, 15, 28, 29 và ngày hạn'), ('Việc mới được giao', 'lần mở app đầu tiên'),
     ('Thư nhắc việc', 'việc sắp đến hạn, báo cáo chưa nộp'), ('Báo cáo KPI tháng', 'ngày 02, trước 16:00'), ('Tự đánh giá quý', 'ngày 09 tháng 9 · 12 · 3 · 6, trước 17:00')]
for i, (t, ds) in enumerate(L):
    y = 1.5 + i * .78
    d.badge(s, .46, y + .04, i + 1, d=.34, sz=12)
    d.text(s, 1.0, y, 3.3, .42, [(t, 14.5, True, 'title')], anchor='m')
    d.text(s, 4.3, y, 3.4, .42, [(ds, 13, False, 'ink')], anchor='m')
    d.hline(s, .46, y + .66, 7.7)
shot(s, 'rh/kpi_reminder_card', 8.1, 1.5, 4.77, 2.6, None, 'auto')
d.text(s, 8.1, 4.3, 4.77, 1.6, [('Popup hiện lần lượt theo thứ tự này, không chồng lên nhau. Nút Hướng dẫn (?) chỉ từng bước ngay trên màn hình.', 13, False, 'ink')])
N(s, 'Liệt kê các nhắc hạn tự động.',
  ['Sáu loại popup, hiện lần lượt theo thứ tự trên bảng.', 'Hai mốc quan trọng nhất cho mọi người: ngày 02 gửi báo cáo KPI tháng và ngày 09 đầu quý tự đánh giá.'], 'Sang phần hai: CRM.')

# ═════════════════════════ PHẦN 2 · CRM
s = divider('2', 'MMH CRM', 'Sổ tay số của đội Sales', ['Ai dùng, bố cục', 'Lịch và đi địa bàn', 'Khách hàng, đơn hàng', 'Mở mới & SKU mới', 'Báo cáo và quản lý'], 'crm/cal_week', '2.1')
N(s, 'Mở phần CRM.', ['CRM là công cụ hằng ngày của đội Sales: gặp khách, ghi kết quả, công tác, đơn hàng, khách hàng mới, báo cáo.'], 'CRM dùng cho ai?')

s = S('2 · CRM', 'Ai dùng CRM', '2 · CRM', 'C36', '2.1', 'Ai dùng CRM, mở ở đâu')
T = [('pri', 'Dental', 'Tiếng Việt'), ('sec', 'Surgical', 'Tiếng Việt'), ('ter', 'Eyeless', 'Tiếng Việt'), ('acc', 'Thailand Surgical', 'Tiếng Anh, không chụp ảnh')]
for i, (c, t, ds) in enumerate(T):
    x = .46 + i * 3.15
    d.box(s, x, 1.45, 2.95, 1.35, c)
    col = 'on_pri' if c in ('pri', 'sec') else 'title'
    d.text(s, x + .2, 1.55, 2.6, .5, [(t, 18, True, col)], anchor='m')
    d.text(s, x + .2, 2.08, 2.6, .55, [(ds, 12.5, False, col)])
d.text(s, .46, 2.95, 12.4, .4, [('Mỗi nhóm một file Google Sheet riêng: khách hàng, đơn hàng, báo cáo tách biệt.', 13.5, False, 'muted')])
d.table(s, .46, 3.6, [3.0, 9.41], [['Vai trò', 'Thấy gì'], ['Sales (PIC)', 'Lịch, khách hàng, đơn hàng, KPI của mình'], ['Team Leader', 'Cả nhóm: chip Cả đội và từng PIC'],
        ['Quản lý', 'Đổi nhóm đang xem, duyệt mở mới, duyệt công tác (Director)']], row_h=.55, head=13, body=13.5, aligns=['l', 'l'])
N(s, 'Giới thiệu 4 nhóm và quyền xem.', ['CRM có bốn nhóm, mỗi nhóm một file dữ liệu riêng. Nhóm Thái dùng toàn bộ tiếng Anh và không chụp ảnh đi địa bàn.',
   'Sales thấy của mình, Team Leader thấy cả nhóm, quản lý đổi được nhóm đang xem.'], 'Mở CRM ra, màn hình trông thế nào?')

s = S('2 · CRM', 'Bố cục màn hình', '2 · CRM', 'C24', '2.2', 'Bố cục màn hình')
shot(s, 'crm/cal_week', .46, 1.45, 7.9, 4.94)
G = [('Thanh trên', 'MMH Calendar · Tìm nhanh · Liên kết · FY68 · Có gì mới'), ('Tổng quan', 'Dashboard · KPI & Mục tiêu'),
     ('Dữ liệu', 'Khách hàng · Key stakeholder · Cập nhật · Mở mới · Đơn hàng · Thuyết trình · CV tháng · ADP'), ('Công cụ', 'Họp tuần · Xem ảnh · Gửi báo cáo')]
steps(s, 8.65, 1.5, 4.22, G, gap=1.2)
N(s, 'Đọc bố cục CRM.', ['Thanh trên có lịch chung, tìm nhanh Ctrl K, liên kết sang app khác và chọn năm tài chính.',
   'Menu trái chia ba nhóm: Tổng quan, Dữ liệu và Công cụ. Mục dùng nhiều nhất là Cập nhật, tức lịch làm việc.'], 'Mình xem lịch và cách thêm việc.')

s = S('2 · CRM', 'Lịch và thêm việc', '2 · CRM', 'C29', '2.3', 'Lịch làm việc và thêm việc')
shot(s, 'crm/cal_month', .46, 1.45, 6.1, 3.3)
shot(s, 'crm/add_menu', 6.77, 1.45, 6.1, 3.3, 'auto')
d.table(s, .46, 4.95, [3.0, 4.6, 4.81], [['Loại việc', 'Ghi vào', 'Khi nào'], ['Đi địa bàn', 'Tab 6. WEEKLY REPORT', 'Gặp khách, bác sĩ, demo'],
        ['Công việc khác', 'Sheet ALL TASK', 'Họp, hỗ trợ NPP, hồ sơ thầu'], ['Đề xuất công tác', 'File Business Trip + email', 'Chuyến đi xa cần duyệt']], row_h=.38, head=12, body=12.5, aligns=['l', 'l', 'l'])
N(s, 'Lịch tuần / tháng và 3 loại việc.', ['Lịch có kiểu Hôm nay, Tuần, Tháng, Danh sách; mỗi loại việc một màu.',
   'Bấm Thêm việc sẽ có ba lựa chọn: đi địa bàn, công việc khác, đề xuất công tác. Mỗi loại ghi vào đúng chỗ của nó.'], 'Quan trọng nhất là đi địa bàn.')

s = S('2 · CRM', 'Đi địa bàn: 4 bước', '2 · CRM', 'C20', '2.4', 'Đi địa bàn, báo kết quả, chụp ảnh')
chev(s, ['Tạo lượt', 'Gặp khách', 'Báo kết quả', 'Chụp ảnh'], y=1.45, h=.65,
     sub=['Account, ngày, loại việc, kế hoạch', 'Hồ sơ 360° để xem lịch sử', 'Người gặp + mẫu 4 mục, trạng thái', 'Ngay tại địa bàn, kèm giờ và vị trí'], subh=.7)
shot(s, 'crm/fieldvisit_form', .46, 3.05, 6.1, 3.35, 'auto')
shot(s, 'crm/fieldvisit_result', 6.77, 3.05, 6.1, 3.35, 'auto')
N(s, 'Luồng đi địa bàn.', ['Bốn bước: tạo lượt, gặp khách, báo kết quả, chụp ảnh.', 'Khung báo kết quả có nút chọn người gặp, nút chèn khung 4 mục và nút nói thay vì gõ. Dòng người gặp là bắt buộc.',
   'Với Dental, Surgical, Eyeless, lượt viếng thăm cần ảnh chụp ngay tại địa bàn.'], 'Trên điện thoại còn nhanh hơn.')

s = S('2 · CRM', 'Trên điện thoại', '2 · CRM', 'C19', '2.11', 'Trên điện thoại')
for i, (sid, t) in enumerate([('crm/m_today', 'Hôm nay'), ('crm/m_photo', 'Chụp ảnh'), ('crm/m_result', 'Báo kết quả'), ('crm/m_more', 'Mọi chức năng')]):
    x = .9 + i * 3.05
    phone(s, sid, x + .25, 1.45, 4.1)
    d.badge(s, x, 5.75, i + 1, d=.34, sz=12)
    d.text(s, x + .45, 5.72, 2.4, .4, [(t, 14.5, True, 'title')])
N(s, 'Thao tác trên điện thoại.', ['Thanh dưới có Lịch, Khách hàng, Hôm nay, Mở mới và Thêm.', 'Trên thẻ đi địa bàn có sẵn hai nút Chụp ảnh và Báo kết quả, nên Sales ghi ngay sau buổi gặp.'],
  'Tiếp theo là công tác.')

s = S('2 · CRM', 'Công tác', '2 · CRM', 'C26', '2.5', 'Công việc khác và công tác')
chev(s, ['Đề xuất', 'Giám đốc duyệt', 'Đi công tác', 'Báo cáo'], y=1.45, h=.65,
     sub=['Ngày, nơi đến, mục đích, chi phí', 'Email phê duyệt / từ chối', 'Hiện trên lịch, nhãn CÔNG TÁC', 'Hoạt động, phát hiện, việc tiếp theo'], subh=.7)
shot(s, 'crm/trip_propose', .46, 3.05, 6.1, 3.35, 'auto')
shot(s, 'crm/trip_report', 6.77, 3.05, 6.1, 3.35, 'auto')
N(s, 'Quy trình công tác.', ['Đề xuất trên app, hệ thống ghi vào file Business Trip và gửi email cho Giám đốc.', 'Giám đốc duyệt ngay trên app. Sau chuyến đi, Sales viết báo cáo ngay trên thẻ công tác.'],
  'Giờ sang dữ liệu khách hàng.')

s = S('2 · CRM', 'Khách hàng 360°', '2 · CRM', 'C07', '2.6', 'Khách hàng, người liên hệ, ADP')
shot(s, 'crm/cust_profile', .46, 1.45, 8.3, 4.9, (.05, .0, .95, .8))
d.box(s, 9.05, 1.45, 3.82, 4.9, 'panel')
for i, (ic, t, ds) in enumerate([('search', 'Danh sách', 'tìm theo tên, mã, tỉnh, PIC; Target → Achieve'), ('eye', 'Hồ sơ 360°', 'dòng thời gian mọi lượt gặp'),
                                  ('users', 'Người liên hệ', 'chức vụ, ảnh hưởng, tin cậy'), ('target', 'ADP', 'khách trọng điểm A · B · C')]):
    y = 1.7 + i * 1.12
    d.icon(s, ic, 9.25, y, .46)
    d.text(s, 9.85, y - .04, 2.9, .32, [(t, 14.5, True, 'title')])
    d.text(s, 9.85, y + .3, 2.9, .7, [(ds, 12, False, 'ink')])
N(s, 'Danh sách khách hàng và hồ sơ 360 độ.', ['Bên trái là danh sách account, lọc theo loại, khu vực, nhóm ADP.', 'Bên phải là hồ sơ 360 độ: tổng quan, biểu đồ lượt gặp, người liên hệ và toàn bộ dòng thời gian làm việc với khách.'],
  'Đơn hàng cũng nằm trong CRM.')

s = S('2 · CRM', 'Đơn hàng', '2 · CRM', 'C10', '2.7', 'Đơn hàng, thuyết trình, công việc trong tháng')
shot(s, 'crm/orders', .46, 1.45, 7.6, 4.9, (.16, .02, 1, .8))
for i, (v, t) in enumerate([('C5', 'Tổng tiền đơn ÷ tỷ giá tháng'), ('C6', 'Số lượng Composite qua NPP'), ('C7', 'Số lượng JIZAI qua NPP')]):
    y = 1.45 + i * 1.25
    d.kpi(s, 8.35, y, 4.52, 1.1, v, t, vsz=26, dark=(i == 0))
note(s, 8.35, 5.25, 4.52, 1.1, 'Chống trùng:', 'mạng chậm bấm lại cũng không ghi đơn hai lần.', fill='soft', line=None)
N(s, 'Đơn hàng và 3 mã KPI tính từ đơn.', ['Đơn hàng nhập trên CRM là nguồn tính ba mã KPI: C5 tổng tiền, C6 số lượng Composite, C7 số lượng JIZAI qua nhà phân phối phụ trách.'], 'Phần tiếp theo: khách hàng mới.')

s = S('2 · CRM', 'Mở mới & SKU mới', '2 · CRM', 'C35', '2.8', 'Mở mới địa bàn & SKU mới')
shot(s, 'crm/newacc_review', .46, 1.45, 7.2, 4.9, 'auto')
d.text(s, 7.95, 1.45, 4.9, .3, [('4 ĐIỀU KIỆN TÍNH KPI C1', 11, True, 'sec')])
steps(s, 7.95, 1.85, 4.92, ['Có trong báo cáo tuần', 'NPP xác nhận', 'Thông tin đúng mẫu', 'Có đơn hàng đầu tiên'], gap=.5, sz=13.5)
d.text(s, 7.95, 4.0, 4.9, .3, [('QUẢN LÝ CHỌN', 11, True, 'sec')])
for i, (t, c, ds) in enumerate([('Xác nhận', 'pos', 'tính KPI'), ('Trả lại bổ sung', 'acc', 'về nháp, giữ chứng từ'), ('Từ chối', 'neg', 'ghi lý do')]):
    y = 4.4 + i * .62
    d.chip(s, 7.95, y, t, fill=c, col='on_pri' if c != 'acc' else 'title', sz=12, h=.36)
    d.text(s, 10.0, y, 2.87, .36, [(ds, 13, False, 'ink')], anchor='m')
N(s, 'Quy trình mở mới.', ['Sales tạo ca và gửi email xin xác nhận. App tự đối chiếu bốn điều kiện.', 'Quản lý có ba nút: Xác nhận để tính KPI, Trả lại để Sales bổ sung, hoặc Từ chối có ghi lý do.'],
  'Cuối tuần, Sales họp và gửi báo cáo.')

s = S('2 · CRM', 'Họp tuần, báo cáo', '2 · CRM', 'C18', '2.10', 'Họp tuần, gửi báo cáo, xem ảnh')
for i, (ic, t, ds) in enumerate([('chart', 'Forecast NPP', 'số liệu tự lấy từ file Turnover'), ('users', 'Quan hệ KH & thị trường', 'theo từng PIC'),
                                  ('file', 'Tình hình thầu', 'gói thầu, giá theo mã sản phẩm'), ('alert', 'Gửi báo cáo', 'tuần · tháng · KPI; xem trước, gửi thử, tải PDF')]):
    y = 1.55 + i * 1.18
    d.icon(s, ic, .46, y, .5)
    d.text(s, 1.1, y - .04, 4.3, .34, [(t, 15, True, 'title')])
    d.text(s, 1.1, y + .32, 4.3, .6, [(ds, 12.5, False, 'ink')])
shot(s, 'crm/weekly_meeting', 5.6, 1.45, 7.27, 4.9, (.16, .03, 1, .9))
N(s, 'Họp tuần và gửi báo cáo.', ['Biên bản họp tuần có bốn phần, số liệu forecast tự lấy từ file doanh số.', 'Gửi báo cáo: chọn kỳ, xem trước đúng email sẽ gửi, rồi gửi thử hoặc gửi chính thức.'], 'Phần dành cho quản lý.')

s = S('2 · CRM', 'Dành cho quản lý', '2 · CRM', 'C28', '2.12', 'Dành cho quản lý và Admin')
for i, (sid, t, ds, bx) in enumerate([('crm/team_switch', 'Đổi nhóm', 'Dental · Surgical · Eyeless · Thái', (.18, .2, .82, .72)),
                                       ('crm/trip_approve', 'Duyệt công tác', 'Giám đốc, chọn nhiều chuyến', (.06, .04, .8, .92)),
                                       ('crm/test_mode', 'Chế độ thử', 'Admin dùng như một Sales', (.22, .25, .78, .85))]):
    x = .46 + i * 4.17
    shot(s, sid, x, 1.45, 4.0, 3.6, bx, 'mid')
    d.text(s, x, 5.2, 4.0, .4, [(t, 16, True, 'title')])
    d.text(s, x, 5.6, 4.0, .4, [(ds, 13, False, 'muted')])
N(s, 'Ba công cụ cho quản lý.', ['Quản lý đổi nhóm đang xem. Giám đốc duyệt nhiều chuyến công tác một lần, mỗi chuyến một email riêng.', 'Admin có chế độ thử: dùng CRM như một bạn Sales nhưng không ghi vào dữ liệu thật.'], 'Cuối phần CRM: dữ liệu đi đâu.')

s = S('2 · CRM', 'Dữ liệu CRM đi đâu', '2 · CRM', 'C21', '2.13', 'Dữ liệu và kết nối')
chev(s, ['Sales ghi trên CRM', 'File CRM nhóm', 'Báo cáo email', 'File KPI FY68'], y=1.5, h=.75,
     sub=['Đi địa bàn, đơn hàng, mở mới, công tác', 'Tab 1 → 16, ALL TASK', 'Tuần, tháng gửi Leader, quản lý', 'C1, C3, C5, C6, C7 tự cập nhật'], subh=.8)
d.table(s, .46, 3.4, [4.0, 8.41], [['Kết nối', 'Đi tới'], ['Đề xuất công tác', 'File Business Trip ⇒ lịch chung MMH Calendar'], ['Doanh số NPP, forecast', 'Đọc từ Turnover Report (kế toán)'],
        ['Khách từ sự kiện', 'Đồng bộ từ file Data khách hàng'], ['Target KPI', 'Từ file KPI tự hiện trong CRM']], row_h=.52, head=13, body=13.5, aligns=['l', 'l'])
N(s, 'Chốt phần CRM bằng luồng dữ liệu.', ['Sales chỉ nhập trên CRM. Dữ liệu vào file của nhóm, báo cáo tự gửi email, các mã KPI tự cập nhật.', 'Ngược lại, Target trong file KPI tự hiện trong CRM.'],
  'Sang phần ba: Report Hub.')

# ═════════════════════════ PHẦN 3 · REPORT HUB
s = divider('3', 'MMH Report Hub', 'Bảng việc chung của các phòng ban', ['Key task, sub task', 'Lịch, báo cáo tuần', 'MMH Calendar', 'Giao việc, công tác', 'Phân quyền'], 'rh/work_cal_week', '3.1')
N(s, 'Mở phần Report Hub.', ['Report Hub là nơi các phòng ban cập nhật việc, tự tạo báo cáo, xem lịch chung của công ty.'], 'Ba phòng ban, một cách làm.')

s = S('3 · Report Hub', '3 phòng ban, 1 cách làm', '3 · REPORT HUB', 'C36', '3.1', 'Ai dùng, bố cục màn hình')
for i, (c, t, ds) in enumerate([('pri', 'Marketing & Sales', 'Product, Event, Digital, Design'), ('sec', 'Back Office', 'Kho, kế toán, nhập hàng'), ('ter', 'Management', 'Việc ban quản lý')]):
    x = .46 + i * 4.17
    d.box(s, x, 1.45, 4.0, 1.2, c)
    col = 'on_pri' if c != 'ter' else 'title'
    d.text(s, x + .22, 1.52, 3.6, .5, [(t, 17, True, col)], anchor='m')
    d.text(s, x + .22, 2.02, 3.6, .5, [(ds, 12.5, False, col)])
shot(s, 'rh/dept_backoffice', .46, 2.9, 7.6, 3.45, (0, 0, 1, .7))
steps(s, 8.4, 2.95, 4.47, ['Cập nhật key task / sub task', 'Báo cáo tuần tự tạo', 'Lịch làm việc, biên bản họp', 'Tìm nhanh Ctrl K'], gap=.78, sz=14)
N(s, 'Ba phòng ban.', ['Marketing và Sales, Back Office, Management. Mỗi phòng một file dữ liệu, nhưng cùng một cách cập nhật.'], 'Cách cập nhật đó là key task và sub task.')

s = S('3 · Report Hub', 'Key task và sub task', '3 · REPORT HUB', 'C24', '3.2', 'Cập nhật công việc')
shot(s, 'rh/task_list', .46, 1.45, 7.4, 4.9, (.18, .0, 1, .75))
steps(s, 8.15, 1.5, 4.72, [('Key task', 'việc lớn: loại việc, PIC, hạn'), ('Sub task', 'bước nhỏ: PIC, hạn, % tiến độ'), ('Cập nhật', 'trạng thái, kết quả, link, file'), ('Báo người nhận', 'tick soạn email khi thêm việc')], gap=1.15)
N(s, 'Hai tầng công việc.', ['Mỗi việc lớn là một key task, chia nhỏ thành sub task có người phụ trách, hạn và phần trăm tiến độ.', 'Cập nhật trạng thái và kết quả ngay trên danh sách hoặc trong khung chi tiết.'], 'Việc hiển thị trên lịch làm việc.')

s = S('3 · Report Hub', 'Lịch làm việc', '3 · REPORT HUB', 'C07', '3.3', 'Lịch làm việc, Pending, Today task')
shot(s, 'rh/work_cal_week', .46, 1.45, 8.3, 4.9, (.15, .0, 1, .9))
d.box(s, 9.05, 1.45, 3.82, 4.9, 'panel')
d.text(s, 9.3, 1.6, 3.4, .3, [('CÁCH XEM', 11, True, 'sec')])
bullets(s, 9.3, 1.95, 3.4, ['Today task · Tuần · Tháng · Danh sách', 'Chỉ tôi / Cả đội', 'Bật / tắt công tác, đào tạo'], 13, 6)
d.text(s, 9.3, 3.75, 3.4, .3, [('PENDING', 11, True, 'sec')])
bullets(s, 9.3, 4.1, 3.4, ['Trễ hạn · đến hạn · đang làm', 'Nút Hoàn thành, Gia hạn'], 13, 6)
N(s, 'Lịch tuần và Pending.', ['Lịch tuần hiện sub task theo ngày, có nhãn quá hạn, đến hạn.', 'Mục Pending gom việc trễ, việc đến hạn hôm nay, ngày mai và việc đang làm, mỗi thẻ có nút Hoàn thành hoặc Gia hạn.'],
  'Cập nhật đủ thì báo cáo tuần tự có.')

s = S('3 · Report Hub', 'Báo cáo tuần tự tạo', '3 · REPORT HUB', 'C42', '3.5', 'Báo cáo tuần, Monthly Report, biên bản họp')
for i, (lb, t) in enumerate([('KẾT QUẢ TUẦN', 'Task có ngày bắt đầu, hạn hoặc hoàn thành trong tuần + task quá hạn chưa xong'),
                              ('PLAN TUẦN SAU', 'Task chưa xong chạm vào tuần sau, kể cả việc quá hạn gánh sang')]):
    y = 1.45 + i * 1.6
    d.box(s, .46, y, 6.3, 1.45, 'panel')
    d.text(s, .7, y + .15, 5.8, .28, [(lb, 11, True, 'sec')])
    d.text(s, .7, y + .48, 5.85, .9, [(t, 15, False, 'title')])
note(s, .46, 4.75, 6.3, 1.6, 'Gửi:', 'gửi thử cho chính mình, hoặc gửi Team Leader, HOD và đồng nghiệp; email kèm PDF.', fill='soft', line=None)
shot(s, 'rh/weekly_send', 7.1, 1.45, 5.77, 4.9, None, 'auto')
N(s, 'Báo cáo tuần lấy việc nào.', ['Không cần gõ lại. Kết quả tuần là mọi task chạm vào tuần cộng việc quá hạn; Plan tuần sau là task chưa xong chạm vào tuần sau.', 'Chọn tuần, chọn gửi thử hoặc gửi chính thức.'],
  'Tiếp theo là lịch chung của công ty.')

s = S('3 · Report Hub', 'MMH Calendar', '3 · REPORT HUB', 'C33', '3.6', 'MMH Calendar')
shot(s, 'rh/mmh_cal_month', .46, 1.45, 7.4, 4.9, (0, 0, 1, .85))
d.table(s, 8.1, 1.45, [1.9, 2.87], [['Loại lịch', 'Nguồn'], ['Công tác', 'File Business Trip'], ['Đào tạo', 'Training Hub'], ['Sự kiện offline', 'Marketing Offline'], ['Bài đăng', 'Digital Marketing']],
        row_h=.55, head=12.5, body=13, aligns=['l', 'l'])
d.text(s, 8.1, 4.45, 4.77, 1.8, [('Nhấp đúp một ngày để tạo mới; bấm một mục để xem, sửa ngay. Có ở cả Report Hub và CRM.', 13, False, 'ink')])
N(s, 'Lịch chung bốn nguồn.', ['MMH Calendar gom bốn loại lịch: công tác, đào tạo, sự kiện offline và bài đăng.', 'Mỗi loại lấy từ file gốc của nó, nên mở một màn hình là thấy hết, tránh trùng lịch.'], 'Từ lịch tạo mới được năm loại.')

s = S('3 · Report Hub', 'Tạo mới từ lịch', '3 · REPORT HUB', 'C18', '3.6', 'MMH Calendar · Tạo mới')
shot(s, 'rh/cal_new', .46, 1.45, 5.4, 4.9, (.38, .48, .62, .98), 'mid')
for i, (ic, t, ds) in enumerate([('wallet', 'Đề xuất công tác', 'File Business Trip, email xin duyệt'), ('book', 'Kế hoạch đào tạo', 'Training Hub, thư mời, tài liệu'),
                                  ('cal', 'Sự kiện offline', 'REPORT_EVENT'), ('file', 'Bài đăng sự kiện', 'EVENT REPORT'), ('box', 'Bài đăng sản phẩm', 'PRODUCT REPORT')]):
    y = 1.5 + i * .95
    d.icon(s, ic, 6.25, y, .5)
    d.text(s, 6.9, y - .03, 6.0, .32, [(t, 15, True, 'title')])
    d.text(s, 6.9, y + .3, 6.0, .3, [(ds, 12.5, False, 'muted')])
N(s, 'Năm lựa chọn khi tạo mới.', ['Nhấp đúp một ngày trên lịch, chọn một trong năm loại. Mỗi loại ghi thẳng vào đúng file nguồn.'], 'Sửa sự kiện và bài đăng cũng ngay trên lịch.')

s = S('3 · Report Hub', 'Sửa ngay trên lịch', '3 · REPORT HUB', 'C29', '3.7', 'Sự kiện offline và bài đăng Marketing')
shot(s, 'rh/offline_edit', .46, 1.45, 6.1, 4.0, (0, 0, 1, .75))
shot(s, 'rh/online_edit', 6.77, 1.45, 6.1, 4.0, (0, 0, 1, .75))
d.text(s, .46, 5.6, 6.1, .6, [('Sự kiện offline · ghi vào REPORT_EVENT', 14, True, 'title')])
d.text(s, 6.77, 5.6, 6.1, .6, [('Bài đăng · ghi vào EVENT REPORT', 14, True, 'title')])
N(s, 'Sửa không cần mở Sheet.', ['Bấm Sửa trên một sự kiện hay bài đăng, khung sửa mở ngay bên phải. App xác nhận đúng dòng và không ghi vào cột có công thức.'], 'Tiếp theo là giao việc.')

s = S('3 · Report Hub', 'Giao việc: 3 bước', '3 · REPORT HUB', 'C20', '3.10', 'Giao việc và trao đổi')
chev(s, ['Giao', 'Nhận', 'Hoàn thành'], y=1.45, h=.65, sub=['Thêm sub task, tick soạn email', 'Popup "Việc mới được giao"', 'Email tự báo người giao'], subh=.6)
for i, (sid, bx) in enumerate([('rh/assign_mail', None), ('rh/assign_popup', (.28, .3, .74, .72)), ('rh/assign_track', None)]):
    shot(s, sid, .46 + i * 4.17, 3.0, 4.0, 3.35, bx, 'mid' if bx else 'top')
N(s, 'Giao việc có theo dõi.', ['Người giao thêm sub task và soạn email ngay. Người nhận thấy popup khi mở app. Xong việc, hệ thống tự email báo người giao.', 'Mục Việc đã giao cho biết ai đã xem, đang làm hay hoàn thành.'],
  'Công tác trên Report Hub giống CRM.')

s = S('3 · Report Hub', 'Công tác', '3 · REPORT HUB', 'C34', '3.9', 'Công tác')
d.table(s, .46, 1.45, [2.4, 3.4, 3.3, 3.31], [['', 'Trước chuyến', 'Duyệt', 'Sau chuyến'], ['Người đi', 'Điền đề xuất, chi phí, lịch trình', 'Nhận email kết quả', 'Báo cáo trong 3 ngày'],
        ['Giám đốc', 'Nhận email xin duyệt', 'Duyệt / Từ chối trên app', 'Nhận báo cáo'], ['Hệ thống', 'Ghi file Business Trip', 'Ghi nhật ký', 'Hiện trên MMH Calendar']],
        row_h=1.15, head=13, body=14, styles={2: 'focus'}, aligns=['l', 'l', 'l', 'l'])
N(s, 'Vai trò từng bên trong công tác.', ['Người đi đề xuất và báo cáo trong ba ngày sau chuyến. Giám đốc duyệt trên app. Hệ thống ghi file Business Trip và hiện chuyến đi trên lịch chung.'], 'Cuối phần ba là phân quyền.')

s = S('3 · Report Hub', 'Phân quyền (Admin)', '3 · REPORT HUB', 'C40', '3.11', 'Nhắc việc, thông báo, quản trị')
shot(s, 'rh/permissions', .46, 1.45, 6.6, 4.9, (0, 0, 1, .9))
P = ['Giao việc & email giao việc', 'Báo cáo tuần / tháng', 'Đề xuất & báo cáo công tác', 'Tạo buổi đào tạo', 'Sửa dữ liệu Marketing', 'Xem Management', 'Product Data: sửa, upload']
for i, t in enumerate(P):
    y = 1.5 + i * .66
    d.icon(s, 'check', 7.4, y, .38, 'pri', circle=False)
    d.text(s, 7.9, y - .02, 4.97, .42, [(t, 14, False, 'ink')], anchor='m')
N(s, 'Bảy quyền Admin bật / tắt.', ['Admin mở menu tài khoản, chọn Phân quyền người dùng. Mỗi người có bảy quyền có thể bật hoặc tắt riêng, ngoài quyền mặc định theo vai trò.'], 'Sang phần bốn: KPI.')

# ═════════════════════════ PHẦN 4 · KPI
s = divider('4', 'KPI', 'Tự điền · tự liên kết · hiển thị', ['3 lớp của KPI', '6 nguồn tự điền', 'Tự động và nhập tay', '% và màu', 'Xem và báo cáo'], 'rh/kpi_tab', '4.1')
N(s, 'Mở phần KPI.', ['Phần KPI trả lời: số lấy từ đâu, máy làm gì, người làm gì, tính thế nào và xem ở đâu.'], 'Bức tranh chung trước.')

s = S('4 · KPI', '3 lớp của KPI', '4 · KPI', 'C26', '4.1', 'Bức tranh chung')
chev(s, ['Nguồn', 'File KPI chung', 'Hiển thị'], y=1.6, h=.85, sub=['File làm việc hằng ngày: CRM, Turnover, Training, Marketing, KBI', 'MMH KPI FY68: Target, Actual, trọng số, % đạt', 'Tab KPI trong CRM, Report Hub; email KPI tháng'], subh=1.0)
for i, (v, t) in enumerate([('0', 'lần chép số tay sang file KPI'), ('1', 'nơi đặt Target, dùng chung mọi app'), ('Chỉ đọc', 'app không sửa KPI, sửa ở file nguồn')]):
    d.kpi(s, .46 + i * 4.17, 4.0, 4.0, 1.9, v, t, vsz=34, dark=(i == 0))
N(s, 'Ba lớp của KPI.', ['Lớp nguồn là các file mọi người làm hằng ngày. Lớp giữa là file KPI chung. Lớp trên là màn hình hiển thị.', 'Không ai chép số tay; Target chỉ đặt một lần; app chỉ đọc.'], 'Sáu nguồn cụ thể là gì?')

s = S('4 · KPI', '6 nguồn tự điền', '4 · KPI', 'C33', '4.3', 'Số liệu đến từ đâu')
d.table(s, .46, 1.45, [3.0, 6.6, 2.81], [['Nguồn', 'Máy tự điền', 'Mã KPI'], ['Turnover Report', 'Doanh số NPP theo NPP, mảng, nước', 'F'], ['4 file CRM', 'Mở mới, viếng thăm, hỗ trợ bán, thuyết trình', 'C1 · C3 · C5 · C6 · C7'],
        ['Training Master', 'Buổi đào tạo Completed', 'L1 · L2 · C10-06'], ['Marketing Online', 'Quảng cáo, bài đăng, thiết kế', 'C17'], ['Marketing Offline', 'Sự kiện, người tham dự, số bán', 'C8 · C9 · C12 · C15'],
        ['Điểm hành vi (KBI)', '100 + điểm cộng + điểm trừ', 'L3']], row_h=.6, head=13, body=13.5, styles={2: 'focus'}, aligns=['l', 'l', 'l'])
d.text(s, .46, 5.85, 12.41, .45, [('Máy chỉ ghi cột Actual, không bao giờ ghi Target.', 13.5, True, 'title')])
N(s, 'Sáu nguồn và mã KPI mỗi nguồn điền.', ['Dòng tô vàng là bốn file CRM, liên quan trực tiếp tới Sales.', 'Máy chỉ ghi cột Actual. Target luôn do người đặt.'], 'Số tự cập nhật nhanh cỡ nào?')

s = S('4 · KPI', 'Tốc độ cập nhật', '4 · KPI', 'C10', '4.4', 'Tự điền diễn ra thế nào')
for i, (v, t, sub) in enumerate([('5-20 giây', 'Sửa tay file nguồn', 'chạy ngay khi file được sửa'), ('1 phút', 'Web app ghi dữ liệu', 'CRM đẩy KPI ngay sau khi lưu'),
                                 ('15 phút', 'Đọc lại toàn bộ', 'để không sót'), ('Mọi tháng', 'Nhập trễ vẫn cập nhật', 'không khoá số tháng cũ')]):
    x = .46 + (i % 2) * 6.27; y = 1.45 + (i // 2) * 2.45
    d.kpi(s, x, y, 6.12, 2.2, v, t, sub, vsz=34, dark=(i == 0))
N(s, 'Tốc độ cập nhật.', ['Sửa tay file nguồn thì 5 tới 20 giây sau số KPI đổi. Dữ liệu từ app được kiểm tra mỗi phút. Cứ 15 phút đọc lại toàn bộ.', 'Nhập trễ cho tháng cũ vẫn được cập nhật.'], 'Phần nào vẫn phải nhập tay?')

s = S('4 · KPI', 'Máy làm, người làm', '4 · KPI', 'C29', '4.5', 'Phần nào còn nhập tay')
for x, t, col, its, ic in ((.46, 'Tự động', 'pri', ['Actual của mã có nguồn', 'Target hiện trong CRM', '% đạt, tổng tháng / quý / năm', 'Tên KPI rút gọn tiếng Việt'], 'check'),
                           (6.75, 'Nhập tay', 'sec', ['Target từng mã', 'Trọng số theo quý', 'Actual chưa có nguồn (Back Office)', 'Review cuối quý'], 'file')):
    d.box(s, x, 1.45, 6.12, .55, col, adj=.2)
    d.text(s, x + .22, 1.45, 5.8, .55, [(t, 17, True, 'on_pri')], anchor='m')
    for i, it in enumerate(its):
        y = 2.3 + i * .85
        d.icon(s, ic, x + .1, y, .45, col)
        d.text(s, x + .7, y, 5.3, .45, [(it, 15, False, 'ink')], anchor='m')
note(s, .46, 5.75, 12.41, .6, 'Kết luận:', 'làm đúng việc của mình trên app, số KPI tự đi theo.', fill='hi')
N(s, 'Phân chia máy và người.', ['Máy điền Actual, tính phần trăm và tổng. Người đặt Target, trọng số và nhập phần chưa có nguồn, chủ yếu KPI Back Office.'], 'Tính phần trăm và tô màu thế nào?')

s = S('4 · KPI', '% và màu', '4 · KPI', 'C16', '4.6', 'Cách tính % và màu')
d.box(s, .46, 1.45, 5.6, 2.2, 'soft')
d.text(s, .7, 1.6, 5.2, .5, [('% đạt = Actual ÷ Target', 22, True, 'title')])
d.text(s, .7, 2.25, 5.2, .5, [('Tổng KPI = Σ (% × trọng số)', 22, True, 'title')])
d.text(s, .7, 2.95, 5.2, .5, [('Q1 = T9 to T11 · Q2 = T12 to T2 · Q3 = T3 to T5 · Q4 = T6 to T8', 12, False, 'muted')])
bullets(s, .46, 4.0, 5.6, ['Target = 0 (kỳ đã tới) ⇒ 100%', 'CRM: tối đa 130% / mã (trừ doanh số F)', 'Chưa có số ⇒ "Chưa có số"'], 13.5, 8)
for i, (t, c, v, ds) in enumerate([('Dưới 90%', 'A9B0B7', 72, 'Chưa đạt'), ('90% đến < 100%', '3A5CAA', 95, 'Gần đạt'), ('100% đến 120%', 'E9B800', 110, 'Đạt'), ('Trên 120%', '3E8E6B', 128, 'Vượt')]):
    y = 1.6 + i * 1.15
    d.text(s, 6.6, y, 2.6, .35, [(t, 15, True, 'title')])
    d.text(s, 6.6, y + .36, 2.6, .3, [(ds, 12.5, False, 'muted')])
    d.box(s, 9.2, y + .1, 3.0, .34, 'panel', adj=.5)
    d.box(s, 9.2, y + .1, 3.0 * v / 130, .34, c, adj=.5)
    d.text(s, 12.25, y + .03, .62, .45, [(f'{v}%', 13, True, 'title')], 'r', 'm')
N(s, 'Công thức và bốn màu.', ['Phần trăm đạt bằng Actual chia Target; tổng KPI cộng phần trăm nhân trọng số.', 'Bốn màu: xám dưới 90, xanh dương 90 tới dưới 100, vàng 100 tới 120, xanh lá trên 120. Các thanh là ví dụ.'], 'Xem KPI ở đâu?')

s = S('4 · KPI', 'Xem KPI trên app', '4 · KPI', 'C24', '4.7', 'Xem KPI trên app')
shot(s, 'rh/kpi_tab', .46, 1.45, 7.9, 4.9, (.06, .0, 1, .9))
d.table(s, 8.65, 1.45, [1.9, 2.32], [['Vai trò', 'Thấy KPI của'], ['PIC', 'Mình'], ['Team Leader', 'Cả nhóm'], ['HOD', 'Phòng ban'], ['Director', 'Toàn công ty']],
        row_h=.55, head=12.5, body=13.5, aligns=['l', 'l'])
d.text(s, 8.65, 4.45, 4.22, 1.8, [('Có ở Report Hub (KPI) và CRM (KPI & Mục tiêu). Chọn Tháng / Quý / Năm; bấm tên để xem thành viên.', 13, False, 'ink')])
N(s, 'Màn KPI.', ['Bốn ô trên cùng là KPI tháng, quý, năm và số mã đạt. Dưới là xu hướng, đóng góp theo nhóm và bảng chi tiết.', 'Mỗi vai trò thấy đúng phạm vi của mình.'], 'Hai mốc báo cáo KPI.')

s = S('4 · KPI', 'Báo cáo tháng, đánh giá quý', '4 · KPI', 'C25', '4.8', 'Báo cáo KPI tháng và tự đánh giá quý')
shot(s, 'crm/kpi_report', .46, 1.45, 6.1, 3.5, 'auto')
shot(s, 'crm/kpi_se', 6.77, 1.45, 6.1, 3.5, 'auto')
d.text(s, .46, 5.1, 6.1, .4, [('Ngày 02 · trước 16:00', 16, True, 'sec')])
d.text(s, .46, 5.5, 6.1, .8, [('Báo cáo KPI tháng: To / CC tự điền, xem trước rồi gửi', 13, False, 'ink')])
d.text(s, 6.77, 5.1, 6.1, .4, [('Ngày 09 tháng 9 · 12 · 3 · 6 · trước 17:00', 16, True, 'sec')])
d.text(s, 6.77, 5.5, 6.1, .8, [('Tự đánh giá quý: KPI điền sẵn, tải file Word', 13, False, 'ink')])
N(s, 'Hai mốc KPI.', ['Ngày 02 hằng tháng gửi báo cáo KPI tháng trước, người nhận tự điền.', 'Ngày 09 tháng đầu quý làm tự đánh giá quý: KPI điền sẵn vào mẫu Word, chỉ viết phần đánh giá.'], 'Cuối phần KPI: Total KPI.')

s = S('4 · KPI', 'Total KPI', '4 · KPI', 'C09', '4.9', 'Total KPI cho quản lý')
shot(s, 'rh/total_kpi', .46, 1.45, 8.2, 4.9, (0, 0, 1, .95))
for i, (t, ds) in enumerate([('Tổng quan', '4 góc F · C · P · L'), ('KPI công ty', 'Tháng · lũy kế · cả năm'), ('Thành viên', '% từng người, tô màu')]):
    y = 1.5 + i * 1.3
    d.text(s, 8.95, y, 3.9, .4, [(t, 17, True, 'title')])
    d.text(s, 8.95, y + .42, 3.9, .4, [(ds, 13, False, 'muted')])
note(s, 8.95, 5.4, 3.92, .95, 'Chỉ:', 'Admin · Director · HOD', fill='soft', line=None)
N(s, 'Total KPI.', ['Màn tổng hợp KPI toàn công ty theo bốn góc nhìn tài chính, khách hàng, quy trình, học hỏi; chỉ Admin, Director, HOD mở được. Số trong ảnh là giả lập.'], 'Sang phần cuối: Training Hub.')

# ═════════════════════════ PHẦN 5 · TRAINING HUB
s = divider('5', 'Training Hub', 'Học · kiểm tra · chứng chỉ', ['3 vai trò', 'Tự học', 'Làm bài, chứng chỉ', 'Trainer', 'Cửa chung cho mọi app'], 'th/home', '5.1')
N(s, 'Mở phần Training Hub.', ['Training Hub vừa là nơi đào tạo nội bộ, vừa là cửa đăng nhập chung của các app.'], 'Ba vai trò.')

s = S('5 · Training Hub', '3 vai trò', '5 · TRAINING HUB', 'C27', '5.1', 'Ai dùng, đăng nhập')
for i, (n, t, ds) in enumerate([('01', 'Học viên', 'Đọc tài liệu, làm bài, nhận chứng chỉ'), ('02', 'Trainer', 'Tạo buổi, tài liệu, thư mời, bài kiểm tra'), ('03', 'Admin', 'Người dùng, cấu hình, chế độ test')]):
    x = .46 + i * 4.17
    d.text(s, x, 1.45, 3.9, .8, [(n, 44, True, 'sec')])
    d.text(s, x, 2.3, 3.9, .45, [(t, 20, True, 'title')])
    d.hline(s, x, 2.85, x + 3.8)
    d.text(s, x, 3.0, 3.85, .8, [(ds, 14, False, 'ink')])
shot(s, 'th/login', .46, 4.05, 6.1, 2.3, (.0, .1, 1, .9), 'mid')
note(s, 6.77, 4.05, 6.1, 2.3, 'Đăng nhập riêng:', 'email + mã PIN. Lần đầu kích hoạt bằng mã 6 số qua email. Sai PIN 8 lần khoá 10 phút.', fill='soft', line=None)
N(s, 'Ba vai trò và đăng nhập.', ['Học viên là mọi nhân sự. Trainer tạo buổi và bài kiểm tra. Admin quản trị.', 'Training Hub đăng nhập bằng email và mã PIN, lần đầu kích hoạt qua email.'], 'Trang đầu tiên của học viên.')

s = S('5 · Training Hub', 'Trang của tôi', '5 · TRAINING HUB', 'C44', '5.2', 'Trang của tôi')
shot(s, 'th/home_learner', .46, 1.45, 7.6, 4.9, (0, 0, 1, .9))
steps(s, 8.35, 1.5, 4.52, [('Việc cần làm', 'bài kiểm tra đang mở, theo hạn'), ('Hành trình học tập', 'Tài liệu → Bài kiểm tra → Chứng chỉ'), ('Sắp diễn ra', 'buổi hôm nay có nút Vào họp'), ('Học tiếp', 'tài liệu đang đọc dở')], gap=1.15)
N(s, 'Trang của tôi.', ['Mở Training Hub là thấy việc cần làm, hành trình học tập của từng buổi, các buổi sắp tới và tài liệu đang đọc dở.'], 'Học viên tự học thế nào?')

s = S('5 · Training Hub', 'Tự học', '5 · TRAINING HUB', 'C19', '5.5', 'Tự học: đọc tài liệu')
chev(s, ['Học tài liệu', 'Làm bài kiểm tra', 'Nhận chứng chỉ'], y=1.45, h=.65, sub=['Đọc tới trang cuối là xong', 'Mở khi học đủ tài liệu bắt buộc', 'PDF gửi email, CC quản lý'], subh=.6)
shot(s, 'th/reader', .46, 2.95, 8.0, 3.4, (0, 0, 1, .85))
bullets(s, 8.75, 3.0, 4.12, ['Đọc ngay trong Hub, không cần quyền Drive', 'PowerPoint, Word tự chuyển PDF', 'Tiến độ đọc được lưu'], 13, 8)
N(s, 'Lộ trình tự học.', ['Ba bước: học tài liệu, làm bài, nhận chứng chỉ. Tài liệu đọc ngay trong Hub, đọc tới trang cuối là tính xong.'], 'Làm bài kiểm tra.')

s = S('5 · Training Hub', 'Làm bài, nhận chứng chỉ', '5 · TRAINING HUB', 'C28', '5.6', 'Làm bài kiểm tra')
for i, (sid, t, bx) in enumerate([('th/quiz_question', 'Làm bài · có đồng hồ', (.1, .05, 1, .7)), ('th/quiz_result', 'Kết quả · thang 100', (.15, .05, .85, .75)), ('th/cert_preview', 'Chứng chỉ PDF', None)]):
    x = .46 + i * 4.17
    shot(s, sid, x, 1.45, 4.0, 3.4, bx, 'mid')
    d.text(s, x, 5.0, 4.0, .4, [(t, 16, True, 'title')])
d.text(s, .46, 5.55, 12.41, .8, [('Hết giờ tự nộp · mất mạng vẫn giữ đáp án · mỗi câu bằng điểm · đạt thì chứng chỉ tự gửi email.', 13.5, False, 'ink')])
N(s, 'Làm bài và chứng chỉ.', ['Bài kiểm tra có đồng hồ, hết giờ tự nộp, lỡ mất mạng vẫn giữ đáp án.', 'Đạt điểm thì chứng chỉ PDF tự tạo và gửi email, có CC quản lý.'], 'Phía trainer làm gì?')

s = S('5 · Training Hub', 'Trainer: tạo buổi', '5 · TRAINING HUB', 'C20', '5.8', 'Trainer: tạo buổi, tài liệu, thư mời')
chev(s, ['Tạo buổi', 'Tải tài liệu', 'Gửi thư mời', 'Tự nhắc'], y=1.45, h=.65, sub=['Thư mục Drive tự tạo', 'Kéo thả cả thư mục, ≤ 1 GB/file', 'Xem trước email thật', '8:30 ngày đào tạo, thứ Hai'], subh=.6)
shot(s, 'th/create_session', .46, 2.95, 6.1, 3.4, 'auto')
shot(s, 'th/invite', 6.77, 2.95, 6.1, 3.4, 'auto')
N(s, 'Trainer tạo buổi.', ['Tạo buổi thì thư mục Drive tự tạo. Kéo thả tài liệu vào, gửi thư mời có xem trước. Hệ thống tự nhắc lúc 8 giờ 30 ngày đào tạo và mỗi thứ Hai.'], 'Soạn bài và xem kết quả.')

s = S('5 · Training Hub', 'Trainer: bài kiểm tra', '5 · TRAINING HUB', 'C10', '5.9', 'Trainer: soạn bài kiểm tra, xem kết quả')
shot(s, 'th/quiz_builder', .46, 1.45, 7.6, 4.9, 'auto')
for i, (v, t) in enumerate([('Dán nhanh', 'câu hỏi từ Word / Excel'), ('Quy định', 'điểm đạt, số phút, số lượt, hạn'), ('Phân tích', 'từng câu; câu khó tô đỏ'), ('Nhắc', 'người chưa đạt; xuất CSV')]):
    d.kpi(s, 8.35, 1.45 + i * 1.25, 4.52, 1.1, v, t, vsz=22, dark=(i == 0))
N(s, 'Soạn bài và xem kết quả.', ['Trainer dán câu hỏi từ Word, đặt điểm đạt, thời gian, số lượt. Kết quả có phân tích từng câu và nút nhắc người chưa đạt.'], 'Cuối cùng: vai trò cửa chung.')

s = S('5 · Training Hub', 'Cửa chung cho mọi app', '5 · TRAINING HUB', 'C34', '5.12', 'Cửa chung cho các app khác')
d.table(s, .46, 1.45, [3.4, 4.6, 1.1, 1.1, 1.1, 1.11], [['Dịch vụ', 'Làm gì', 'Report Hub', 'CRM', 'Product Data', 'KPI'],
        ['Đăng nhập', 'Email + mã 6 số, 30 ngày', '●', '●', '●', '●'], ['Phân quyền', 'Bật / tắt quyền từng người', '●', '●', '●', '●'],
        ['Giao việc', 'Danh sách, trao đổi, hoàn thành', '●', '', '', ''], ['Email cập nhật', 'Báo bản mới tới mọi người', '●', '●', '', '']],
        row_h=.72, head=12, body=13.5, styles={1: 'focus'}, aligns=['l', 'l', 'c', 'c', 'c', 'c'])
note(s, .46, 5.35, 12.41, .95, 'Một tài khoản, mọi app.', 'Nghỉ việc hay đổi vị trí: Admin sửa một dòng, quyền ở mọi app đổi theo. Buổi Completed tự vào KPI đào tạo.', fill='soft', line=None)
N(s, 'Dịch vụ dùng chung.', ['Backend Training Hub giữ đăng nhập, phân quyền, giao việc và email cập nhật cho các app khác.', 'Vì vậy Admin chỉ cần sửa một nơi.'], 'Mình tổng kết.')

# ═════════════════════════ TỔNG KẾT
s = S('Tổng kết', '6 điều cần nhớ', 'TỔNG KẾT', 'C45', '1.1', 'Sách hướng dẫn đầy đủ')
SUM = [('1 hệ thống', '4 app dùng chung Google Sheets của công ty'), ('1 lần đăng nhập', 'Email + mã 6 số, 30 ngày'), ('Nhập 1 lần', 'Báo cáo, lịch, KPI tự tạo'),
       ('KPI tự điền', '6 nguồn; người chỉ đặt Target'), ('4 màu', 'Xám · xanh dương · vàng · xanh lá'), ('2 mốc', 'Ngày 02 báo cáo · ngày 09 đầu quý')]
for i, (t, ds) in enumerate(SUM):
    x = .46 + (i % 3) * 4.17; y = 1.45 + (i // 3) * 2.5
    d.box(s, x, y, 4.0, 2.3, 'hi' if i == 3 else 'panel', line='hi_line' if i == 3 else None)
    d.badge(s, x + .25, y + .25, i + 1)
    d.text(s, x + .25, y + .8, 3.55, .45, [(t, 19, True, 'title')])
    d.text(s, x + .25, y + 1.3, 3.55, .9, [(ds, 14, False, 'ink')])
N(s, 'Chốt 6 ý.', ['Một hệ thống, một lần đăng nhập, nhập một lần, KPI tự điền, bốn màu KPI và hai mốc ngày 02, ngày 09.'], 'Tra cứu chi tiết ở đâu?')

s = S('Tổng kết', 'Tra cứu chi tiết: IFU', 'TỔNG KẾT', 'C41', '1.1', 'Sách hướng dẫn đầy đủ')
F = [('MMH_IFU_He_thong_du_lieu.pdf', 'Sách đầy đủ · 139 trang', '1'), ('MMH_IFU_01_CRM.pdf', 'CRM · 40 trang', '2'), ('MMH_IFU_02_Report_Hub.pdf', 'Report Hub · 40 trang', '3'),
     ('MMH_IFU_03_KPI.pdf', 'KPI · 18 trang', '4'), ('MMH_IFU_04_Training_Hub.pdf', 'Training Hub · 33 trang', '5')]
for i, (fn, ds, part) in enumerate(F):
    x = .46 + (i % 2) * 6.27; y = 1.45 + (i // 2) * 1.35
    d.box(s, x, y, 6.12, 1.15, 'panel')
    d.icon(s, 'book', x + .2, y + .3, .55)
    tb = d.text(s, x + .95, y + .17, 5.0, .4, [(fn, 15, True, 'title')])
    tb.text_frame.paragraphs[0].runs[0].hyperlink.address = IFU_FILE[part][1]
    d.text(s, x + .95, y + .6, 5.0, .4, [(ds, 13, False, 'muted')])
note(s, 6.73, 4.15, 6.14, 1.15, 'Bấm tên file', 'để mở. Sách dựng lại mỗi khi hệ thống thay đổi.', fill='soft', line=None)
d.text(s, .46, 5.55, 12.4, .6, [('Tài liệu nội bộ: có link tới dữ liệu thật của công ty, không chia sẻ ra ngoài.', 13, False, 'muted', True)])
N(s, 'Chỉ chỗ tra cứu.', ['IFU có năm file: sách đầy đủ và bốn phần riêng. Bấm tên file để mở.', 'Sách được dựng lại mỗi khi hệ thống thay đổi, nên luôn khớp app.'], 'Cảm ơn mọi người.')

footer(d.prs.slides[2], '1.1', 'Sách hướng dẫn sử dụng đầy đủ', y=4.6, w=9.6, credit=False)
d.notes(d.prs.slides[2], 'Kết thúc, mời hỏi đáp.', ['Cảm ơn mọi người. Cần làm theo từng bước thì mở link IFU ở cuối mỗi slide, hoặc bấm nút Hướng dẫn trong app.', 'Mình mời câu hỏi.'])

# ── kiểm tra chống lặp (deck > 30 slide: tối đa 3 lần / pattern)
from collections import Counter
bad = [f'{a[2]}/{b[2]} liền kề cùng {a[1]}' for a, b in zip(LOG, LOG[1:]) if a[1] == b[1]]
bad += [f'{k} dùng {n} lần' for k, n in Counter(p for _, p, _ in LOG).items() if n > 3 and k != 'C04']
print('pattern issues:', bad)
print('long titles:', [t for _, _, t in LOG if len(t) > 30])
print('missing notes:', [i for i, x in enumerate(d.prs.slides, 1) if i != 2 and not x.notes_slide.notes_text_frame.text.strip()])
out = d.save('/home/claude/outputs/20261013_MMH_He_thong_du_lieu_noi_bo_v2.pptx'); print(out, len(d.prs.slides))
