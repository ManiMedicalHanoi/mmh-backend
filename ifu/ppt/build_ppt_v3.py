"""PPT giới thiệu hệ thống dữ liệu nội bộ MMH · bản 3: khung Dữ liệu → Chức năng, KPI theo "Hướng dẫn hệ thống KPI FY68", chế độ TEST, ảnh máy tính + điện thoại."""
# Bộ dựng PPT bản 3 (ghép từ các phần trong phiên dựng 13/10/2026). Đường dẫn tuyệt đối (/home/claude/...) ở đầu file: chỉnh theo máy đang chạy.
import sys, os
sys.path.insert(0, '/home/claude')
from style_c import *
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE

SH = '/home/claude/ifu/shots/'
ICO = '/home/claude/img/icons/'
CACHE = '/home/claude/img/crop3'; os.makedirs(CACHE, exist_ok=True)
TPL = '/home/claude/build/MMH_PPT_TEMPLATE.pptx'

# ── Link IFU: đổi ở đây khi có link file trên Drive (mỗi phần 1 file) ──
IFU_FOLDER = 'https://drive.google.com/drive/folders/1g1lfTxLU51qTkxIl-bRqJwVyvN5r7j6q'
IFU_FILE = {'1': ('MMH_IFU_He_thong_du_lieu.pdf', IFU_FOLDER), '2': ('MMH_IFU_01_CRM.pdf', IFU_FOLDER),
            '3': ('MMH_IFU_02_Report_Hub.pdf', IFU_FOLDER), '4': ('MMH_IFU_03_KPI.pdf', IFU_FOLDER),
            '5': ('MMH_IFU_04_Training_Hub.pdf', IFU_FOLDER), '6': ('MMH_IFU_05_Che_do_TEST.pdf', IFU_FOLDER), '7': ('MMH_IFU_He_thong_du_lieu.pdf', IFU_FOLDER)}
import json as _json
_P = _json.load(open('/home/claude/build/ifu_pages.json'))
PAGE = dict(_P['all']); [PAGE.update(_P[k]) for k in ('2', '3', '4', '5', '6')]   # số trang đọc từ mục lục của PDF IFU vừa dựng

SEC = ['TỔNG QUAN', '1 · NỀN TẢNG', '2 · CRM', '3 · REPORT HUB', '4 · KPI', '5 · TRAINING HUB', '6 · CHẾ ĐỘ TEST', 'TỔNG KẾT']
d = Deck('mmh', template=TPL, palette='classic', sections=SEC)
ensure_icons(d, names=('users', 'target', 'check', 'chart', 'book', 'shield', 'clock', 'file', 'alert', 'bulb', 'layers',
                       'arrow', 'box', 'globe', 'award', 'cal', 'eye', 'building', 'search', 'wallet', 'truck'), roles=('sec', 'pri', 'on_pri'))
LOG = []

# ───────────────────────── helpers
def modal_box(im):
    """Tìm khung hộp thoại trắng trên nền mờ (ảnh có lớp phủ); không thấy thì trả toàn ảnh."""
    import numpy as np
    g = np.asarray(im.convert('L').resize((im.width // 4, im.height // 4)), dtype=float)
    m = g >= 249
    rows = np.where(m.mean(1) > .18)[0]; cols = np.where(m.mean(0) > .18)[0]
    if len(rows) < 5 or len(cols) < 5: return None
    H, W = g.shape; y0, y1, x0, x1 = rows[0], rows[-1], cols[0], cols[-1]
    if (x1 - x0) * (y1 - y0) > .92 * W * H: return None
    pad = .012
    return (max(0, x0 / W - pad), max(0, y0 / H - pad), min(1, x1 / W + pad), min(1, y1 / H + pad))

def crop(path, w, h, box=None, anchor='top', maxw=2000, fit=True):
    """box = vùng cắt (tỉ lệ) hoặc 'auto' (khung hộp thoại). fit=True: giữ đủ vùng, thêm lề trắng cho đúng tỉ lệ khung."""
    im = Image.open(path).convert('RGB')
    if box == 'auto': box = modal_box(im)
    if box:
        W, H = im.size; im = im.crop((int(box[0] * W), int(box[1] * H), int(box[2] * W), int(box[3] * H)))
    W, H = im.size; r = w / h
    if fit:
        if abs(W / H - r) > .01:
            if W / H > r: nw, nh = W, int(W / r)
            else: nw, nh = int(H * r), H
            bg = Image.new('RGB', (nw, nh), (255, 255, 255))
            bg.paste(im, ((nw - W) // 2, 0 if anchor == 'top' and W / H > r else (nh - H) // 2)); im = bg
    else:
        if W / H > r:
            nw = int(H * r); x0 = (W - nw) // 2 if anchor != 'left' else 0; im = im.crop((x0, 0, x0 + nw, H))
        else:
            nh = int(W / r); y0 = 0 if anchor in ('top', 'left') else (H - nh) // 2; im = im.crop((0, y0, W, y0 + nh))
    if im.width > maxw: im = im.resize((maxw, int(maxw * im.height / im.width)), Image.LANCZOS)
    out = f'{CACHE}/{abs(hash((path, w, h, box, anchor, fit)))}.jpg'; im.save(out, quality=86); return out

AUTO = {'crm/team_switch', 'crm/trip_approve', 'crm/test_mode', 'crm/test_pick', 'rh/test_pick', 'rh/assign_mail', 'rh/assign_popup', 'rh/assign_track', 'th/quiz_question', 'th/quiz_result', 'crm/newacc_review'}
def shot(sl, sid, x, y, w, h, box=None, anchor='top'):
    p = SH + sid + '.jpg'
    if sid in AUTO: box = 'auto'
    d.box(sl, x - .03, y - .03, w + .06, h + .06, 'line', kind=MSO_SHAPE.RECTANGLE)
    sl.shapes.add_picture(crop(p, w, h, box, anchor), I(x), I(y), I(w), I(h))

def phone(sl, sid, x, y, h, box=None):
    im = Image.open(SH + sid + '.jpg'); r = im.width / im.height
    if box: r = r * (box[2] - box[0]) / (box[3] - box[1])
    w = h * r
    d.box(sl, x - .06, y - .06, w + .12, h + .12, 'ink', adj=.08)
    sl.shapes.add_picture(crop(SH + sid + '.jpg', w, h, box), I(x), I(y), I(w), I(h)); return w

def logo(sl, name, x, y, s=.4):
    sl.shapes.add_picture(ICO + name + '.png', I(x), I(y), I(s), I(s))

def footer(sl, sec, title, y=6.64, w=9.6, credit=True):
    """Dòng cuối slide: link tới đúng mục trong IFU."""
    part = sec.split('.')[0]; fname, url = IFU_FILE[part]; pg = PAGE.get(sec)
    tb = sl.shapes.add_textbox(I(.46), I(y), I(w), I(.27)); tf = tb.text_frame
    for a in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, a, 0)
    p = tf.paragraphs[0]
    r0 = p.add_run(); r0.text = 'Chi tiết  '; r0.font.size = Pt(10); r0.font.name = d.font; r0.font.color.rgb = d.c('muted')
    r = p.add_run(); r.text = f'IFU {sec} · {title}' + (f' · {fname} tr. {pg}' if pg else '') + '  ↗'
    r.font.size = Pt(10); r.font.bold = True; r.font.name = d.font; r.font.color.rgb = d.c('sec'); r.hyperlink.address = url
    if credit: d.text(sl, 10.2, 6.64, 2.67, .27, [('Ảnh thật · dữ liệu giả lập', 9, False, 'muted', True)], 'r')

def S(kicker, title, section, pattern, sec=None, sectitle=None):
    sl = d.slide(kicker, title, section, pattern)
    LOG.append((d.page, pattern, title))
    if sec: footer(sl, sec, sectitle or title)
    return sl

def steps(sl, x, y, w, items, gap=.62, sz=13.5, start=1):
    for i, t in enumerate(items):
        yy = y + i * gap
        d.badge(sl, x, yy, i + start, d=.34, sz=12)
        if isinstance(t, tuple):
            d.text(sl, x + .5, yy - .04, w - .5, .3, [(t[0], 14, True, 'title')])
            d.text(sl, x + .5, yy + .27, w - .5, gap - .3, [(t[1], 12.5, False, 'ink')])
        else:
            d.text(sl, x + .5, yy - .03, w - .5, gap, [(t, sz, False, 'ink')])

def bullets(sl, x, y, w, items, sz=13.5, gap=5):
    d.text(sl, x, y, w, .35 * len(items), [[('•  ', sz, True, 'sec'), (t, sz, False, 'ink')] for t in items], spc=gap)

def chev(sl, items, y=1.5, h=.7, x0=.46, W=12.41, last_dark=True, sub=None, subh=.9):
    n = len(items); w = (W + .1 * (n - 1)) / n
    for i, t in enumerate(items):
        x = x0 + i * (w - .1)
        fill = 'pri' if (last_dark and i == n - 1) else ('sec' if i == 0 else 'ter')
        d.box(sl, x, y, w, h, fill, kind=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
        d.text(sl, x + .32, y, w - .55, h, [(t, 13, True, 'on_pri' if fill != 'ter' else 'title')], 'c', 'm')
        if sub: d.text(sl, x + .1, y + h + .15, w - .35, subh, [(sub[i], 12, False, 'ink')])

def note(sl, x, y, w, h, head, body, fill='hi', line='hi_line'):
    d.box(sl, x, y, w, h, fill, line=line)
    d.text(sl, x + .22, y + .1, w - .44, h - .2, [[(head + '  ', 13.5, True, 'title'), (body, 13, False, 'ink')]], anchor='m')

def divider(no, name, sub, points, img, sec, pattern='C04'):
    sl = d.slide(f'Phần {no}', f'Phần {no} · {name}', SEC[int(no)], pattern)
    LOG.append((d.page, pattern, name))
    d.text(sl, .46, 1.5, 5.4, .45, [(f'PHẦN {no}', 16, True, 'sec')])
    d.text(sl, .46, 1.95, 5.4, .85, [(name, 34, True, 'title')])
    d.text(sl, .46, 2.8, 5.4, .55, [(sub, 15, False, 'muted', True)])
    steps(sl, .46, 3.65, 5.4, points, gap=.56, sz=14)
    d.box(sl, 6.15, 1.38, 6.72, 5.12, 'pri')
    if img.endswith('_m'):
        pass
    shot(sl, img, 6.42, 1.65, 6.18, 4.58)
    footer(sl, sec, f'Chương {no} · {name}')
    return sl

def N(sl, purpose, talk, transition=None, reasoning=None):
    d.notes(sl, purpose, talk, reasoning=reasoning, transition=transition)

def pair(sl, dsid, msid, x, y, W, h, box=None, anchor='top', mbox=None):
    """Ảnh máy tính (trái) + điện thoại (phải) cùng chiều cao."""
    pw = h * 390 / 844
    dw = W - pw - .3
    shot(sl, dsid, x, y, dw, h, box, anchor)
    phone(sl, msid, x + dw + .3 + .06, y, h, mbox)
    d.text(sl, x, y + h + .06, dw, .25, [('Máy tính', 10, True, 'muted')])
    d.text(sl, x + dw + .3, y + h + .06, pw + .1, .25, [('Điện thoại', 10, True, 'muted')], 'c')

def link_run(p, text, url, sz=12, bold=True, col='sec'):
    r = p.add_run(); r.text = text; r.font.size = Pt(sz); r.font.bold = bold; r.font.name = d.font; r.font.color.rgb = d.c(col)
    if url: r.hyperlink.address = url
    return r


# ═════════════════════════ dữ liệu KPI (dùng chung với chương KPI của IFU)
_src = open('/home/claude/ifu/tools/kpi_chapter.py', encoding='utf-8').read()
_ns = {}; exec(_src[:_src.index('\nout = []')], _ns)
KF, KG, KPEOPLE, kstat = _ns['F'], _ns['GROUPS'], _ns['PEOPLE'], _ns['pstat']
STC = {'A': '2E8B4E', 'S': '9AA6B2', 'W': 'E0A800', 'M': 'C9473B', 'N': 'D5DCE4'}
STN = {'A': 'Đã tự động kết nối', 'S': 'Tự cộng từ dòng con', 'W': 'Đã kết nối · chưa nhận số', 'M': 'Chưa kết nối · nhập tay', 'N': 'Chưa có file nguồn'}

def divider(no, name, sub, points, img, sec, mimg=None, pattern='C04'):
    sl = d.slide(f'Phần {no}', f'Phần {no} · {name}', SEC[int(no)], pattern)
    LOG.append((d.page, pattern, name))
    d.text(sl, .46, 1.5, 5.4, .45, [(f'PHẦN {no}', 16, True, 'sec')])
    d.text(sl, .46, 1.95, 5.4, .85, [(name, 32, True, 'title')])
    d.text(sl, .46, 2.85, 5.4, .55, [(sub, 15, False, 'muted', True)])
    steps(sl, .46, 3.65, 5.4, points, gap=.5, sz=14)
    d.box(sl, 6.15, 1.38, 6.72, 5.12, 'pri')
    if mimg: pair(sl, img, mimg, 6.4, 1.62, 6.25, 4.35)
    else: shot(sl, img, 6.42, 1.65, 6.18, 4.58)
    footer(sl, sec, f'Chương {no} · {name}')
    return sl

def dtable(sl, x, y, cols, rows, row_h=.42, head=11.5, body=11.5):
    d.table(sl, x, y, cols, rows, row_h=row_h, head=head, body=body, aligns=['l'] * len(cols))

# ═════════════════════════ BÌA & MỤC LỤC
s1 = d.cover_mmh('GIỚI THIỆU HỆ THỐNG · DÀNH CHO TOÀN CÔNG TY', 'HỆ THỐNG DỮ LIỆU', 'NỘI BỘ MMH', 'Product Team · 13/10/2026')
LOG.append((1, 'C01', 'Bìa'))
footer(s1, '1.1', 'Sách hướng dẫn sử dụng đầy đủ', y=6.45, w=6.3, credit=False)
N(s1, 'Mở đầu, nêu mục tiêu và cách bài trình bày đi theo cuốn hướng dẫn IFU.',
  ['Chào mọi người. Hôm nay mình giới thiệu hệ thống dữ liệu nội bộ của MMH: CRM, Report Hub, KPI, Training Hub và chế độ TEST để kiểm thử.',
   'Mỗi app trình bày theo cùng một khung: dữ liệu nằm ở đâu, được tạo, sửa, lưu thế nào, rồi các chức năng chính.',
   'Cuối mỗi slide có link tới đúng mục trong cuốn hướng dẫn IFU. Ảnh là ảnh chụp thật từ app, có cả bản máy tính và điện thoại, dữ liệu giả lập.'], 'Mình xem mục lục.')

s = S('Tổng quan', 'Nội dung', 'TỔNG QUAN', 'C03', '1.1', 'Sách hướng dẫn đầy đủ')
rows = [('1', 'Nền tảng chung', 'Đăng nhập, quyền, lưu dữ liệu, nhắc hạn'), ('2', 'MMH CRM', 'Dữ liệu · địa bàn · đơn hàng · lịch · báo cáo · thống kê'),
        ('3', 'MMH Report Hub', 'Dữ liệu · công việc · lịch chung · báo cáo · thống kê'), ('4', 'Hệ thống KPI FY68', 'Nguồn · quy tắc · xem ở đâu · theo phòng ban'),
        ('5', 'Training Hub', 'Học, kiểm tra, chứng chỉ'), ('6', 'Chế độ TEST', 'Admin thử mọi chức năng, xem luồng dữ liệu'), ('', 'Tổng kết', 'Điều cần nhớ, tra cứu IFU')]
for i, (n, t, ds) in enumerate(rows):
    y = 1.42 + i * .72
    d.text(s, .6, y, .8, .5, [(n, 22, True, 'sec')], anchor='m')
    d.text(s, 1.5, y, 4.2, .5, [(t, 19, True, 'title')], anchor='m')
    d.text(s, 5.8, y, 7.0, .5, [(ds, 14, False, 'muted')], anchor='m')
    d.hline(s, .6, y + .62, 12.8)
N(s, 'Cho thấy mạch bài.', ['Sáu phần theo đúng sáu chương IFU: nền tảng, CRM, Report Hub, KPI, Training Hub và chế độ TEST, rồi tổng kết.'], 'Bắt đầu với phần nền tảng chung.')
# ═════════════════════════ PHẦN 1 · NỀN TẢNG
s = divider('1', 'Nền tảng chung', 'Phần chung của mọi app', ['4 app, 1 hệ thống', 'Dữ liệu đi đâu', 'Đăng nhập bằng email', 'Ai thấy gì', 'Lưu không phải chờ'], 'rh/main', '1.1', 'rh/m_task_list')
N(s, 'Mở phần 1.', ['Phần một nói về những gì giống nhau ở mọi app: đăng nhập, quyền xem, cách lưu dữ liệu, nhắc hạn.', 'Hiểu phần này một lần thì các app sau dùng rất nhanh.'], 'Trước hết, hệ thống gồm những gì.')

s = S('1 · Nền tảng', '4 app, 1 hệ thống', '1 · NỀN TẢNG', 'C27', '1.1', 'Hệ thống gồm những gì')
apps = [('MMH CRM', 'Đội Sales', 'crm/cal_week'), ('Report Hub', 'Các phòng ban', 'rh/main'), ('KPI', 'Mọi người · quản lý', 'rh/kpi_tab'), ('Training Hub', 'Học viên · trainer', 'th/home')]
for i, (t, u, img) in enumerate(apps):
    x = .46 + i * 3.15
    shot(s, img, x, 1.45, 2.95, 1.85)
    d.text(s, x, 3.5, 2.95, .4, [(t, 18, True, 'title')])
    d.text(s, x, 3.92, 2.95, .3, [(u, 13, False, 'muted')])
d.box(s, .46, 4.6, 12.41, 1.75, 'soft')
for i, (lg, t, ds) in enumerate([('googlesheets', 'Google Sheets', 'dữ liệu gốc'), ('googleappsscript', 'Apps Script', 'cầu nối'), ('googledrive', 'Drive', 'tài liệu, ảnh'), ('gmail', 'Gmail', 'email tự động'), ('github', 'GitHub Pages', 'nơi đặt app')]):
    x = .75 + i * 2.45
    logo(s, lg, x, 4.95, .5)
    d.text(s, x + .62, 4.9, 1.8, .3, [(t, 13.5, True, 'title')])
    d.text(s, x + .62, 5.2, 1.8, .3, [(ds, 12, False, 'muted')])
d.text(s, .75, 5.75, 11.9, .4, [('Mọi app mở bằng trình duyệt, không cài đặt; dữ liệu nằm trong Google Sheets của công ty.', 13.5, False, 'ink')])
N(s, 'Giới thiệu 4 app và nền tảng Google bên dưới.',
  ['Hàng trên là bốn app: CRM cho Sales, Report Hub cho các phòng ban, KPI cho mọi người, Training Hub cho đào tạo.',
   'Hàng dưới là nền tảng: dữ liệu gốc nằm trong Google Sheets của công ty, Apps Script là cầu nối, Drive giữ tài liệu và ảnh, Gmail gửi email tự động, app đặt trên GitHub Pages.',
   'Mọi app mở bằng trình duyệt, không cần cài.'], 'Vậy dữ liệu đi qua đâu?')

s = S('1 · Nền tảng', 'Dữ liệu đi đâu', '1 · NỀN TẢNG', 'C19', '1.2', 'Dữ liệu nằm ở đâu, đi qua đâu')
chev(s, ['Người dùng', 'App web', 'Apps Script', 'Google Sheets'], y=1.6, h=.8,
     sub=['Thao tác trên máy tính hoặc điện thoại', 'Hiện kết quả ngay, gửi lệnh ở phía sau', 'Kiểm tra phiên đăng nhập và quyền, rồi ghi', 'Dữ liệu gốc của công ty; Drive, Gmail'], subh=1.0)
note(s, .46, 4.1, 12.41, 1.0, 'Nhập một lần.', 'Sales báo kết quả trên CRM ⇒ số lượt tự sang KPI. Buổi đào tạo xong ⇒ KPI đào tạo tự đếm. Công tác được duyệt ⇒ tự hiện trên lịch chung.')
d.text(s, .46, 5.4, 12.41, .8, [('Mã backend quản lý tập trung trong kho mmh-backend: sửa xong tự cập nhật, đường link app không đổi, lỗi thì tự quay về bản trước.', 13, False, 'muted')])
N(s, 'Giải thích đường đi của dữ liệu và nguyên tắc nhập một lần.',
  ['Đọc từ trái sang phải: người dùng thao tác trên app; app hiện kết quả ngay và gửi lệnh ở phía sau; Apps Script kiểm tra đăng nhập, quyền rồi ghi vào Google Sheets.',
   'Ô vàng là nguyên tắc quan trọng nhất: nhập một lần. Dữ liệu nhập ở app này tự xuất hiện ở app khác.'], 'Tiếp theo là cách đăng nhập.')

s = S('1 · Nền tảng', 'Đăng nhập bằng email', '1 · NỀN TẢNG', 'C30', '1.3', 'Đăng nhập bằng email công ty')
shot(s, 'rh/login_email', .46, 1.45, 4.6, 3.0, 'auto')
phone(s, 'crm/m_login', 5.45, 1.45, 3.0)
steps(s, 8.0, 1.5, 4.87, [('Nhập email công ty', '@manimedicalhanoi.com hoặc @mani.inc'), ('Lấy mã 6 số', 'trong email vừa nhận'), ('Nhập mã', 'app tự nhận tên, nhóm, quyền')], gap=.95)
d.table(s, .46, 4.75, [3.6, 3.6], [['Quy định', 'Giá trị'], ['Mã hết hạn', '10 phút'], ['Nhập sai', '5 lần'], ['Dùng được', '30 ngày / máy']], row_h=.4, head=12, body=12.5, aligns=['l', 'l'])
note(s, 8.0, 4.75, 4.87, 1.6, 'Một lần, mọi app.', 'Đăng nhập Report Hub rồi thì mở CRM, Product Data không phải đăng nhập lại.', fill='soft', line=None)
N(s, 'Hướng dẫn đăng nhập.',
  ['Không cần mật khẩu. Nhập email công ty, hệ thống gửi mã 6 số, nhập mã là vào.', 'Mã hết hạn sau 10 phút, sai 5 lần phải xin mã mới, một lần đăng nhập dùng 30 ngày trên máy đó.',
   'Đăng nhập một app thì các app khác trên cùng trình duyệt cũng nhận luôn.', 'Riêng Training Hub có đăng nhập bằng mã PIN, mình nói ở phần năm.'], 'Đăng nhập xong, mỗi người thấy gì?')

s = S('1 · Nền tảng', 'Ai thấy gì', '1 · NỀN TẢNG', 'C33', '1.4', 'Quyền xem và quyền sửa')
d.table(s, .46, 1.45, [2.6, 4.9, 4.91], [['Vai trò', 'Thấy', 'Làm thêm'], ['PIC / Sales', 'Việc, lịch, KPI của mình', 'Sửa việc của mình'],
        ['Team Leader', 'Cả nhóm', 'Duyệt một số mục của nhóm'], ['HOD', 'Cả phòng ban', 'Theo phân quyền'],
        ['Director', 'Toàn công ty, Total KPI', 'Duyệt công tác'], ['Admin', 'Toàn bộ', 'Phân quyền, chế độ TEST, email cập nhật']],
        row_h=.62, head=13, body=13.5, styles={4: 'focus'}, aligns=['l', 'l', 'l'])
note(s, .46, 5.4, 12.41, .85, 'Mặc định:', 'chuyến công tác của người khác bị ẩn; buổi đào tạo chỉ hiện với trainer và người được mời.', fill='soft', line=None)
N(s, 'Giải thích quyền xem theo vai trò.',
  ['PIC chỉ thấy việc của mình; Team Leader thấy cả nhóm; HOD thấy phòng ban; Director thấy toàn công ty và duyệt công tác; Admin phân quyền.',
   'Admin có thể bật thêm hoặc tắt bớt từng quyền cho từng người.'], 'Tiếp theo là cách app lưu dữ liệu.')

s = S('1 · Nền tảng', 'Lưu không phải chờ', '1 · NỀN TẢNG', 'C18', '1.5', 'Lưu không phải chờ')
steps(s, .46, 1.5, 6.3, [('Bấm Lưu', 'khung đóng ngay, thay đổi hiện trên màn hình'), ('App gửi ở phía sau', 'lần lượt từng lệnh, ảnh gửi sau chữ'),
                          ('Mạng chập chờn', 'tự thử lại, không tạo bản ghi trùng'), ('Lỗi thật', 'chip đỏ trên thanh trên cùng: Thử lại / Bỏ')], gap=1.0)
shot(s, 'rh/outbox', 7.1, 1.5, 3.7, 3.83, 'auto')
phone(s, 'crm/m_result', 11.1, 1.5, 3.83)
note(s, .46, 5.6, 12.41, .7, 'Lưu ý:', 'còn chip đỏ thì không xoá dữ liệu trình duyệt; đóng tab vẫn giữ, lần mở sau app gửi tiếp.', fill='hi')
N(s, 'Giải thích cơ chế lưu không phải chờ.',
  ['Bấm Lưu là xong, khung đóng ngay. App tự gửi lên Sheet ở phía sau, mạng yếu thì tự thử lại.', 'Chỉ khi có lỗi thật, ví dụ dòng đã bị xoá trên Sheet, app mới hiện chip đỏ để mình chọn Thử lại hoặc Bỏ.'],
  'Cuối phần một: các nhắc hạn tự động.')

s = S('1 · Nền tảng', 'App tự nhắc hạn', '1 · NỀN TẢNG', 'C40', '1.6', 'Nhắc hạn, thông báo và hướng dẫn')
L = [('Thông báo cập nhật', '7 ngày sau mỗi bản mới'), ('Hạn chứng từ kế toán', 'ngày 14, 15, 28, 29 và ngày hạn'), ('Việc mới được giao', 'lần mở app đầu tiên'),
     ('Thư nhắc việc', 'việc sắp đến hạn, báo cáo chưa nộp'), ('Báo cáo KPI tháng', 'ngày 02, trước 16:00'), ('Tự đánh giá quý', 'ngày 09 tháng 9 · 12 · 3 · 6, trước 17:00')]
for i, (t, ds) in enumerate(L):
    y = 1.5 + i * .78
    d.badge(s, .46, y + .04, i + 1, d=.34, sz=12)
    d.text(s, 1.0, y, 3.3, .42, [(t, 14.5, True, 'title')], anchor='m')
    d.text(s, 4.3, y, 3.4, .42, [(ds, 13, False, 'ink')], anchor='m')
    d.hline(s, .46, y + .66, 7.7)
shot(s, 'rh/kpi_reminder_card', 8.1, 1.5, 2.9, 2.6, None, 'auto')
phone(s, 'crm/m_today', 11.25, 1.5, 3.4)
d.text(s, 8.1, 5.05, 4.77, 1.3, [('Popup hiện lần lượt theo thứ tự này, không chồng lên nhau. Nút Hướng dẫn (?) chỉ từng bước ngay trên màn hình.', 13, False, 'ink')])
N(s, 'Liệt kê các nhắc hạn tự động.',
  ['Sáu loại popup, hiện lần lượt theo thứ tự trên bảng.', 'Hai mốc quan trọng nhất cho mọi người: ngày 02 gửi báo cáo KPI tháng và ngày 09 đầu quý tự đánh giá.'], 'Sang phần hai: CRM.')


# ═════════════════════════ PHẦN 2 · CRM
s = divider('2', 'MMH CRM', 'Sổ tay số của đội Sales', ['Dữ liệu: nguồn, cách lưu, lưu ở đâu', 'Cập nhật địa bàn', 'Đơn hàng · MMH Calendar', 'Báo cáo · biểu đồ', 'Quản lý'], 'crm/cal_week', '2.1', 'crm/m_cal')
N(s, 'Mở phần CRM.', ['CRM là công cụ hằng ngày của đội Sales. Mình đi theo khung: dữ liệu ở đâu, rồi năm nhóm chức năng chính.'], 'CRM dùng cho ai?')

s = S('2 · CRM', 'Ai dùng CRM', '2 · CRM', 'C36', '2.1', 'Ai dùng CRM, mở ở đâu')
T = [('pri', 'Dental', 'Tiếng Việt'), ('sec', 'Surgical', 'Tiếng Việt'), ('ter', 'Eyeless', 'Tiếng Việt'), ('acc', 'Thailand Surgical', 'Tiếng Anh, không chụp ảnh')]
for i, (c, t, ds) in enumerate(T):
    x = .46 + i * 3.15
    d.box(s, x, 1.45, 2.95, 1.35, c)
    col = 'on_pri' if c in ('pri', 'sec') else 'title'
    d.text(s, x + .2, 1.55, 2.6, .5, [(t, 18, True, col)], anchor='m')
    d.text(s, x + .2, 2.08, 2.6, .55, [(ds, 12.5, False, col)])
d.text(s, .46, 2.95, 12.4, .4, [('Mỗi nhóm một file Google Sheet riêng: khách hàng, đơn hàng, báo cáo tách biệt.', 13.5, False, 'muted')])
d.table(s, .46, 3.6, [3.0, 9.41], [['Vai trò', 'Thấy gì'], ['Sales (PIC)', 'Lịch, khách hàng, đơn hàng, KPI của mình'], ['Team Leader', 'Cả nhóm: chip Cả đội và từng PIC'],
        ['Quản lý', 'Đổi nhóm đang xem, duyệt mở mới, duyệt công tác (Director)']], row_h=.55, head=13, body=13.5, aligns=['l', 'l'])
N(s, 'Bốn nhóm và quyền xem.', ['Bốn nhóm, mỗi nhóm một file dữ liệu riêng. Nhóm Thái dùng tiếng Anh, không chụp ảnh đi địa bàn.', 'Sales thấy của mình, Team Leader thấy cả nhóm, quản lý đổi nhóm đang xem.'], 'Dữ liệu CRM nằm ở đâu?')

s = S('2 · CRM', 'Dữ liệu: lưu thế nào', '2 · CRM', 'C19', '2.2', 'Dữ liệu của CRM: nguồn, cách lưu, lưu ở đâu')
chev(s, ['Tạo / sửa trên app', 'Hàng đợi trên máy', 'Backend của nhóm', 'Đi tiếp'], y=1.45, h=.7,
     sub=['Máy tính hoặc điện thoại; màn hình cập nhật ngay', 'Chữ đi trước, ảnh sau; mất mạng vẫn giữ, tự gửi lại', 'Apps Script kiểm tra, chống ghi trùng, ghi đúng tab', 'Tab 9 tính KPI ⇒ file KPI · lịch · báo cáo · email'], subh=.85)
note(s, .46, 3.35, 12.41, .75, 'App không giữ dữ liệu.', 'Mọi thứ nằm trong file Google Sheet của nhóm; KPI, lịch, báo cáo đọc lại từ chính các tab đó.', fill='soft', line=None)
dtable(s, .46, 4.3, [3.2, 4.2, 5.01], [['Dữ liệu', 'Lưu ở (file CRM nhóm ▸ tab)', 'Đi tiếp tới'], ['Đi địa bàn + ảnh', '6. WEEKLY REPORT · ảnh ở Drive', 'Lịch, báo cáo tuần, KPI viếng thăm C3'],
       ['Khách hàng · người liên hệ', '1. CUSTOMER CODE · 2. CBC', 'Hồ sơ 360°'], ['Đơn hàng', '3. SALE ORDER', 'KPI bán hàng hỗ trợ C5 – C7'], ['Mở mới & SKU mới', '16. NEW ACCOUNT & SKU', 'KPI mở mới C1 sau khi quản lý xác nhận']], row_h=.38)
N(s, 'Đường đi của dữ liệu CRM.', ['Bấm lưu là màn hình cập nhật ngay, app gửi ở phía sau; mất mạng thì giữ trên máy rồi tự gửi.', 'Backend của nhóm ghi vào đúng tab. Từ đó tab KPI tự tính và đẩy sang file KPI, lịch và báo cáo cũng đọc từ đây.'], 'Bảng đầy đủ các loại dữ liệu.')

s = S('2 · CRM', 'Dữ liệu: lưu ở đâu', '2 · CRM', 'C33', '2.2', 'Dữ liệu của CRM: nguồn, cách lưu, lưu ở đâu')
dtable(s, .46, 1.42, [2.75, 3.1, 3.55, 3.01], [['Dữ liệu', 'Tạo / sửa ở đâu', 'Lưu ở', 'Đi tiếp tới'],
       ['Đi địa bàn, kết quả', 'Lịch ▸ + · Ghi nhanh · Báo kết quả', '6. WEEKLY REPORT', 'Lịch, báo cáo tuần, họp tuần, C3'],
       ['Ảnh + GPS', 'Thẻ đi địa bàn ▸ Chụp ảnh', 'Drive (PIC_ngày_account_STT.jpg) + cột Photo', 'Xem ảnh, điều kiện KPI'],
       ['Khách hàng · CBC · ADP', 'Khách hàng · Key stakeholder · ADP', '1. CUSTOMER · 2. CBC · 8. ADP', 'Hồ sơ 360°, Dashboard'],
       ['Đơn hàng (Dental)', 'Đơn hàng ▸ + Đơn hàng mới', '3. SALE ORDER', 'C5 – C7, Dashboard'],
       ['Thuyết trình', 'Thuyết trình ▸ + Thêm buổi', '7. PRODUCT PRESENTATION', 'KPI sự kiện C10'],
       ['Công việc trong tháng', 'Công việc trong tháng', '5. MONTHLY REPORT', 'Báo cáo tháng'],
       ['Mở mới & SKU mới', 'Mở mới ▸ + Thêm case', '16. NEW ACCOUNT & SKU + Drive', 'Quản lý duyệt ⇒ C1'],
       ['Công việc khác', 'Lịch ▸ Công việc khác', 'ALL TASK', 'Báo cáo tuần / tháng'],
       ['Công tác', 'Đề xuất · Duyệt · báo cáo', 'File Business Trip + Business Trip Log', 'MMH Calendar, email'],
       ['KPI (chỉ xem)', 'KPI & Mục tiêu', 'MMH_KPI_FY68', 'Báo cáo KPI tháng']], row_h=.43, head=11.5, body=11)
N(s, 'Bảng tra dữ liệu.', ['Mỗi loại dữ liệu có một chỗ tạo trên app và một tab trong file của nhóm. Ảnh và chứng từ nằm trong Drive, tên file có quy tắc.', 'KPI chỉ xem trên CRM, số gốc ở file KPI.'], 'Năm nhóm chức năng chính.')

s = S('2 · CRM', '5 nhóm chức năng', '2 · CRM', 'C27', '2.4', 'Chức năng chính')
FN = [('pri', 'Cập nhật địa bàn', 'Đi địa bàn, chụp ảnh · khách hàng, người liên hệ, ADP, mở mới', 'crm/fieldvisit_result'),
      ('sec', 'Thông tin đơn hàng', 'Đơn hàng Dental sale theo tháng, PIC, nhóm sản phẩm', 'crm/orders'),
      ('ter', 'MMH Calendar', 'Kế hoạch cá nhân / đội: công tác, sự kiện, bài đăng, đào tạo', 'crm/mmhcal'),
      ('acc', 'Hệ thống báo cáo', 'Báo cáo tuần, tháng, họp tuần, báo cáo KPI', 'crm/send_report'),
      ('pri', 'Biểu đồ thống kê', 'Tiến độ công việc, KPI, Dashboard', 'crm/dashboard')]
for i, (c, t, ds, img) in enumerate(FN):
    x = .46 + i * 2.5
    shot(s, img, x, 1.45, 2.35, 1.6, 'auto' if img == 'crm/fieldvisit_result' else None)
    d.box(s, x, 3.2, 2.35, .08, c)
    d.text(s, x, 3.38, 2.35, .45, [(t, 15, True, 'title')])
    d.text(s, x, 3.85, 2.35, 1.5, [(ds, 12.5, False, 'ink')])
note(s, .46, 5.55, 12.41, .75, 'Trên điện thoại:', 'thanh dưới Lịch · Khách hàng · Hôm nay · Mở mới · Thêm; nút Ghi nhanh nổi góc phải.', fill='soft', line=None)
N(s, 'Năm nhóm chức năng.', ['Cập nhật địa bàn, đơn hàng, MMH Calendar, báo cáo và biểu đồ thống kê. Mỗi nhóm có một slide với ảnh máy tính và điện thoại.'], 'Bắt đầu với đi địa bàn.')

s = S('2 · CRM', 'Đi địa bàn', '2 · CRM', 'C20', '2.5', 'Cập nhật địa bàn: hoạt động đi địa bàn')
chev(s, ['Tạo lượt', 'Gặp khách', 'Báo kết quả', 'Chụp ảnh'], y=1.42, h=.55, sub=None)
pair(s, 'crm/fieldvisit_result', 'crm/m_result', .46, 2.2, 8.4, 3.9, 'auto')
steps(s, 9.2, 2.25, 3.67, [('Lưu ở', '6. WEEKLY REPORT'), ('Bắt buộc', 'người gặp + kết quả'), ('Ảnh', 'Drive, tên PIC_ngày_…'), ('Đi tiếp', 'lịch · báo cáo · KPI C3')], gap=.95)
N(s, 'Luồng đi địa bàn.', ['Tạo lượt, gặp khách, báo kết quả, chụp ảnh. Trên điện thoại thẻ đi địa bàn có sẵn nút Chụp ảnh và Báo kết quả.', 'Dữ liệu vào tab 6, ảnh vào Drive, và tự sang lịch, báo cáo tuần, KPI viếng thăm.'], 'Thông tin khách hàng.')

s = S('2 · CRM', 'Khách hàng 360°', '2 · CRM', 'C07', '2.6', 'Cập nhật địa bàn: thông tin khách hàng')
pair(s, 'crm/cust_profile', 'crm/m_cust_profile', .46, 1.45, 8.4, 4.65, (.05, .0, .95, .8))
for i, (ic, t, ds) in enumerate([('search', 'Danh sách', '1. CUSTOMER CODE'), ('users', 'Người liên hệ', '2. CBC'), ('target', 'ADP', '8. ADP'), ('eye', 'Hồ sơ 360°', 'mọi lượt gặp, đơn, thuyết trình')]):
    y = 1.55 + i * 1.12
    d.icon(s, ic, 9.2, y, .46)
    d.text(s, 9.8, y - .04, 3.07, .32, [(t, 14.5, True, 'title')])
    d.text(s, 9.8, y + .3, 3.07, .6, [(ds, 12, False, 'ink')])
N(s, 'Khách hàng và hồ sơ 360 độ.', ['Account ở tab 1, người liên hệ ở tab 2, khách trọng điểm ADP ở tab 8.', 'Hồ sơ 360 độ gom mọi lượt gặp, đơn hàng, buổi thuyết trình của một khách.'], 'Khách hàng mới tính KPI thế nào?')

s = S('2 · CRM', 'Mở mới & SKU mới', '2 · CRM', 'C35', '2.6', 'Cập nhật địa bàn: thông tin khách hàng')
pair(s, 'crm/newacc_review', 'crm/m_newacc_review', .46, 1.45, 7.6, 4.65, 'auto')
d.text(s, 8.4, 1.45, 4.47, .3, [('4 ĐIỀU KIỆN TÍNH KPI C1', 11, True, 'sec')])
steps(s, 8.4, 1.85, 4.47, ['Có trong báo cáo tuần', 'NPP xác nhận', 'Thông tin đúng mẫu', 'Có đơn hàng đầu tiên'], gap=.5, sz=13)
d.text(s, 8.4, 4.0, 4.47, .3, [('QUẢN LÝ CHỌN', 11, True, 'sec')])
for i, (t, c, ds) in enumerate([('Xác nhận', 'pos', 'tính KPI'), ('Trả lại bổ sung', 'acc', 'về nháp, giữ chứng từ'), ('Từ chối', 'neg', 'ghi lý do')]):
    y = 4.4 + i * .58
    d.chip(s, 8.4, y, t, fill=c, col='on_pri' if c != 'acc' else 'title', sz=12, h=.36)
    d.text(s, 10.45, y, 2.42, .36, [(ds, 12.5, False, 'ink')], anchor='m')
N(s, 'Mở mới.', ['Sales tạo case, tải chứng từ, gửi xin xác nhận. Dữ liệu ở tab 16, chứng từ ở Drive.', 'Quản lý xác nhận, trả lại hoặc từ chối. Chỉ case được xác nhận mới tính KPI C1.'], 'Đơn hàng.')

s = S('2 · CRM', 'Đơn hàng (Dental)', '2 · CRM', 'C10', '2.7', 'Thông tin đơn hàng (Dental sale)')
pair(s, 'crm/orders', 'crm/m_orders', .46, 1.45, 8.0, 4.65, (.16, .02, 1, .8))
for i, (v, t) in enumerate([('C5', 'Tổng tiền đơn ÷ tỷ giá tháng'), ('C6', 'SL Composite qua NPP'), ('C7', 'SL JIZAI qua NPP')]):
    d.kpi(s, 8.8, 1.45 + i * 1.25, 4.07, 1.1, v, t, vsz=26, dark=(i == 0))
note(s, 8.8, 5.25, 4.07, .85, 'Lưu ở', '3. SALE ORDER · chống ghi trùng.', fill='soft', line=None)
N(s, 'Đơn hàng và 3 mã KPI.', ['Đơn hàng nhập trên CRM vào tab 3 là nguồn tính C5, C6, C7.'], 'Kế hoạch trên MMH Calendar.')

s = S('2 · CRM', 'MMH Calendar', '2 · CRM', 'C24', '2.8', 'MMH Calendar: kế hoạch cá nhân và đội nhóm')
pair(s, 'crm/mmhcal', 'crm/m_mmhcal', .46, 1.45, 8.0, 4.65)
dtable(s, 8.75, 1.45, [1.75, 2.37], [['Loại lịch', 'Nguồn'], ['Đi địa bàn', '6. WEEKLY REPORT'], ['Công việc khác', 'ALL TASK'], ['Công tác', 'File Business Trip'],
       ['Đào tạo', 'Training Master'], ['Sự kiện · bài đăng', 'MKT Offline / Online']], row_h=.5, head=12, body=12)
d.text(s, 8.75, 4.7, 4.12, 1.4, [('Team Leader chọn Cả đội hoặc từng PIC; công tác người khác ẩn, quản lý bật bằng công tắc.', 12.5, False, 'ink')])
N(s, 'Lịch chung.', ['MMH Calendar gom kế hoạch của mình và của đội: đi địa bàn, công việc khác, công tác, đào tạo, sự kiện, bài đăng. Mỗi loại đọc từ file gốc của nó.'], 'Kế hoạch công tác.')

s = S('2 · CRM', 'Kế hoạch & báo cáo công tác', '2 · CRM', 'C26', '2.8', 'MMH Calendar: kế hoạch cá nhân và đội nhóm')
chev(s, ['Đề xuất', 'Giám đốc duyệt', 'Đi công tác', 'Báo cáo'], y=1.42, h=.55)
pair(s, 'crm/trip_propose', 'crm/m_trip_propose', .46, 2.2, 6.1, 3.9, 'auto')
pair(s, 'crm/trip_report', 'crm/m_trip_report', 6.77, 2.2, 6.1, 3.9, 'auto')
N(s, 'Công tác.', ['Đề xuất ghi vào file Business Trip và gửi email cho Giám đốc. Duyệt xong, thẻ công tác trên lịch đổi trạng thái; sau chuyến viết báo cáo ngay trên thẻ.'], 'Hệ thống báo cáo.')

s = S('2 · CRM', 'Báo cáo tuần, tháng', '2 · CRM', 'C29', '2.9', 'Hệ thống báo cáo: tuần, tháng, KPI')
pair(s, 'crm/send_report', 'crm/m_send_report', .46, 1.45, 7.8, 4.65)
steps(s, 8.6, 1.5, 4.27, [('Tự dựng', 'từ đi địa bàn, công việc khác, CV tháng'), ('Người nhận', 'theo sơ đồ tổ chức'), ('Kèm PDF', 'lưu Drive'), ('Gửi thử', 'cho chính mình trước')], gap=1.1)
N(s, 'Gửi báo cáo.', ['Báo cáo dựng từ chính dữ liệu đã ghi, không gõ lại. Chọn kỳ, xem trước, gửi thử hoặc gửi chính thức, kèm PDF.'], 'Họp tuần.')

s = S('2 · CRM', 'Họp tuần', '2 · CRM', 'C18', '2.9', 'Hệ thống báo cáo: tuần, tháng, KPI')
for i, (ic, t, ds) in enumerate([('chart', 'Forecast NPP', 'số liệu tự lấy từ file Turnover'), ('users', 'Quan hệ KH & thị trường', 'theo từng PIC'), ('file', 'Tình hình thầu', 'gói thầu, giá theo mã'), ('check', 'So sánh tuần trước', '4 phần biên bản')]):
    y = 1.55 + i * 1.12
    d.icon(s, ic, .46, y, .5)
    d.text(s, 1.1, y - .04, 3.6, .34, [(t, 15, True, 'title')])
    d.text(s, 1.1, y + .32, 3.6, .6, [(ds, 12.5, False, 'ink')])
pair(s, 'crm/weekly_meeting', 'crm/m_meet', 4.9, 1.45, 7.97, 4.65, (.16, .03, 1, .9))
N(s, 'Họp tuần.', ['Biên bản họp tuần có bốn phần; forecast tự lấy từ file doanh số.'], 'Báo cáo KPI.')

s = S('2 · CRM', 'KPI & báo cáo KPI', '2 · CRM', 'C25', '2.9', 'Hệ thống báo cáo: tuần, tháng, KPI')
pair(s, 'crm/kpi_tab', 'crm/m_kpi', .46, 1.45, 6.3, 4.65)
pair(s, 'crm/kpi_report', 'crm/m_kpi_report', 6.97, 1.45, 5.9, 4.65, 'auto')
N(s, 'KPI và báo cáo KPI tháng.', ['Trang KPI & Mục tiêu đọc số từ file KPI. Báo cáo KPI tháng điền sẵn người nhận, xem trước rồi gửi; tự đánh giá quý tải file Word.'], 'Biểu đồ thống kê.')

s = S('2 · CRM', 'Biểu đồ thống kê', '2 · CRM', 'C09', '2.10', 'Biểu đồ thống kê: tiến độ công việc, KPI')
pair(s, 'crm/dashboard', 'crm/m_dashboard', .46, 1.45, 8.3, 4.65)
for i, (t, ds) in enumerate([('Dashboard', 'Cả FY · quý · tháng; doanh số, viếng thăm, đơn'), ('Tiến độ', 'Công việc khác: key task, quá hạn'), ('KPI', 'xu hướng theo tháng, đóng góp theo nhóm')]):
    y = 1.5 + i * 1.5
    d.text(s, 9.1, y, 3.77, .4, [(t, 17, True, 'title')])
    d.text(s, 9.1, y + .42, 3.77, .9, [(ds, 13, False, 'muted')])
N(s, 'Biểu đồ.', ['Dashboard tổng quan theo năm, quý, tháng; tiến độ công việc khác; biểu đồ KPI.'], 'Dành cho quản lý.')

s = S('2 · CRM', 'Dành cho quản lý', '2 · CRM', 'C28', '2.12', 'Dành cho quản lý và Admin')
pair(s, 'crm/trip_approve', 'crm/m_trip_approve', .46, 1.45, 7.6, 4.65, 'auto')
for i, (t, ds) in enumerate([('Đổi nhóm', 'Dental · Surgical · Eyeless · Thái'), ('Duyệt công tác', 'Giám đốc, nhiều chuyến một lần'), ('Duyệt mở mới', 'Team Leader, HOD'), ('Chế độ TEST', 'Admin thử mọi chức năng (Phần 6)')]):
    y = 1.5 + i * 1.12
    d.badge(s, 8.4, y, i + 1, d=.34, sz=12)
    d.text(s, 8.9, y - .04, 3.97, .34, [(t, 15, True, 'title')])
    d.text(s, 8.9, y + .32, 3.97, .5, [(ds, 12.5, False, 'muted')])
N(s, 'Công cụ quản lý.', ['Đổi nhóm đang xem, Giám đốc duyệt công tác, quản lý duyệt mở mới, Admin có chế độ TEST mình nói ở phần sáu.'], 'Sang phần ba: Report Hub.')

# ═════════════════════════ PHẦN 3 · REPORT HUB
s = divider('3', 'MMH Report Hub', 'Bảng việc chung của các phòng ban', ['Dữ liệu: nguồn, cách lưu, lưu ở đâu', 'Cập nhật công việc', 'MMH Calendar', 'Báo cáo · biểu đồ', 'Giao việc · công tác'], 'rh/work_cal_week', '3.1', 'rh/m_work_cal')
N(s, 'Mở phần Report Hub.', ['Report Hub là nơi các phòng ban cập nhật việc, tự tạo báo cáo, xem lịch chung. Cùng khung: dữ liệu trước, chức năng sau.'], 'Ba phòng ban, một cách làm.')

s = S('3 · Report Hub', '3 phòng ban, 1 cách làm', '3 · REPORT HUB', 'C36', '3.1', 'Ai dùng, bố cục màn hình')
for i, (c, t, ds) in enumerate([('pri', 'Marketing & Sales', 'Product, Event, Digital, Design'), ('sec', 'Back Office', 'Kho, kế toán, nhập hàng'), ('ter', 'Management', 'Việc ban quản lý')]):
    x = .46 + i * 4.17
    d.box(s, x, 1.45, 4.0, 1.2, c)
    col = 'on_pri' if c != 'ter' else 'title'
    d.text(s, x + .22, 1.52, 3.6, .5, [(t, 17, True, col)], anchor='m')
    d.text(s, x + .22, 2.02, 3.6, .5, [(ds, 12.5, False, col)])
pair(s, 'rh/main', 'rh/m_task_list', .46, 2.9, 8.1, 3.2, (0, 0, 1, .8))
steps(s, 8.9, 2.95, 3.97, ['Mỗi phòng một file dữ liệu', 'Cùng cách cập nhật', 'Báo cáo tự tạo', 'Tìm nhanh Ctrl K'], gap=.75, sz=13.5)
N(s, 'Ba phòng ban.', ['Marketing và Sales, Back Office, Management: mỗi phòng một file, cùng một cách cập nhật.'], 'Dữ liệu Report Hub nằm ở đâu?')

s = S('3 · Report Hub', 'Dữ liệu: lưu ở đâu', '3 · REPORT HUB', 'C33', '3.2', 'Dữ liệu của Report Hub: nguồn, cách lưu, lưu ở đâu')
chev(s, ['Tạo / sửa trên app', 'Hàng đợi ghi', 'Backend đúng loại', 'Tổng hợp'], y=1.42, h=.6)
dtable(s, .46, 2.25, [2.9, 3.3, 3.5, 2.71], [['Dữ liệu', 'Tạo / sửa ở đâu', 'Lưu ở', 'Đi tiếp tới'],
       ['Key task / Sub-task', '+ Thêm việc · thẻ việc · kiểu Bảng', 'File của phòng ▸ tab theo dõi tuần', 'Lịch, tiến độ, báo cáo tuần'],
       ['Hạng mục tháng · biên bản họp', 'Monthly Report · Biên bản họp', 'File của phòng', 'Báo cáo tháng, email'],
       ['Giao việc + trao đổi', 'Thêm việc cho người khác · Việc đã giao', 'Training Master ▸ RH_Assign', 'Popup, email'],
       ['Công tác', 'Đề xuất · Duyệt · báo cáo', 'File Business Trip', 'MMH Calendar, email'],
       ['Sự kiện offline', 'MMH Calendar ▸ Sửa', 'MKT Offline ▸ REPORT_EVENT', 'KPI C12, C15'],
       ['Bài đăng, thiết kế', 'MMH Calendar ▸ Sửa', 'MKT Online ▸ EVENT / PRODUCT REPORT', 'KPI C17'],
       ['Buổi đào tạo', 'MMH Calendar ▸ Tạo mới', 'Training Master', 'Training Hub, KPI L1 · L2'],
       ['KPI (chỉ xem)', 'KPI', 'MMH_KPI_FY68', 'Báo cáo KPI tháng']], row_h=.45, head=11.5, body=11)
N(s, 'Report Hub không có cơ sở dữ liệu riêng.', ['Mỗi loại dữ liệu ghi thẳng vào file nguồn của nó: việc của phòng ở file phòng, công tác ở file Business Trip, sự kiện và bài đăng ở file Marketing, đào tạo và giao việc ở Training Master.'], 'Bốn nhóm chức năng.')

s = S('3 · Report Hub', '4 nhóm chức năng', '3 · REPORT HUB', 'C27', '3.3', 'Chức năng chính')
FN = [('pri', 'Cập nhật công việc', 'Key task, sub task, tiến độ, kết quả, giao việc', 'rh/task_list'), ('sec', 'MMH Calendar', 'Lịch làm việc, công tác, sự kiện, bài đăng, đào tạo', 'rh/mmh_cal_month'),
      ('ter', 'Hệ thống báo cáo', 'Báo cáo tuần, Monthly Report, biên bản họp, báo cáo KPI', 'rh/weekly_report'), ('acc', 'Biểu đồ thống kê', 'Tiến độ, Tổng quan, Tình hình đội, KPI', 'rh/overview')]
for i, (c, t, ds, img) in enumerate(FN):
    x = .46 + i * 3.15
    shot(s, img, x, 1.45, 2.95, 1.9)
    d.box(s, x, 3.5, 2.95, .08, c)
    d.text(s, x, 3.68, 2.95, .45, [(t, 16, True, 'title')])
    d.text(s, x, 4.15, 2.95, 1.3, [(ds, 13, False, 'ink')])
N(s, 'Bốn nhóm chức năng.', ['Cập nhật công việc, MMH Calendar, hệ thống báo cáo và biểu đồ thống kê.'], 'Cập nhật công việc.')

s = S('3 · Report Hub', 'Cập nhật công việc', '3 · REPORT HUB', 'C24', '3.4', 'Cập nhật công việc: key task và sub task')
pair(s, 'rh/task_list', 'rh/m_task_list', .46, 1.45, 8.0, 4.65, (.18, .0, 1, .75))
steps(s, 8.8, 1.5, 4.07, [('Key task', 'loại việc, PIC, hạn'), ('Sub task', 'PIC, hạn, % tiến độ'), ('Cập nhật', 'trạng thái, kết quả, file'), ('Lưu ở', 'file của phòng')], gap=1.1)
N(s, 'Hai tầng công việc.', ['Mỗi việc lớn là một key task, chia thành sub task có người phụ trách, hạn và phần trăm tiến độ.'], 'Việc hiện trên lịch làm việc.')

s = S('3 · Report Hub', 'Lịch làm việc', '3 · REPORT HUB', 'C07', '3.5', 'MMH Calendar: lịch làm việc, Pending, Today task')
pair(s, 'rh/work_cal_week', 'rh/m_work_cal', .46, 1.45, 8.4, 4.65, (.15, .0, 1, .9))
d.box(s, 9.15, 1.45, 3.72, 4.65, 'panel')
d.text(s, 9.4, 1.6, 3.3, .3, [('CÁCH XEM', 11, True, 'sec')])
bullets(s, 9.4, 1.95, 3.3, ['Today task · Tuần · Tháng · Danh sách', 'Chỉ tôi / Cả đội', 'Bật / tắt công tác, đào tạo'], 13, 6)
d.text(s, 9.4, 3.75, 3.3, .3, [('PENDING', 11, True, 'sec')])
bullets(s, 9.4, 4.1, 3.3, ['Trễ hạn · đến hạn · đang làm', 'Hoàn thành, Gia hạn'], 13, 6)
N(s, 'Lịch tuần và Pending.', ['Lịch hiện sub task theo ngày; Pending gom việc trễ, đến hạn và đang làm.'], 'Biểu đồ thống kê.')

s = S('3 · Report Hub', 'Biểu đồ thống kê', '3 · REPORT HUB', 'C09', '3.6', 'Biểu đồ thống kê: tiến độ, tổng quan, tình hình đội')
pair(s, 'rh/overview', 'rh/m_overview', .46, 1.45, 6.2, 4.65, (0, 0, 1, .6))
pair(s, 'rh/task_progress', 'rh/m_progress', 6.87, 1.45, 6.0, 4.65)
N(s, 'Tổng quan và tiến độ.', ['Tổng quan có số việc, hiệu suất theo người, biểu đồ theo loại và trạng thái. Tiến độ công việc là một thanh cho mỗi key task.'], 'Báo cáo tuần tự tạo.')

s = S('3 · Report Hub', 'Báo cáo tuần tự tạo', '3 · REPORT HUB', 'C42', '3.7', 'Hệ thống báo cáo: tuần, tháng, biên bản họp')
for i, (lb, t) in enumerate([('KẾT QUẢ TUẦN', 'Task chạm vào tuần + task quá hạn chưa xong'), ('PLAN TUẦN SAU', 'Task chưa xong chạm vào tuần sau')]):
    y = 1.45 + i * 1.3
    d.box(s, .46, y, 4.6, 1.15, 'panel')
    d.text(s, .7, y + .12, 4.2, .28, [(lb, 11, True, 'sec')])
    d.text(s, .7, y + .44, 4.2, .7, [(t, 14, False, 'title')])
note(s, .46, 4.15, 4.6, 1.95, 'Gửi:', 'gửi thử cho chính mình, hoặc Team Leader, HOD và đồng nghiệp; email kèm PDF.', fill='soft', line=None)
pair(s, 'rh/weekly_report', 'rh/m_weekly_report', 5.35, 1.45, 7.52, 4.65)
N(s, 'Báo cáo tuần.', ['Không gõ lại: kết quả tuần và plan tuần sau lấy từ task.'], 'Monthly Report và biên bản họp.')

s = S('3 · Report Hub', 'Monthly · biên bản họp', '3 · REPORT HUB', 'C29', '3.7', 'Hệ thống báo cáo: tuần, tháng, biên bản họp')
pair(s, 'rh/monthly_report', 'rh/m_monthly', .46, 1.45, 6.1, 4.65)
pair(s, 'rh/minutes', 'rh/m_minutes', 6.77, 1.45, 6.1, 4.65)
N(s, 'Báo cáo tháng và biên bản họp.', ['Monthly Report mỗi mục một thẻ; biên bản họp có thành phần, nội dung theo mục, follow-up, gửi email.'], 'Lịch chung của công ty.')

s = S('3 · Report Hub', 'MMH Calendar', '3 · REPORT HUB', 'C34', '3.8', 'MMH Calendar: lịch chung của công ty')
pair(s, 'rh/mmh_cal_month', 'rh/m_mmh_cal', .46, 1.45, 8.0, 4.65, (0, 0, 1, .85))
dtable(s, 8.75, 1.45, [1.8, 2.32], [['Loại lịch', 'Nguồn'], ['Công tác', 'File Business Trip'], ['Đào tạo', 'Training Master'], ['Sự kiện offline', 'MKT Offline'], ['Bài đăng', 'MKT Online']], row_h=.55, head=12, body=12.5)
d.text(s, 8.75, 4.45, 4.12, 1.6, [('Nhấp đúp một ngày để tạo mới; bấm một mục để xem, sửa ngay.', 13, False, 'ink')])
N(s, 'Lịch chung bốn nguồn.', ['MMH Calendar gom công tác, đào tạo, sự kiện và bài đăng; mỗi loại đọc và ghi vào file gốc.'], 'Sự kiện, bài đăng sửa ngay trên lịch.')

s = S('3 · Report Hub', 'Sự kiện · bài đăng · đào tạo', '3 · REPORT HUB', 'C10', '3.10', 'Đào tạo trên lịch')
shot(s, 'rh/offline_edit', .46, 1.45, 4.0, 4.3, (0, 0, 1, .75))
shot(s, 'rh/online_edit', 4.63, 1.45, 4.0, 4.3, (0, 0, 1, .75))
phone(s, 'rh/m_training', 8.85, 1.45, 4.3)
d.text(s, .46, 5.85, 4.0, .4, [('Sự kiện · REPORT_EVENT', 13, True, 'title')])
d.text(s, 4.63, 5.85, 4.0, .4, [('Bài đăng · EVENT REPORT', 13, True, 'title')])
d.text(s, 8.85, 5.85, 3.0, .4, [('Đào tạo · điện thoại', 13, True, 'title')])
N(s, 'Sửa không cần mở Sheet.', ['Sự kiện và bài đăng sửa ngay trên lịch; buổi đào tạo xem trên điện thoại, kéo thả tài liệu trên máy tính.'], 'Công tác.')

s = S('3 · Report Hub', 'Công tác', '3 · REPORT HUB', 'C20', '3.11', 'MMH Calendar: kế hoạch và báo cáo công tác')
chev(s, ['Đề xuất', 'Giám đốc duyệt', 'Lịch chung', 'Báo cáo'], y=1.42, h=.55)
pair(s, 'rh/trip_propose', 'rh/m_trip_propose', .46, 2.2, 6.1, 3.9, 'auto')
pair(s, 'rh/trip_approve', 'rh/m_trip_approve', 6.77, 2.2, 6.1, 3.9, 'auto')
N(s, 'Công tác.', ['Đề xuất ghi file Business Trip, email Giám đốc; Giám đốc duyệt trên app, mỗi người nhận một email riêng; chuyến đi hiện trên lịch chung.'], 'Giao việc.')

s = S('3 · Report Hub', 'Giao việc', '3 · REPORT HUB', 'C40', '3.12', 'Giao việc và trao đổi')
chev(s, ['Giao', 'Nhận', 'Hoàn thành'], y=1.42, h=.55)
shot(s, 'rh/assign_mail', .46, 2.2, 4.6, 3.9, 'auto')
pair(s, 'rh/assign_track', 'rh/m_assign_track', 5.3, 2.2, 7.57, 3.9, 'auto')
N(s, 'Giao việc có theo dõi.', ['Người giao soạn email ngay; người nhận thấy popup; xong việc hệ thống báo người giao. Lưu ở Training Master, tab RH_Assign.'], 'Sang phần bốn: KPI.')

# ═════════════════════════ PHẦN 4 · KPI (theo "Hướng dẫn hệ thống KPI FY68")
import re as _re
def plain(t): return _re.sub(r'<[^>]+>', '', t)
def stchip(sl, x, y, code, w=None, sz=10):
    t = STN[code]; w = w or (.25 + .07 * len(t))
    fill = {'A': 'E3F3E8', 'S': 'ECEFF3', 'W': 'FFF3CC', 'M': 'FBE5E2', 'N': 'FFFFFF'}[code]; col = {'A': '1E6B34', 'S': '4A5866', 'W': '8A6200', 'M': 'A23B30', 'N': '7A8794'}[code]
    d.box(sl, x, y, w, .28, fill, line='D5DCE4' if code == 'N' else None, adj=.5); d.text(sl, x, y, w, .28, [(t, sz, True, col)], 'c', 'm'); return w

s = divider('4', 'Hệ thống KPI FY68', 'Số KPI lấy từ đâu, tính thế nào, xem ở đâu', ['17 nhóm · 22 người · 4 phòng ban', 'Nguồn dữ liệu', 'Quy tắc tính', 'Xem KPI ở đâu', 'Theo phòng ban · việc cần chốt'], 'rh/kpi_tab', '4.1', 'rh/m_kpi')
N(s, 'Mở phần KPI.', ['Phần này đi theo tài liệu Hướng dẫn hệ thống KPI FY68 của Product Team: mỗi KPI ghi rõ lấy dữ liệu từ file nào, sheet nào, tính thế nào, điều kiện được tính và trạng thái kết nối.'], 'KPI gồm những nhóm nào.')

s = S('4 · KPI', '17 nhóm KPI FY68', '4 · KPI', 'C27', '4.1', 'Phần 1 · Về hệ thống KPI')
G17 = [('Kinh doanh', 'pri', [('Doanh số', 'F0 – F5'), ('Doanh số SP mới', 'F6'), ('Mở mới', 'C1 – C2'), ('Viếng thăm KH', 'C3 – C4'), ('Bán hàng hỗ trợ', 'C5 – C7')]),
       ('Marketing', 'acc', [('Marketing hỗ trợ bán', 'C8 – C9'), ('Sự kiện', 'C10 – C14'), ('Người tham dự', 'C15'), ('Ra mắt sản phẩm', 'C16'), ('Digital & Design', 'C17')]),
       ('Back Office', 'sec', [('Tài liệu', 'P1'), ('Đơn hàng & giao nhận', 'P2'), ('Tồn kho', 'P3'), ('Dự án', 'P4'), ('Tài chính & hành chính', 'P5')]),
       ('Con người', 'ter', [('Đào tạo', 'L1 – L2'), ('Hành vi (KBI)', 'L3')])]
for r, (area, c, items) in enumerate(G17):
    y = 1.42 + r * 1.2
    d.box(s, .46, y, 1.9, 1.05, c); d.text(s, .56, y, 1.7, 1.05, [(area, 15, True, 'on_pri' if c in ('pri', 'sec') else 'title')], 'l', 'm')
    for i, (t, code) in enumerate(items):
        x = 2.5 + i * 2.08
        d.box(s, x, y, 1.98, 1.05, 'panel')
        d.text(s, x + .12, y + .1, 1.78, .55, [(t, 13, True, 'title')])
        d.text(s, x + .12, y + .68, 1.78, .3, [(code, 11, True, 'sec')])
d.text(s, .46, 6.25, 12.4, .3, [('Tên nhóm đúng như trên Report Hub / CRM · mã nhóm để đối chiếu với file KPI · 22 người có KPI · 4 phòng ban', 11.5, False, 'muted')])
N(s, '17 nhóm KPI.', ['Kinh doanh 5 nhóm, Marketing 5, Back Office 5, Con người 2. Điểm hành vi có ở cả 22 người.'], 'Số KPI chạy thế nào?')

s = S('4 · KPI', 'Số KPI đi từ đâu tới đâu', '4 · KPI', 'C19', '4.1', 'Phần 1 · Về hệ thống KPI')
chev(s, ['Nhập hằng ngày', 'Tự đổ vào file KPI', 'Tính % và tổng', 'Xem trên app'], y=1.45, h=.72,
     sub=['Sales ghi đơn, đi địa bàn, mở mới trên CRM; kế toán chốt doanh số; Marketing ghi sự kiện; Back Office chấm điểm',
          'Sửa file nguồn ⇒ vài giây; lưu trên app ⇒ ~1 phút; 15 phút đọc lại toàn bộ. Ghi vào MMH_KPI_FY68 ▸ 3. Detail KPI',
          '% đạt = thực đạt ÷ mục tiêu; × trọng số ⇒ Tổng KPI tháng, quý, năm (sheet 5, 6)', 'Report Hub và CRM hiện đúng số của file KPI: của bạn, của nhóm, Total KPI, email KPI tháng'], subh=1.6)
note(s, .46, 4.15, 12.41, 1.0, 'App chỉ đọc số từ file KPI.', 'Số trên app sai ⇒ sửa ở file nguồn (bước 1), không sửa trên app.', fill='hi')
d.table(s, .46, 5.35, [3.1, 3.1, 3.1, 3.11], [['Giám đốc', 'Trưởng phòng', 'Team Leader', 'Nhân viên'], ['Xem tất cả', 'Cả phòng', 'Cả nhóm', 'KPI của mình']], row_h=.42, head=12, body=12.5, aligns=['c'] * 4)
tb = d.text(s, .46, 6.3, 6, .25, [('', 10, False, 'muted')]); link_run(tb.text_frame.paragraphs[0], 'Mở MMH_KPI_FY68 ▸ 3. Detail KPI ↗', KF['KPI'][1], 10)
N(s, 'Bốn bước.', ['Mọi người chỉ nhập dữ liệu vào file hoặc app quen thuộc. Máy tự đổ số vào file KPI, tự tính phần trăm và tổng, app hiện lại đúng số đó.'], 'Đã tự động được bao nhiêu?')

s = S('4 · KPI', 'Đã tự động được bao nhiêu', '4 · KPI', 'C16', '4.1', 'Phần 1 · Về hệ thống KPI')
for i, (v, c, t, ds) in enumerate([('92', 'A', 'Đã tự động kết nối', 'Doanh số, CRM VN, Marketing VN, Đào tạo, Hành vi'), ('13', 'S', 'Tự cộng từ dòng con', 'VD: Viếng thăm · VN = cộng từng sales'),
                                    ('6', 'W', 'Đã kết nối · chưa nhận số', 'Sự kiện Surgical Thái Lan'), ('38', 'M', 'Chưa kết nối · nhập tay', 'Back Office, Dental Thái Lan, coaching')]):
    x = .46 + i * 3.15
    d.box(s, x, 1.45, 2.95, 2.1, {'A': 'E3F3E8', 'S': 'ECEFF3', 'W': 'FFF3CC', 'M': 'FBE5E2'}[c])
    d.text(s, x + .2, 1.55, 2.6, .8, [(v, 40, True, STC[c] if c != 'S' else '4A5866')])
    d.text(s, x + .2, 2.35, 2.6, .35, [(t, 13.5, True, 'title')])
    d.text(s, x + .2, 2.72, 2.6, .8, [(ds, 11.5, False, 'ink')])
d.text(s, .46, 3.75, 12.4, .3, [('Trên toàn bộ 149 dòng KPI của sheet 3. Detail KPI · tỉ lệ trọng số tự động theo phòng ban:', 12.5, True, 'title')])
DEP = {}
for dept, sub_, ppl in KPEOPLE:
    key = 'Back Office' if 'Back' in dept else 'Thái Lan' if 'Thái' in dept else 'Kinh doanh & Marketing VN'
    for name, title, ks in ppl:
        for k, w in ks: DEP.setdefault(key, {}).setdefault(kstat(k), 0); DEP[key][kstat(k)] += w
for i, (dep, m) in enumerate(DEP.items()):
    y = 4.2 + i * .62; tot = sum(m.values()); x = 3.6
    d.text(s, .46, y, 3.0, .45, [(dep, 13.5, True, 'title')], anchor='m')
    for c in 'ASWMN':
        if not m.get(c): continue
        w = 8.7 * m[c] / tot; d.box(s, x, y + .06, w, .34, STC[c], kind=MSO_SHAPE.RECTANGLE)
        if w > .55: d.text(s, x, y + .06, w, .34, [(f'{round(100 * m[c] / tot)}%', 10.5, True, 'on_pri' if c in 'AM' else 'title')], 'c', 'm')
        x += w
lx = .46
for c in 'ASWMN':
    d.box(s, lx, 6.2, .2, .2, STC[c], kind=MSO_SHAPE.RECTANGLE); d.text(s, lx + .27, 6.13, 2.3, .32, [(STN[c], 10.5, False, 'muted')], anchor='m'); lx += 2.48
N(s, 'Mức độ tự động.', ['92 dòng đã tự động, 13 dòng tự cộng, 6 dòng đã nối nhưng chưa có số, 38 dòng còn nhập tay.', 'Kinh doanh và Marketing Việt Nam gần như 100 phần trăm tự động; Back Office và Dental Thái Lan còn nhiều nhập tay.'], 'Các nguồn dữ liệu cụ thể.')

def srcslide(title, sec, rows, heads, cw, kind):
    s = S('4 · KPI', title, '4 · KPI', kind, sec, 'Phần 2 · Nguồn dữ liệu')
    x0 = .46; y = 1.42
    for j, h in enumerate(heads):
        d.text(s, x0 + sum(cw[:j]), y, cw[j] - .1, .3, [(h.upper(), 10.5, True, 'sec')])
    d.hline(s, .46, y + .34, 12.87, 'line')
    for i, r in enumerate(rows):
        yy = y + .45 + i * .6
        for j, v in enumerate(r):
            x = x0 + sum(cw[:j])
            if isinstance(v, tuple) and v and v[0] == 'ST': stchip(s, x, yy + .1, v[1]); continue
            tb = d.text(s, x, yy, cw[j] - .12, .58, [('', 10.5, False, 'ink')]); p = tb.text_frame.paragraphs[0]
            for k, part in enumerate(v if isinstance(v, list) else [v]):
                if k: link_run(p, ' · ', None, 10.5, False, 'muted')
                if isinstance(part, tuple): link_run(p, part[0], part[1], 10.5 if j == 0 else 10.5, j == 0, 'sec')
                else: link_run(p, part, None, 10.5, j == 0, 'title' if j == 0 else 'ink')
        d.hline(s, .46, yy + .58, 12.87, 'line')
    return s
L_ = lambda k, nm=None: ((nm or KF[k][0]) + ' ↗', KF[k][1])
s = srcslide('Nguồn đã tự động kết nối', '4.2', [
    [[L_('TURN')], 'Kế toán', 'Doanh số · Doanh số sản phẩm mới', ('ST', 'A')],
    [[L_('DEN', 'Dental CRM'), L_('SUR', 'Surgical CRM'), L_('EYE', 'Eyeless CRM')], 'Sales VN (app CRM)', 'Mở mới · Viếng thăm · Bán hàng hỗ trợ · Sự kiện', ('ST', 'A')],
    [[L_('THS')], 'Miew (app CRM)', 'Mở mới · Viếng thăm (Manipler)', ('ST', 'A')],
    [[L_('MKO')], 'Marketing', 'Marketing hỗ trợ bán · Sự kiện KOL · Người tham dự', ('ST', 'A')],
    [[L_('DMA', 'Digital ▸ ADS'), L_('DME', 'EVENT + PRODUCT'), L_('DMO', 'ORDER REPORT')], 'Minh Trang, Duc Anh', 'Digital & Design', ('ST', 'A')],
    [[L_('TRN')], 'Người đứng lớp', 'Đào tạo · Đào tạo khách hàng', ('ST', 'A')],
    [[L_('KBI')], 'Quản lý trực tiếp', 'Hành vi', ('ST', 'A')],
    [[L_('THM')], 'Surgical Thái Lan', 'Sự kiện Surgical Thái Lan', ('ST', 'W')]],
    ['Nguồn dữ liệu (bấm để mở)', 'Ai nhập', 'Đổ số cho nhóm KPI', 'Trạng thái'], [5.1, 2.1, 3.3, 1.9], 'C33')
N(s, 'Nguồn đã kết nối.', ['Tám nguồn đang tự đổ số. Bấm tên file để mở đúng sheet. Riêng sự kiện Surgical Thái Lan đã nối nhưng chưa nhận số.'], 'Các nguồn chưa kết nối.')
s = srcslide('Nguồn chưa kết nối', '4.2', [
    [[L_('BO')], 'Báo cáo, xử lý đơn, giao hàng, Bravo, ngân hàng, chi phí, tồn kho', 'Back Office', 'Đã có, chưa nối'],
    [[L_('OPT', 'Order Processing Tracking'), L_('DAF', 'Distributor forecast')], 'Xác nhận lịch giao, đơn hàng đúng hạn', 'Dam Ha, Dam Viet', 'Đã có, chưa nối'],
    [[L_('IMX'), L_('TND', 'Tender documents KPI')], 'Thông quan, hồ sơ thầu', 'Hoa, Dam Viet, Hau', 'Đã có, chưa nối'],
    [[L_('TDA', 'Thailand Dental Action FY67'), L_('TDE', 'Dental Event Report FY67')], 'Đi kèm, đào tạo, giới thiệu SP, sự kiện (Dental Thái)', 'Yong, Man, Sui', 'File FY67, chưa nối'],
    ['Sổ SOP · chuyển văn phòng · số đăng ký SP', 'Tài liệu, dự án, số đăng ký', 'Dam Ha, Hoa', 'Chưa thấy file'],
    ['Báo cáo NPP Thái · KH trọng điểm · bài test · digital Thái', 'Mở mới KH trọng điểm, bài test, nội dung, ra mắt SP', 'Yong, Man, Sui', 'Chưa thấy file'],
    ['Form đánh giá Field Coaching', 'Đi kèm coaching Sales', 'Tuyen, Giang', 'Chưa thấy file']],
    ['Nguồn dữ liệu', 'Dùng cho KPI', 'Người phụ trách', 'Tình trạng file'], [4.6, 4.0, 2.1, 1.7], 'C34')
N(s, 'Nguồn chưa kết nối.', ['Các KPI này đang gõ tay vào file KPI. File đã có thì bấm để mở; các file chưa thấy cần người phụ trách cung cấp.'], 'Quy tắc tính chung.')

s = S('4 · KPI', 'Quy tắc chung', '4 · KPI', 'C42', '4.3', 'Phần 3 · Quy tắc tính (Rules)')
RU = [('% đạt', 'Thực đạt ÷ Mục tiêu của tháng. VD 27 / 30 lượt ⇒ 90%'), ('Tối đa 130%', 'Mỗi KPI (tháng, quý, năm); Doanh số và SP mới không giới hạn'), ('Mục tiêu = 0', 'Tháng đã qua ⇒ tính 100%, đủ trọng số'),
      ('Tổng KPI', 'Cộng (% đạt × trọng số); tổng trọng số 1 người = 100%'), ('Quý và năm', 'Cộng các tháng; KPI % hoặc điểm lấy trung bình'), ('Hạn báo cáo', 'Trước ngày 2 tháng sau, gửi trưởng bộ phận + Giám đốc, CC Ms. Hoa'),
      ('Trễ báo cáo', '10% trễ / thiếu ⇒ trừ 20% thưởng; > 50% ⇒ không xét thưởng (MMH-009)'), ('Màu kết quả', '< 90% xám · 90 – < 100% xanh dương · 100 – 120% vàng · > 120% xanh lá')]
for i, (t, ds) in enumerate(RU):
    x = .46 + (i % 4) * 3.15; y = 1.45 + (i // 4) * 2.35
    d.box(s, x, y, 2.95, 2.15, 'hi' if i == 0 else 'panel', line='hi_line' if i == 0 else None)
    d.text(s, x + .2, y + .18, 2.6, .45, [(t, 17, True, 'title')])
    d.text(s, x + .2, y + .7, 2.6, 1.4, [(ds, 13, False, 'ink')])
N(s, 'Tám quy tắc chung theo sheet 4. Rule.', ['Phần trăm đạt bằng thực đạt chia mục tiêu, tối đa 130 phần trăm trừ doanh số; mục tiêu bằng 0 thì tính đủ 100.', 'Tổng KPI là cộng phần trăm nhân trọng số. Hạn báo cáo trước ngày 2 tháng sau.'], 'Từng nhóm KPI.')

def grpslide(title, names, kind):
    s = S('4 · KPI', title, '4 · KPI', kind, '4.3', 'Phần 3 · Quy tắc tính (Rules)')
    gs = [g for g in KG if g[0] in names]
    cw = [2.3, 3.5, 3.4, 1.95, 1.26]; heads = ['Nhóm KPI', 'Lấy dữ liệu từ', 'Cách tính', 'Điều kiện', 'Trạng thái']
    for j, h in enumerate(heads): d.text(s, .46 + sum(cw[:j]), 1.42, cw[j] - .1, .3, [(h.upper(), 10.5, True, 'sec')])
    d.hline(s, .46, 1.76, 12.87)
    rh = min(1.12, 4.6 / len(gs))
    for i, (g, sub_, rows, (how, cond)) in enumerate(gs):
        y = 1.86 + i * rh
        codes = _re.search(r'Mã ([^·]+)', sub_); d.text(s, .46, y, cw[0] - .1, rh - .05, [[(g, 13, True, 'title')], [((codes.group(1).strip() if codes else ''), 10.5, True, 'sec')]])
        srcs = []
        for r in rows:
            for m_ in _re.finditer(r'<a href="([^"]+)">([^<]+)</a>', r[3]):
                if m_.group(1) not in [u for _, u in srcs]: srcs.append((m_.group(2), m_.group(1)))
        tb = d.text(s, .46 + sum(cw[:1]), y, cw[1] - .12, rh - .05, [('', 10, False, 'ink')]); p = tb.text_frame.paragraphs[0]
        for k, (nm, u) in enumerate(srcs[:3]):
            if k: link_run(p, '\n', None, 10, False)
            link_run(p, nm[:52] + ('…' if len(nm) > 52 else '') + ' ↗', u, 10)
        if not srcs: link_run(p, 'Chưa có file nguồn', None, 10, False, 'muted')
        h1 = plain(how[0]); h1 = h1 if len(h1) < 150 else h1[:147] + '…'
        c1 = plain(cond[0]); c1 = c1 if len(c1) < 80 else c1[:77] + '…'
        d.text(s, .46 + sum(cw[:2]), y, cw[2] - .12, rh - .05, [(h1, 10.5, False, 'ink')])
        d.text(s, .46 + sum(cw[:3]), y, cw[3] - .12, rh - .05, [(c1, 10, False, 'muted')])
        cnt = {}
        for r in rows: cnt[r[4]] = cnt.get(r[4], 0) + 1
        yy = y
        for c in 'AWSMN':
            if cnt.get(c): d.box(s, .46 + sum(cw[:4]), yy + .05, .16, .16, STC[c], kind=MSO_SHAPE.OVAL); d.text(s, .46 + sum(cw[:4]) + .22, yy, 1.05, .26, [(f'{cnt[c]} · ' + {'A': 'tự động', 'W': 'chưa có số', 'S': 'tự cộng', 'M': 'nhập tay', 'N': 'chưa có file'}[c], 9.5, False, 'ink')]); yy += .26
        d.hline(s, .46, y + rh - .04, 12.87)
    return s
s = grpslide('Quy tắc: Kinh doanh', ['Doanh số', 'Doanh số sản phẩm mới', 'Mở mới', 'Viếng thăm khách hàng', 'Bán hàng hỗ trợ'], 'C21')
N(s, 'Năm nhóm KPI kinh doanh.', ['Doanh số đọc từ file Turnover của kế toán; mở mới, viếng thăm, bán hàng hỗ trợ đọc từ file CRM. Bấm tên file để mở đúng sheet.'], 'Nhóm Marketing.')
s = grpslide('Quy tắc: Marketing', ['Marketing hỗ trợ bán · Người tham dự', 'Sự kiện', 'Digital & Design · Ra mắt sản phẩm'], 'C35')
N(s, 'Nhóm Marketing.', ['Marketing hỗ trợ bán và người tham dự đọc từ Marketing Offline; Digital đọc từ file Digital Marketing; Dental Thái Lan chưa có nguồn.'], 'Back Office và Con người.')
s = grpslide('Quy tắc: Back Office · Con người', ['Đơn hàng & giao nhận', 'Tài chính & hành chính', 'Tài liệu · Tồn kho · Dự án', 'Đào tạo · Hành vi'], 'C23')
N(s, 'Back Office và Con người.', ['Back Office chấm điểm trong file theo dõi riêng, chưa nối vào file KPI nên đang nhập tay. Đào tạo và hành vi đã tự động.'], 'Xem KPI ở đâu.')

s = S('4 · KPI', 'Xem KPI: Report Hub', '4 · KPI', 'C24', '4.4', 'Phần 4 · Xem KPI ở đâu')
pair(s, 'rh/kpi_tab', 'rh/m_kpi', .46, 1.45, 8.4, 4.65, (.06, .0, 1, .9))
steps(s, 9.2, 1.5, 3.67, ['Chọn người (quản lý)', '4 ô: tháng · quý · năm · số KPI đạt', 'Xu hướng · đóng góp theo nhóm', 'Bảng chi tiết · "Chưa chốt"', 'Báo cáo KPI tháng · Setting Expectation'], gap=.88, sz=12.5)
N(s, 'Trang KPI của Report Hub.', ['Chọn người, đọc bốn ô tổng, biểu đồ, bảng chi tiết. Chưa có số thì hiện Chưa chốt, không tính là 0 phần trăm.'], 'Trên CRM.')

s = S('4 · KPI', 'Xem KPI: CRM', '4 · KPI', 'C07', '4.4', 'Phần 4 · Xem KPI ở đâu')
pair(s, 'crm/kpi_tab', 'crm/m_kpi', .46, 1.45, 8.4, 4.65)
bullets(s, 9.2, 1.5, 3.67, ['Trung bình nhóm, số người ≥ 100%, < 90%', 'KPI từng người như Report Hub', 'Lưu đơn, lượt đi, buổi thuyết trình ⇒ số KPI tự nhảy', 'Mở mới chỉ tính khi quản lý xác nhận', 'Nhóm Thái: tiếng Anh'], 12.5, 8)
N(s, 'KPI trên CRM.', ['Cùng cách xem như Report Hub, thêm phần tổng nhóm. Sales lưu dữ liệu thì sau vài giây số KPI tự nhảy.'], 'Total KPI và email KPI.')

s = S('4 · KPI', 'Total KPI · email KPI tháng', '4 · KPI', 'C30', '4.4', 'Phần 4 · Xem KPI ở đâu')
shot(s, 'rh/total_kpi', .46, 1.45, 6.8, 4.65, (0, 0, 1, .95))
phone(s, 'crm/m_kpi_report', 7.55, 1.45, 4.65)
d.text(s, 9.95, 1.5, 2.92, .4, [('Total KPI', 16, True, 'title')])
d.text(s, 9.95, 1.9, 2.92, 1.6, [('Admin · Giám đốc · Trưởng phòng. 4 mảng F · C · P · L, KPI công ty, thành viên.', 12, False, 'ink')])
d.text(s, 9.95, 3.6, 2.92, .4, [('Email KPI tháng', 16, True, 'title')])
d.text(s, 9.95, 4.0, 2.92, 1.8, [('Người nhận tự điền theo sơ đồ tổ chức, bảng KPI, nhận xét ⇒ xem trước ⇒ gửi.', 12, False, 'ink')])
N(s, 'Total KPI và email KPI tháng.', ['Total KPI tổng hợp toàn công ty cho quản lý. Email KPI tháng soạn sẵn, xem trước rồi gửi; số trong ảnh là giả lập.'], 'KPI theo phòng ban.')

def ppl_slide(title, groups, kind):
    s = S('4 · KPI', title, '4 · KPI', kind, '4.5', 'Phần 5 · KPI theo phòng ban')
    ppl = [(n, t, ks) for dept, sub_, P in KPEOPLE if dept in groups for n, t, ks in P]
    d.text(s, .46, 1.38, 12.4, .3, [('Trọng số quý 1 của từng người theo trạng thái kết nối của KPI', 12, True, 'title')])
    for i, (n, t, ks) in enumerate(ppl):
        y = 1.8 + i * .4; x = 3.9
        d.text(s, .46, y, 1.35, .36, [(n, 12.5, True, 'title')], anchor='m')
        d.text(s, 1.8, y, 2.05, .36, [(t[:30], 9.5, False, 'muted')], anchor='m')
        tot = sum(w for _, w in ks)
        for c in 'ASWMN':
            w_ = sum(w for k, w in ks if kstat(k) == c)
            if not w_: continue
            bw = 8.0 * w_ / tot; d.box(s, x, y + .05, bw, .27, STC[c], kind=MSO_SHAPE.RECTANGLE)
            if bw > .45: d.text(s, x, y + .05, bw, .27, [(f'{w_}%', 9.5, True, 'on_pri' if c in 'AM' else 'title')], 'c', 'm')
            x += bw
    lx = .46
    for c in 'ASWMN':
        d.box(s, lx, 6.25, .2, .2, STC[c], kind=MSO_SHAPE.RECTANGLE); d.text(s, lx + .27, 6.18, 2.3, .32, [(STN[c], 10.5, False, 'muted')], anchor='m'); lx += 2.48
    return s
s = ppl_slide('Theo phòng ban: KD & MKT VN', ['Kinh doanh & Marketing Việt Nam · Sales', 'Kinh doanh & Marketing Việt Nam · Quản lý & Marketing'], 'C10')
N(s, 'Kinh doanh và Marketing Việt Nam.', ['Mỗi thanh là 100 phần trăm trọng số của một người, tô theo trạng thái kết nối. Gần như toàn bộ đã tự động, trừ đi kèm coaching của Tuyen, Giang.'], 'Back Office và Thái Lan.')
s = ppl_slide('Theo phòng ban: BO & Thái Lan', ['Back Office', 'Thái Lan · Surgical & Dental'], 'C18')
N(s, 'Back Office và Thái Lan.', ['Back Office còn 30 tới 65 phần trăm trọng số nhập tay; Dental Thái Lan chưa có nguồn; sự kiện Surgical Thái Lan đã nối nhưng chưa nhận số.'], 'Việc cần chốt.')

s = S('4 · KPI', 'Việc cần chốt', '4 · KPI', 'C40', '4.6', 'Phần 6 · Việc cần chốt')
ISS = [('Số tháng tồn kho tính sai %', 'Mục tiêu theo tháng, thực đạt nhập điểm ⇒ đổi sang điểm, mục tiêu 100', 'Back Office'), ('Sự kiện Surgical Thái Lan chưa nhận số', 'Dao 40%, Miew 15% trọng số trống tháng 09 – 10', 'Product Team'),
       ('Dental Thái Lan chưa có nguồn', 'Sui 70%, Man 65%, Yong 50% đang gõ tay', 'Dao'), ('Mức trừ điểm xử lý đơn khác nhau', 'File Back Office trừ 1, 3 hoặc 5 điểm', 'Ms. Hoa'),
       ('Điểm báo cáo Back-office: số người khác nhau', 'Quy tắc 3 người, file KPI 4 người; VAT chưa phân công', 'Back Office'), ('Hạn ngày 2 hằng tháng', 'Hệ thống đã bỏ khoá từ 03/10 — giữ hay khoá lại', 'Giám đốc'),
       ('Còn thiếu quy tắc / mục tiêu', 'Tổng KH 3 NPP Thái, nội dung digital Thái, thang điểm thầu, form coaching', 'Người phụ trách'), ('File đang đọc mang tên FY67', 'Training Master, KBI — xác nhận dùng tiếp cho FY68', 'Product Team')]
for i, (t, ds, who) in enumerate(ISS):
    x = .46 + (i % 2) * 6.27; y = 1.45 + (i // 2) * 1.2
    d.box(s, x, y, .1, 1.05, '#C9473B'.lstrip('#'), kind=MSO_SHAPE.RECTANGLE)
    d.text(s, x + .25, y, 5.8, .35, [(t, 14, True, 'title')])
    d.text(s, x + .25, y + .38, 4.4, .65, [(ds, 12, False, 'ink')])
    d.chip(s, x + 4.75, y + .45, who, fill='soft', sz=10.5)
N(s, 'Tám điểm cần chốt, xếp theo ảnh hưởng tới điểm KPI.', ['Nặng nhất là số tháng tồn kho đang tính sai phần trăm, sự kiện Surgical Thái Lan chưa có số và Dental Thái Lan chưa có nguồn.'], 'Bước tiếp theo.')

s = S('4 · KPI', 'Bước tiếp theo', '4 · KPI', 'C45', '4.6', 'Phần 6 · Việc cần chốt')
NX = [('Product Team làm được ngay', ['Nối sự kiện Surgical Thái Lan vào file KPI', 'Nối Back Office KPI Tracking FY68 (các sheet đã có điểm)', 'Kiểm tra file Đào tạo / Hành vi đúng năm FY68']),
      ('Back Office chốt', ['Cách tính Số tháng tồn kho', 'Mức trừ điểm xử lý đơn, điểm báo cáo; danh sách người', 'Thêm KPI Phân bổ chi phí, phân công Hoá đơn VAT']),
      ('Giám đốc / Dao chốt', ['Cập nhật liên tục hay khoá số sau ngày 2', 'File nguồn FY68 cho Dental Thái Lan', 'Form coaching, tổng KH 3 NPP Thái, mục tiêu digital Thái'])]
for i, (t, its) in enumerate(NX):
    x = .46 + i * 4.17
    d.box(s, x, 1.45, 4.0, 4.7, 'hi' if i == 0 else 'panel', line='hi_line' if i == 0 else None)
    d.text(s, x + .25, 1.65, 3.5, .45, [(t, 17, True, 'title')])
    steps(s, x + .25, 2.35, 3.55, its, gap=1.1, sz=13)
N(s, 'Ai làm gì tiếp.', ['Product Team nối các nguồn đã có file; Back Office chốt cách chấm điểm; Giám đốc và Dao chốt quy định và file nguồn cho Thái Lan.'], 'Sang phần năm: Training Hub.')
# ═════════════════════════ PHẦN 5 · TRAINING HUB
s = divider('5', 'Training Hub', 'Học · kiểm tra · chứng chỉ', ['3 vai trò', 'Tự học', 'Làm bài, chứng chỉ', 'Trainer', 'Cửa chung cho mọi app'], 'th/home', '5.1', 'th/m_home')
N(s, 'Mở phần Training Hub.', ['Training Hub vừa là nơi đào tạo nội bộ, vừa là cửa đăng nhập chung của các app.'], 'Ba vai trò.')

s = S('5 · Training Hub', '3 vai trò', '5 · TRAINING HUB', 'C27', '5.1', 'Ai dùng, đăng nhập')
for i, (n, t, ds) in enumerate([('01', 'Học viên', 'Đọc tài liệu, làm bài, nhận chứng chỉ'), ('02', 'Trainer', 'Tạo buổi, tài liệu, thư mời, bài kiểm tra'), ('03', 'Admin', 'Người dùng, cấu hình, chế độ test')]):
    x = .46 + i * 4.17
    d.text(s, x, 1.45, 3.9, .8, [(n, 44, True, 'sec')])
    d.text(s, x, 2.3, 3.9, .45, [(t, 20, True, 'title')])
    d.hline(s, x, 2.85, x + 3.8)
    d.text(s, x, 3.0, 3.85, .8, [(ds, 14, False, 'ink')])
shot(s, 'th/login', .46, 4.05, 4.9, 2.3, (.0, .1, 1, .9), 'mid')
phone(s, 'th/m_more', 5.55, 3.9, 2.45)
note(s, 6.77, 4.05, 6.1, 2.3, 'Đăng nhập riêng:', 'email + mã PIN. Lần đầu kích hoạt bằng mã 6 số qua email. Sai PIN 8 lần khoá 10 phút.', fill='soft', line=None)
N(s, 'Ba vai trò và đăng nhập.', ['Học viên là mọi nhân sự. Trainer tạo buổi và bài kiểm tra. Admin quản trị.', 'Training Hub đăng nhập bằng email và mã PIN, lần đầu kích hoạt qua email.'], 'Trang đầu tiên của học viên.')

s = S('5 · Training Hub', 'Trang của tôi', '5 · TRAINING HUB', 'C44', '5.2', 'Trang của tôi')
pair(s, 'th/home_learner', 'th/m_home', .46, 1.45, 7.6, 4.6, (0, 0, 1, .9))
steps(s, 8.35, 1.5, 4.52, [('Việc cần làm', 'bài kiểm tra đang mở, theo hạn'), ('Hành trình học tập', 'Tài liệu → Bài kiểm tra → Chứng chỉ'), ('Sắp diễn ra', 'buổi hôm nay có nút Vào họp'), ('Học tiếp', 'tài liệu đang đọc dở')], gap=1.15)
N(s, 'Trang của tôi.', ['Mở Training Hub là thấy việc cần làm, hành trình học tập của từng buổi, các buổi sắp tới và tài liệu đang đọc dở.'], 'Học viên tự học thế nào?')

s = S('5 · Training Hub', 'Tự học', '5 · TRAINING HUB', 'C19', '5.5', 'Tự học: đọc tài liệu')
chev(s, ['Học tài liệu', 'Làm bài kiểm tra', 'Nhận chứng chỉ'], y=1.45, h=.65, sub=['Đọc tới trang cuối là xong', 'Mở khi học đủ tài liệu bắt buộc', 'PDF gửi email, CC quản lý'], subh=.6)
pair(s, 'th/reader', 'th/m_reader', .46, 2.95, 8.0, 3.2, (0, 0, 1, .85))
bullets(s, 8.75, 3.0, 4.12, ['Đọc ngay trong Hub, không cần quyền Drive', 'PowerPoint, Word tự chuyển PDF', 'Tiến độ đọc được lưu'], 13, 8)
N(s, 'Lộ trình tự học.', ['Ba bước: học tài liệu, làm bài, nhận chứng chỉ. Tài liệu đọc ngay trong Hub, đọc tới trang cuối là tính xong.'], 'Làm bài kiểm tra.')

s = S('5 · Training Hub', 'Làm bài, nhận chứng chỉ', '5 · TRAINING HUB', 'C28', '5.6', 'Làm bài kiểm tra')
for i, (sid, t, bx) in enumerate([('th/quiz_question', 'Làm bài · có đồng hồ', (.1, .05, 1, .7)), ('th/quiz_result', 'Kết quả · thang 100', (.15, .05, .85, .75))]):
    x = .46 + i * 4.17
    shot(s, sid, x, 1.45, 4.0, 3.4, bx, 'mid')
    d.text(s, x, 5.0, 4.0, .4, [(t, 16, True, 'title')])
phone(s, 'th/m_quiz', 9.4, 1.45, 3.4); d.text(s, 8.8, 5.0, 4.0, .4, [('Điện thoại', 16, True, 'title')], 'c')
d.text(s, .46, 5.55, 12.41, .8, [('Hết giờ tự nộp · mất mạng vẫn giữ đáp án · mỗi câu bằng điểm · đạt thì chứng chỉ tự gửi email.', 13.5, False, 'ink')])
N(s, 'Làm bài và chứng chỉ.', ['Bài kiểm tra có đồng hồ, hết giờ tự nộp, lỡ mất mạng vẫn giữ đáp án.', 'Đạt điểm thì chứng chỉ PDF tự tạo và gửi email, có CC quản lý.'], 'Phía trainer làm gì?')

s = S('5 · Training Hub', 'Trainer: tạo buổi', '5 · TRAINING HUB', 'C20', '5.8', 'Trainer: tạo buổi, tài liệu, thư mời')
chev(s, ['Tạo buổi', 'Tải tài liệu', 'Gửi thư mời', 'Tự nhắc'], y=1.45, h=.65, sub=['Thư mục Drive tự tạo', 'Kéo thả cả thư mục, ≤ 1 GB/file', 'Xem trước email thật', '8:30 ngày đào tạo, thứ Hai'], subh=.6)
shot(s, 'th/create_session', .46, 2.95, 5.0, 3.4, 'auto')
shot(s, 'th/invite', 5.65, 2.95, 5.0, 3.4, 'auto')
phone(s, 'th/m_session', 11.25, 2.95, 3.4)
N(s, 'Trainer tạo buổi.', ['Tạo buổi thì thư mục Drive tự tạo. Kéo thả tài liệu vào, gửi thư mời có xem trước. Hệ thống tự nhắc lúc 8 giờ 30 ngày đào tạo và mỗi thứ Hai.'], 'Soạn bài và xem kết quả.')

s = S('5 · Training Hub', 'Trainer: bài kiểm tra', '5 · TRAINING HUB', 'C10', '5.9', 'Trainer: soạn bài kiểm tra, xem kết quả')
shot(s, 'th/quiz_builder', .46, 1.45, 5.4, 4.6, 'auto')
phone(s, 'th/m_library', 6.1, 1.45, 4.6)
for i, (v, t) in enumerate([('Dán nhanh', 'câu hỏi từ Word / Excel'), ('Quy định', 'điểm đạt, số phút, số lượt, hạn'), ('Phân tích', 'từng câu; câu khó tô đỏ'), ('Nhắc', 'người chưa đạt; xuất CSV')]):
    d.kpi(s, 8.35, 1.45 + i * 1.25, 4.52, 1.1, v, t, vsz=22, dark=(i == 0))
N(s, 'Soạn bài và xem kết quả.', ['Trainer dán câu hỏi từ Word, đặt điểm đạt, thời gian, số lượt. Kết quả có phân tích từng câu và nút nhắc người chưa đạt.'], 'Cuối cùng: vai trò cửa chung.')

s = S('5 · Training Hub', 'Cửa chung cho mọi app', '5 · TRAINING HUB', 'C34', '5.12', 'Cửa chung cho các app khác')
d.table(s, .46, 1.45, [3.4, 4.6, 1.1, 1.1, 1.1, 1.11], [['Dịch vụ', 'Làm gì', 'Report Hub', 'CRM', 'Product Data', 'KPI'],
        ['Đăng nhập', 'Email + mã 6 số, 30 ngày', '●', '●', '●', '●'], ['Phân quyền', 'Bật / tắt quyền từng người', '●', '●', '●', '●'],
        ['Giao việc', 'Danh sách, trao đổi, hoàn thành', '●', '', '', ''], ['Email cập nhật', 'Báo bản mới tới mọi người', '●', '●', '', '']],
        row_h=.72, head=12, body=13.5, styles={1: 'focus'}, aligns=['l', 'l', 'c', 'c', 'c', 'c'])
note(s, .46, 5.35, 12.41, .95, 'Một tài khoản, mọi app.', 'Nghỉ việc hay đổi vị trí: Admin sửa một dòng, quyền ở mọi app đổi theo. Buổi Completed tự vào KPI đào tạo.', fill='soft', line=None)
N(s, 'Dịch vụ dùng chung.', ['Backend Training Hub giữ đăng nhập, phân quyền, giao việc và email cập nhật cho các app khác.', 'Vì vậy Admin chỉ cần sửa một nơi.'], 'Mình tổng kết.')

# ═════════════════════════ PHẦN 6 · CHẾ ĐỘ TEST
s = divider('6', 'Chế độ TEST', 'Admin thử mọi chức năng · thấy luồng dữ liệu', ['Ghi thật, tên TEST', 'Email chỉ về Admin', 'Đóng vai Giám đốc, Quản lý', 'Bảng Luồng dữ liệu', 'Xoá sạch sau khi thử'], 'crm/test_flow', '6.1', 'crm/m_test_flow')
N(s, 'Mở phần chế độ TEST.', ['Admin cần thử mọi chức năng trên dữ liệu thật mà không làm phiền ai. Chế độ TEST cho phép dùng app như một người tên TEST, lưu thật, và app vẽ ngay dữ liệu đi đâu.'], 'Ba nguyên tắc.')

s = S('6 · Chế độ TEST', 'Ba nguyên tắc', '6 · CHẾ ĐỘ TEST', 'C27', '6.2', 'Ba nguyên tắc')
for i, (ic, t, ds) in enumerate([('file', 'Ghi thật, tên TEST', 'Lưu vào đúng file Sheet / Drive đang dùng. PIC, tên file bắt đầu bằng TEST (ảnh: TEST_20261007_…_01.jpg).'),
                                 ('shield', 'Email chỉ về Admin', 'Báo cáo, xin duyệt, phê duyệt, KPI… chỉ gửi Admin, tiêu đề có [TEST], đầu thư ghi người nhận thật.'),
                                 ('eye', 'Không ai khác thấy', 'Ẩn với mọi người, không vào báo cáo của người thật, không tính KPI.')]):
    x = .46 + i * 4.17
    d.box(s, x, 1.45, 4.0, 3.3, 'panel')
    d.icon(s, ic, x + .3, 1.75, .6)
    d.text(s, x + .3, 2.55, 3.45, .45, [(t, 19, True, 'title')])
    d.text(s, x + .3, 3.1, 3.45, 1.5, [(ds, 13.5, False, 'ink')])
note(s, .46, 5.05, 12.41, 1.2, 'Chỉ Admin, chỉ hồ sơ TEST.', 'Máy chủ kiểm chữ ký phiên ở mọi lệnh. Đóng vai Giám đốc / Quản lý chỉ duyệt được hồ sơ của TEST; hồ sơ thật vẫn cần người thật.')
N(s, 'Ba nguyên tắc an toàn.', ['Ghi thật nhưng mọi thứ mang tên TEST. Email chỉ về Admin. Người khác không thấy và KPI không bị ảnh hưởng.', 'Máy chủ kiểm quyền Admin ở mọi lệnh, không chỉ trên giao diện.'], 'Bật thế nào?')

s = S('6 · Chế độ TEST', 'Bật chế độ TEST', '6 · CHẾ ĐỘ TEST', 'C44', '6.3', 'Bật chế độ TEST')
pair(s, 'crm/test_pick', 'crm/m_test_pick', .46, 1.45, 6.1, 4.0, 'auto')
pair(s, 'rh/test_pick', 'rh/m_test_pick', 6.77, 1.45, 6.1, 4.0, 'auto')
d.text(s, .46, 5.85, 6.1, .55, [[('CRM  ', 13.5, True, 'sec'), ('Chế độ TEST trên thanh trên · điện thoại: Thêm ▸ Tiện ích', 12.5, False, 'ink')]])
d.text(s, 6.77, 5.85, 6.1, .55, [[('Report Hub  ', 13.5, True, 'sec'), ('Bấm tên ▸ Chế độ TEST · luồng dữ liệu · Ctrl + K "test"', 12.5, False, 'ink')]])
N(s, 'Lối vào.', ['CRM: nút Chế độ TEST trên thanh trên cùng, chọn nhóm và vai. Report Hub: bấm tên mình, chọn phòng và vai.', 'Bấm Bắt đầu, app tải lại với dải vàng đen trên đỉnh và nút TEST vàng ở góc trái dưới.'], 'Bảng Luồng dữ liệu.')

s = S('6 · Chế độ TEST', 'Bảng Luồng dữ liệu', '6 · CHẾ ĐỘ TEST', 'C25', '6.4', 'Bảng Luồng dữ liệu')
pair(s, 'crm/test_flow', 'crm/m_test_flow', .46, 1.45, 7.4, 4.75)
for i, (t, ds) in enumerate([('① Thao tác', 'app, tên thao tác, vai đang dùng'), ('② Backend', 'Apps Script nhận lệnh, thời gian chạy'), ('③ Lưu ở', 'file ▸ tab ▸ dòng, ảnh / PDF trên Drive'), ('④ Tiếp theo', 'email lẽ ra gửi ai, dữ liệu đi tiếp đâu')]):
    y = 1.5 + i * 1.17
    d.text(s, 8.2, y, 4.67, .4, [(t, 16, True, 'title')])
    d.text(s, 8.2, y + .4, 4.67, .6, [(ds, 13, False, 'muted')])
    if i < 3: d.hline(s, 8.2, y + 1.02, 12.87)
N(s, 'Đọc một thẻ luồng.', ['Mỗi lần lưu là một thẻ: thao tác gì, backend nào nhận, lưu ở file nào tab nào dòng nào, email lẽ ra gửi ai.', 'Nhật ký dùng chung cho CRM và Report Hub nên thấy cả luồng đi qua hai app.'], 'Kịch bản kiểm thử.')

s = S('6 · Chế độ TEST', 'Kịch bản CRM', '6 · CHẾ ĐỘ TEST', 'C34', '6.5', 'Kịch bản kiểm thử CRM')
dtable(s, .46, 1.42, [.45, 2.6, 4.6, 4.76], [['#', 'Thử', 'Lưu ở', 'Kiểm tra'],
       ['1', 'Đi địa bàn + chụp ảnh', '6. WEEKLY REPORT · ảnh TEST_….jpg trên Drive', 'Thẻ lịch, mở link ảnh'],
       ['2', 'Khách hàng · đơn hàng', '1. CUSTOMER · 3. SALE ORDER', 'KPI không cộng'],
       ['3', 'Đề xuất công tác', 'Business Trip + tab Business Trip Log', 'Email xin duyệt về Admin'],
       ['4', 'Duyệt (vai Giám đốc)', 'Cột Approval, folder công tác', 'Email [Approved] về Admin'],
       ['5', 'Báo cáo tuần / tháng', 'PDF trong folder báo cáo', 'Email về Admin'],
       ['6', 'KPI tháng · tự đánh giá quý', 'Không ghi Sheet · file Word', 'Email KPI về Admin'],
       ['7', 'Mở mới → Quản lý xác nhận', '16. NEW ACCOUNT & SKU, chứng từ TEST_…', 'KPI mở mới không cộng'],
       ['8', 'Công việc khác', 'File CRM nhóm ▸ ALL TASK', 'Thẻ trên tab']], row_h=.56, head=12, body=12)
N(s, 'Kịch bản CRM.', ['Bảng kịch bản tự đánh dấu chức năng đã thử thành công: đi địa bàn, chụp ảnh, khách hàng, đơn hàng, công tác, báo cáo, KPI, mở mới.', 'Mỗi dòng ghi rõ dữ liệu lưu ở đâu và kiểm tra cái gì.'], 'Report Hub.')

s = S('6 · Chế độ TEST', 'Kịch bản Report Hub', '6 · CHẾ ĐỘ TEST', 'C29', '6.6', 'Kịch bản kiểm thử Report Hub')
pair(s, 'rh/test_flow', 'rh/m_test_flow', .46, 1.45, 7.4, 4.75)
steps(s, 8.2, 1.5, 4.67, [('Key task, tiến độ', 'giao người khác ⇒ tự đổi thành TEST'), ('Công tác', 'đề xuất → duyệt → báo cáo'), ('Báo cáo', 'tuần, tháng, KPI, biên bản họp'), ('Giao việc', 'vai Quản lý ⇒ RH_Assign')], gap=1.15)
N(s, 'Kịch bản Report Hub.', ['Thêm key task, cập nhật tiến độ, đề xuất và duyệt công tác, gửi báo cáo, giao việc.', 'Giao cho người khác thì người nhận tự đổi thành TEST, thẻ luồng ghi rõ.'], 'Đóng vai Giám đốc.')

s = S('6 · Chế độ TEST', 'Đóng vai Giám đốc', '6 · CHẾ ĐỘ TEST', 'C10', '6.7', 'Đóng vai Giám đốc và Quản lý')
pair(s, 'crm/test_tfa', 'crm/m_test_tfa', .46, 1.45, 7.6, 4.75, 'auto')
for i, (t, ds) in enumerate([('Giám đốc', 'Duyệt công tác: chỉ thẻ TEST chọn được; email phê duyệt về Admin'), ('Quản lý', 'Mở mới: Xác nhận · Trả lại · Từ chối case TEST'), ('Đổi vai', 'ngay trong bảng Luồng dữ liệu')]):
    d.kpi(s, 8.4, 1.45 + i * 1.6, 4.47, 1.4, t, ds, vsz=20, dark=(i == 0))
N(s, 'Đóng vai người duyệt.', ['Đề xuất công tác ở vai Sales TEST, rồi đổi sang vai Giám đốc để duyệt luôn. Đề xuất thật bị khoá ô chọn.', 'Tương tự vai Quản lý xác nhận mở mới.'], 'Xoá dữ liệu sau khi thử.')

s = S('6 · Chế độ TEST', 'Xoá dữ liệu test', '6 · CHẾ ĐỘ TEST', 'C19', '6.8', 'Xoá dữ liệu test')
chev(s, ['Bảng Luồng dữ liệu', 'Xoá dữ liệu test', 'Máy chủ dọn', 'Thẻ kết quả'], y=1.45, h=.65,
     sub=['Nút vàng TEST góc trái dưới', 'Xác nhận một lần', 'Từng backend: file nhóm, Business Trip, Drive', 'Bao nhiêu dòng, bao nhiêu file'], subh=.7)
for i, (t, ds) in enumerate([('Xoá nội dung dòng TEST', 'giữ công thức, không xoá dòng ⇒ số dòng người khác không đổi'), ('Business Trip', 'dòng TEST + folder công tác vào thùng rác'), ('File Drive TEST_…', 'ảnh, chứng từ, báo cáo vào thùng rác')]):
    x = .46 + i * 4.17
    d.box(s, x, 3.3, 4.0, 1.9, 'panel')
    d.text(s, x + .25, 3.45, 3.5, .45, [(t, 16, True, 'title')])
    d.text(s, x + .25, 3.95, 3.5, 1.1, [(ds, 13, False, 'ink')])
note(s, .46, 5.45, 12.41, .85, 'Dữ liệu thật không bị đụng tới.', 'Chỉ dòng có PIC "TEST" và file tên TEST_… bị dọn.', fill='soft', line=None)
N(s, 'Dọn dẹp.', ['Một nút Xoá dữ liệu test: máy chủ xoá nội dung dòng TEST, giữ công thức, chuyển file TEST vào thùng rác.', 'Không xoá dòng nên số dòng của người khác không đổi.'], 'Training Hub có chế độ riêng.')

s = S('6 · Chế độ TEST', 'Training Hub & kỹ thuật', '6 · CHẾ ĐỘ TEST', 'C41', '6.9', 'Training Hub')
shot(s, 'th/test_mode', .46, 1.45, 5.6, 3.6, 'auto')
d.text(s, .46, 5.2, 5.6, 1.1, [('Training Hub: Admin ▸ menu tài khoản ▸ Bật Chế độ Test · đổi vai Trainer / Học viên · kịch bản 7 bước · buổi test ở sheet HUB_TestSessions.', 13, False, 'ink')])
dtable(s, 6.35, 1.45, [2.1, 4.42], [['Thành phần', 'Việc'],
       ['testflow.js', 'CRM + Report Hub: gắn tm=1 vào mọi lệnh, vẽ bảng luồng'],
       ['MMH_TestMode.gs', '9 backend: kiểm Admin, email về Admin, ẩn dữ liệu TEST, xoá'],
       ['MMH_TripFlow.gs', 'Vai Giám đốc chỉ duyệt đề xuất của TEST']], row_h=.75, head=12, body=12.5)
N(s, 'Training Hub và phần kỹ thuật.', ['Training Hub có chế độ test riêng từ trước, 7 bước từ tạo buổi tới chứng chỉ.', 'Phía kỹ thuật: một file testflow.js giống nhau ở CRM và Report Hub, một file MMH_TestMode.gs ở chín backend.'], 'Tổng kết.')

# ═════════════════════════ TỔNG KẾT
s = S('Tổng kết', '6 điều cần nhớ', 'TỔNG KẾT', 'C45', '1.1', 'Sách hướng dẫn đầy đủ')
SUM = [('1 hệ thống', 'Các app dùng chung Google Sheets của công ty'), ('1 lần đăng nhập', 'Email + mã 6 số, 30 ngày'), ('Nhập 1 lần', 'Báo cáo, lịch, KPI tự tạo'),
       ('KPI tự điền', 'Nguồn đã kết nối tự đổ số; người chỉ đặt Target'), ('4 màu KPI', 'Xám · xanh dương · vàng · xanh lá'), ('Chế độ TEST', 'Admin thử thật, xem luồng, xoá sạch')]
for i, (t, ds) in enumerate(SUM):
    x = .46 + (i % 3) * 4.17; y = 1.45 + (i // 3) * 2.5
    d.box(s, x, y, 4.0, 2.3, 'hi' if i == 3 else 'panel', line='hi_line' if i == 3 else None)
    d.badge(s, x + .25, y + .25, i + 1)
    d.text(s, x + .25, y + .8, 3.55, .45, [(t, 19, True, 'title')])
    d.text(s, x + .25, y + 1.3, 3.55, .9, [(ds, 14, False, 'ink')])
N(s, 'Chốt 6 ý.', ['Một hệ thống, một lần đăng nhập, nhập một lần, KPI tự điền, bốn màu KPI, và chế độ TEST để Admin kiểm thử an toàn.'], 'Tra cứu chi tiết ở đâu?')

s = S('Tổng kết', 'Tra cứu chi tiết: IFU', 'TỔNG KẾT', 'C41', '1.1', 'Sách hướng dẫn đầy đủ')
F = [('MMH_IFU_He_thong_du_lieu.pdf', 'Sách đầy đủ · 163 trang', '1'), ('MMH_IFU_01_CRM.pdf', 'CRM · 41 trang', '2'), ('MMH_IFU_02_Report_Hub.pdf', 'Report Hub · 41 trang', '3'),
     ('MMH_IFU_03_KPI.pdf', 'KPI · 29 trang', '4'), ('MMH_IFU_04_Training_Hub.pdf', 'Training Hub · 33 trang', '5'), ('MMH_IFU_05_Che_do_TEST.pdf', 'Chế độ TEST · 14 trang', '6')]
for i, (fn, ds, part) in enumerate(F):
    x = .46 + (i % 2) * 6.27; y = 1.45 + (i // 2) * 1.3
    d.box(s, x, y, 6.12, 1.12, 'panel')
    d.icon(s, 'book', x + .2, y + .28, .55)
    tb = d.text(s, x + .95, y + .15, 5.0, .4, [(fn, 15, True, 'title')])
    tb.text_frame.paragraphs[0].runs[0].hyperlink.address = IFU_FILE[part][1]
    d.text(s, x + .95, y + .58, 5.0, .4, [(ds, 13, False, 'muted')])
note(s, .46, 5.4, 12.41, .75, 'Bấm tên file để mở.', 'Sách dựng lại mỗi khi hệ thống thay đổi. Tài liệu nội bộ: có link tới dữ liệu thật, không chia sẻ ra ngoài.', fill='soft', line=None)
N(s, 'Chỉ chỗ tra cứu.', ['IFU có sáu file: sách đầy đủ và năm phần riêng, trong đó có phần chế độ TEST. Bấm tên file để mở.'], 'Cảm ơn mọi người.')

footer(d.prs.slides[2], '1.1', 'Sách hướng dẫn sử dụng đầy đủ', y=4.6, w=9.6, credit=False)
d.notes(d.prs.slides[2], 'Kết thúc, mời hỏi đáp.', ['Cảm ơn mọi người. Cần làm theo từng bước thì mở link IFU ở cuối mỗi slide, hoặc bấm nút Hướng dẫn trong app.', 'Mình mời câu hỏi.'])

from collections import Counter
bad = [f'{a[2]}/{b[2]} liền kề cùng {a[1]}' for a, b in zip(LOG, LOG[1:]) if a[1] == b[1]]
bad += [f'{k} dùng {n} lần' for k, n in Counter(p for _, p, _ in LOG).items() if n > 3 and k != 'C04']
print('pattern issues:', bad)
print('long titles:', [t for _, _, t in LOG if len(t) > 30])
print('missing notes:', [i for i, x in enumerate(d.prs.slides, 1) if i != 2 and not x.notes_slide.notes_text_frame.text.strip()])
out = d.save('/home/claude/outputs/20261013_MMH_He_thong_du_lieu_noi_bo_v3.pptx'); print(out, len(d.prs.slides))
