"""Shared MMH brand styles + small helpers for openpyxl."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter as L

FONT = 'Aptos'
NAVY, BLUE, BLUE_MID, LIGHT, PALE = '003047', '3A5CAA', '7D97CE', 'CFE2F3', 'F0F7FF'
YELLOW, CREAM, INK, MUTED, LINE = 'FFE100', 'FFF8E1', '1A1A1A', '666666', 'D9D9D9'
TAB = '1155CC'
RED, ORANGE, GREEN = 'C62828', 'B45309', '2E7D32'
RED_BG, ORANGE_BG, GREEN_BG = 'FDECEA', 'FFF3E0', 'E8F5E9'
SECTION = 'F5F5F5'

thin = Side(style='thin', color='BFBFBF')
hair = Side(style='thin', color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BORDER_LIGHT = Border(left=hair, right=hair, top=hair, bottom=hair)
YB = Side(style='medium', color=YELLOW)

Q_FMT = '#,##0;-#,##0;\\-'          # quantity, zero shows "-"
Q_IN_FMT = '#,##0'                   # input quantity, zero shows 0
TB_FMT = '[=0]\\-;[<10]0.0;#,##0'    # average: 1 decimal when small
M_FMT = '0'                          # yyyymm
MON_FMT = '0.0;-0.0;\\-'             # months of cover
PCT_FMT = '0%'
VND_FMT = '#,##0;-#,##0;\\-'


def f(size=11, bold=False, color=INK, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(c):
    return PatternFill('solid', fgColor=c)


def al(h='left', v='center', wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


def title(ws, text, cell='B1', size=16):
    c = ws[cell]
    c.value = text
    c.font = f(size, True, '000000')
    c.alignment = al('left', 'center')
    ws.row_dimensions[c.row].height = 26


def header_row(ws, row, col1, labels, height=36, fill_color=LIGHT, size=11):
    for i, lab in enumerate(labels):
        c = ws.cell(row, col1 + i, lab)
        c.font = f(size, True, '000000')
        c.fill = fill(fill_color)
        c.alignment = al('center', 'center', True)
        c.border = BORDER
    ws.row_dimensions[row].height = height


def put(ws, cell, value, font=None, fill_color=None, align=None, fmt=None, border=None):
    c = ws[cell] if isinstance(cell, str) else ws.cell(*cell)
    c.value = value
    if font is not None:
        c.font = font
    if fill_color:
        c.fill = fill(fill_color)
    if align is not None:
        c.alignment = align
    if fmt:
        c.number_format = fmt
    if border is not None:
        c.border = border
    return c


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def cf_fill(ws, rng, formula, bg=None, color=None, bold=False, stop=True):
    font = Font(name=FONT, color=color, bold=bold) if (color or bold) else None
    fll = PatternFill('solid', bgColor=bg, fgColor=bg) if bg else None
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], font=font, fill=fll, stopIfTrue=stop))


def band_cf(ws, rng, key_col_letter, first_row):
    """Pale-blue alternate rows (brand) only where the row has data."""
    cf_fill(ws, rng, f'AND(${key_col_letter}{first_row}<>"",MOD(ROW(),2)=0)', bg=PALE, stop=False)


def light_cf(ws, rng, first_cell_rel):
    """Traffic-light text cells: XANH / VÀNG / ĐỎ / 6T CHƯA LẤY / ĐẠT / CAO / QUÁ CAO / CẦN XÁC NHẬN."""
    c = first_cell_rel
    rules = [
        (f'{c}="6T CHƯA LẤY"', RED, 'FFFFFF'),
        (f'OR({c}="ĐỎ",{c}="QUÁ CAO",{c}="THIẾU")', RED_BG, RED),
        (f'OR({c}="VÀNG",{c}="CAO",LEFT({c},2)="DƯ")', ORANGE_BG, ORANGE),
        (f'OR({c}="XANH",{c}="ĐẠT",{c}="KHỚP")', GREEN_BG, GREEN),
        (f'OR({c}="CẦN XÁC NHẬN",{c}="CHƯA CÓ DỰ TRÙ")', PALE, BLUE),
    ]
    for formula, bg, color in rules:
        cf_fill(ws, rng, formula, bg=bg, color=color, bold=True)


def protect(ws, password=None):
    p = ws.protection
    p.sheet = True
    if password:
        p.password = password
    p.selectLockedCells = False
    p.selectUnlockedCells = False
    p.formatColumns = False
    p.formatRows = False
    p.autoFilter = False
    p.sort = True
    p.insertRows = True
    p.deleteRows = True


def unlock(cell):
    cell.protection = Protection(locked=False)


def lbl(ref):
    """Excel formula text giving 'mm/yyyy' for a yyyymm number reference."""
    return f'IF(MOD({ref},100)<10,"0","")&MOD({ref},100)&"/"&INT({ref}/100)'


def ym_add_formula(ref, k):
    return (f'YEAR(EDATE(DATE(INT({ref}/100),MOD({ref},100),1),{k}))*100'
            f'+MONTH(EDATE(DATE(INT({ref}/100),MOD({ref},100),1),{k}))')


def ym_add(ym, k):
    y, m = divmod(int(ym), 100)
    idx = y * 12 + (m - 1) + k
    return (idx // 12) * 100 + idx % 12 + 1


def ym_lbl(ym):
    return f'{int(ym) % 100:02d}/{int(ym) // 100}'
