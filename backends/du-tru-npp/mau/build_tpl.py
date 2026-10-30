"""Excel template for the NPP forecast file (same layout as the approved demo, no data).
The Apps Script fills it: {{...}} placeholders in shared strings, the hidden DỮ LIỆU sheet, and the end row of DS_MA.
Usage: python3 -I build_tpl.py <out.xlsx>
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.styles import Border
from sty import *

OUT = sys.argv[1]
NHIST = 10            # months in the history panel (the last one is the forecast month T)
R0 = 11               # first data row on DỮ LIỆU
MAXR = 3000           # formula ranges cover R0 .. R0 + MAXR - 1 (catalogue size limit)
NR = 500              # input rows on DỰ TRÙ
SIG = 'MMH_DUTRU_NPP_V1'

wb = Workbook()
wb.remove(wb.active)
S_HD, S_P, S_D = [wb.create_sheet(n) for n in ['HƯỚNG DẪN', 'DỰ TRÙ', 'DỮ LIỆU']]
for ws in (S_HD, S_P):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 2.5
S_P.sheet_properties.tabColor = TAB
F_, F_IN = f(10), f(11, True, TAB)

# ------------------------------------------------------------------ DỮ LIỆU (veryHidden; the script rewrites the whole sheet)
D = S_D
D['A1'], D['B1'] = 'Ký hiệu', SIG
D.cell(R0 - 1, 4, 'Nhãn chọn')
D.cell(R0, 4, '-')
DL = lambda j: f"'DỮ LIỆU'!${CL(4 + j)}${R0}:${CL(4 + j)}${R0 + MAXR - 1}"
wb.defined_names['DS_MA'] = DefinedName('DS_MA', attr_text=f"'DỮ LIỆU'!$D${R0}:$D${R0}")
D.sheet_state = 'veryHidden'

# ------------------------------------------------------------------ DỰ TRÙ
ws = S_P
ws['B1'] = 'DỰ TRÙ THÁNG {{T}} – {{NPP}}'
ws['B1'].font = f(16, True, '000000'); ws.row_dimensions[1].height = 26
ws['B2'] = 'Tháng gửi: {{CM}}   |   Dự trù cho tháng {{T}}   |   Hạn gửi MMH: {{HAN}}'
ws['B2'].font = f(12, True, BLUE)
ws['B3'] = '{{B3}}'
ws['B3'].font = f(10, False, MUTED)
ws['B4'] = 'Nếu bảng chưa hiển thị số liệu, vui lòng bấm "Enable Editing" (Cho phép chỉnh sửa) trên thanh thông báo màu vàng.'
ws['B4'].font = f(10, True, RED)

P0 = 22
LAST = P0 + NR - 1
rng = lambda c: f'${c}${P0}:${c}${LAST}'

ws.merge_cells('B5:M5')
n_in = f'COUNT({rng("L")})'
n_hi = f'SUMPRODUCT(--(ISNUMBER({rng("L")})),--(ISNUMBER({rng("J")})),--({rng("L")}>{rng("J")}))'
n_cao = f'SUMPRODUCT(--({rng("K")}="Cao"),--({rng("L")}>0))'
ws['B5'] = (f'=IF({n_in}=0,"Tình trạng: chưa nhập số dự trù. Chọn mã ở cột B và nhập số lượng ở cột L.",'
            f'"Tình trạng: đã nhập "&{n_in}&" mã, tổng "&TEXT(SUM({rng("L")}),"#,##0")&". Cao hơn đề xuất: "&{n_hi}&" mã. '
            f'Mã có cảnh báo Cao vẫn dự trù: "&{n_cao}&" mã.")')
ws['B5'].font = f(12, True, NAVY); ws['B5'].alignment = al('left', 'center', indent=1); ws.row_dimensions[5].height = 26
cf_fill(ws, 'B5:M5', f'{n_in}=0', bg=ORANGE_BG, color=ORANGE, bold=True)
cf_fill(ws, 'B5:M5', f'{n_in}>0', bg=GREEN_BG, color=GREEN, bold=True)

# ---- panel: one code (rows 7-18)
ws['B7'] = 'XEM CHI TIẾT MỘT MÃ'; ws['B7'].font = f(12, True, NAVY)
ws['B8'] = 'Mã cần xem:'; ws['B8'].font = f(11, True)
ws.merge_cells('C8:F8')
c = ws['C8']; c.fill = fill(CREAM); c.font = F_IN; c.border = BORDER; unlock(c)
for cc in 'DEF':
    ws[f'{cc}8'].fill = fill(CREAM); ws[f'{cc}8'].border = BORDER
ws['G8'] = '▾ chọn mã; để trống thì hiển thị mã vừa chọn gần nhất trong danh sách bên dưới.'
ws['G8'].font = f(9, False, MUTED, True)
dvp = DataValidation(type='list', formula1='=DS_MA', allow_blank=True)
dvp.error, dvp.errorTitle, dvp.showErrorMessage = 'Vui lòng chọn mã trong danh sách.', 'Mã hàng', True
ws.add_data_validation(dvp); dvp.add('C8')
ws['Q8'] = f'=IF($C$8<>"",$C$8,IFERROR(LOOKUP(2,1/({rng("B")}<>""),{rng("B")}),""))'
ws['Q9'] = f'=IFERROR(MATCH($Q$8,DS_MA,0),"")'
X = lambda j: f'IF($Q$9="","",INDEX({DL(j)},$Q$9))'
ws['B9'] = f'=IF($Q$9="","Chưa chọn mã.",{X(0)}&"   ·   "&{X(3)})'
ws['B9'].font = f(12, True, BLUE)
kn = [('B', 'Số lượng đặt TB/tháng', 4, '#,##0'), ('D', 'Dự trù đã gửi {{K0}}–{{LASTP}}', 5, Q_FMT), ('F', 'Đã đặt hàng {{K0}}–{{LASTP}}', 6, Q_FMT),
      ('H', 'Dự trù chưa đặt hàng đến hết {{LASTP}}', 7, Q_FMT), ('J', 'Dự trù đề xuất {{T}}', 12, '#,##0'), ('L', 'Cảnh báo', 9, '@')]
for col_, lb, j, fm in kn:
    c2 = CL(CI(col_) + (0 if col_ == 'L' else 1))
    ws[f'{col_}10'] = lb; ws[f'{col_}10'].font = f(9, True, BLUE); ws[f'{col_}10'].alignment = al('center', 'center', True)
    ws[f'{col_}11'] = f'={X(j)}'; ws[f'{col_}11'].font = f(14, True, NAVY); ws[f'{col_}11'].alignment = al('center')
    ws[f'{col_}11'].number_format = fm
    for r in (10, 11):
        for cc in range(CI(col_), CI(c2) + 1):
            ws.cell(r, cc).fill = fill(LIGHT if r == 10 else PALE)
        if c2 != col_:
            ws.merge_cells(f'{col_}{r}:{c2}{r}')
ws.row_dimensions[10].height = 26; ws.row_dimensions[11].height = 26
cf_fill(ws, 'L11', 'L11="Cao"', bg=RED_BG, color=RED, bold=True)
cf_fill(ws, 'L11', 'L11="Theo dõi"', bg=ORANGE_BG, color=ORANGE, bold=True)
cf_fill(ws, 'L11', 'L11="Ổn"', bg=GREEN_BG, color=GREEN, bold=True)
ws.merge_cells('B12:M12')
ws['B12'] = f'={X(13)}'
ws['B12'].font = f(10, True, INK); ws['B12'].alignment = al('left', 'center', True); ws.row_dimensions[12].height = 30
ws['B14'] = 'Tháng'
ws['B15'] = 'Dự trù đã gửi'
ws['B16'] = 'Đặt hàng (đã đặt / dự kiến)'
ws['B17'] = 'Dự trù chưa đặt hàng (cộng dồn)'
for r in range(14, 18):
    ws[f'B{r}'].font = f(10, True, NAVY if r == 14 else INK)
    ws[f'B{r}'].fill = fill(LIGHT if r == 14 else 'FFFFFF'); ws[f'B{r}'].border = BORDER_LIGHT
for k in range(NHIST):
    colk = CL(3 + k)
    ws[f'{colk}14'] = '{{HM%d}}' % (k + 1)
    ws[f'{colk}14'].font = f(9, True, NAVY); ws[f'{colk}14'].alignment = al('center', 'center', True)
    ws[f'{colk}14'].fill = fill(CREAM if k == NHIST - 1 else LIGHT); ws[f'{colk}14'].border = BORDER_LIGHT
    for j, r in enumerate((15, 16, 17)):
        cc = ws[f'{colk}{r}']
        cc.value = f'={X(15 + 3 * k + j)}'
        cc.number_format = Q_FMT; cc.alignment = al('center'); cc.font = f(10, r == 17); cc.border = BORDER_LIGHT
ws.row_dimensions[14].height = 28
ws['B18'] = 'Kỳ mới bắt đầu từ tháng 09: phần dự trù chưa đặt hàng được tính lại từ 0. Tháng đặt nhiều hơn dự trù không được trừ vào phần chưa đặt của các tháng trước.'
ws['B18'].font = f(9, False, MUTED, True)
cf_fill(ws, f'C17:{CL(2 + NHIST)}17', 'N(C17)>0', color=RED, bold=True)

# ---- list
ws['B20'] = 'DANH SÁCH DỰ TRÙ THÁNG {{T}}'; ws['B20'].font = f(12, True, NAVY)
ws['F20'] = 'Chọn mã ở cột B (bấm vào ô rồi chọn trong danh sách), thông tin của mã hiện ra ngay; sau đó nhập số lượng ở cột L. Mỗi mã một dòng.'
ws['F20'].font = f(9, False, MUTED, True)
heads = ['Mã hàng (chọn trong danh sách)', 'Nhóm hàng', 'Số lượng đặt TB/tháng (6 tháng)', 'Dự trù đã gửi {{K0}}–{{LASTP}}',
         'Đã đặt hàng {{K0}}–{{LASTP}}', 'Dự trù chưa đặt hàng đến hết {{LASTP}}', 'Dự trù đã gửi cho {{EDITR}}',
         'Dự trù chưa đặt hàng dự kiến đến hết {{TM1}}', 'Dự trù đề xuất {{T}}', 'Cảnh báo', 'SỐ DỰ TRÙ {{T}} (nhập)', 'Thông tin cần lưu ý']
header_row(ws, 21, 2, heads, 58)
for cc in ('B', 'L'):
    ws[f'{cc}21'].fill = fill(CREAM)
    ws[f'{cc}21'].border = Border(left=thin, right=thin, top=thin, bottom=YB)
for i in range(NR):
    r = P0 + i
    cb = ws[f'B{r}']
    cb.font = F_IN; cb.fill = fill(CREAM); unlock(cb)
    ws[f'P{r}'] = f'=IF($B{r}="","",IFERROR(MATCH($B{r},DS_MA,0),"x"))'
    Y = lambda j: f'IF(OR($P{r}="",$P{r}="x"),"",INDEX({DL(j)},$P{r}))'
    ws[f'C{r}'] = f'={Y(3)}'
    for colx, j in (('D', 4), ('E', 5), ('F', 6), ('G', 7), ('H', 10), ('I', 11), ('J', 12), ('K', 9)):
        ws[f'{colx}{r}'] = f'={Y(j)}'
    ws[f'Q{r}'] = f'=IF($B{r}="",0,COUNTIF({rng("B")},$B{r}))'
    L_, J_ = f'$L{r}', f'$J{r}'
    ws[f'M{r}'] = (f'=IF($B{r}="",IF({L_}="","","Chưa chọn mã hàng."),IF($P{r}="x","Mã không có trong danh mục – vui lòng chọn lại.",'
                   f'IF($Q{r}>1,"Mã bị chọn 2 lần – vui lòng giữ 1 dòng. ","")&'
                   f'IF(AND(ISNUMBER({L_}),ISNUMBER({J_})),IF({L_}>{J_},"Cao hơn đề xuất "&TEXT({L_}-{J_},"#,##0")&": phần dư thành dự trù chưa đặt hàng. ",'
                   f'IF({L_}<{J_},"Thấp hơn đề xuất "&TEXT({J_}-{L_},"#,##0")&". ","")),"")&{Y(14)}))')
    cl_ = ws[f'L{r}']; cl_.font = F_IN; cl_.fill = fill(CREAM); unlock(cl_); cl_.number_format = Q_IN_FMT
    for colx in 'DEFGHIJ':
        ws[f'{colx}{r}'].number_format = '#,##0' if colx in 'DJ' else Q_FMT
        ws[f'{colx}{r}'].alignment = al('center')
    ws[f'K{r}'].alignment = al('center'); ws[f'L{r}'].alignment = al('center')
    ws[f'M{r}'].alignment = al('left', 'center', False)
    for colx in 'CDEFGHIJKM':
        ws[f'{colx}{r}'].font = F_
dvb = DataValidation(type='list', formula1='=DS_MA', allow_blank=True)
dvb.error, dvb.errorTitle, dvb.showErrorMessage = 'Vui lòng chọn mã trong danh sách.', 'Mã hàng', True
ws.add_data_validation(dvb); dvb.add(f'B{P0}:B{LAST}')
dvl = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', allow_blank=True)
dvl.error, dvl.errorTitle, dvl.showErrorMessage = 'Chỉ nhập số nguyên từ 0 trở lên.', 'Số lượng', True
ws.add_data_validation(dvl); dvl.add(f'L{P0}:L{LAST}')
K1 = f'$K{P0}'
cf_fill(ws, f'K{P0}:K{LAST}', f'{K1}="Cao"', bg=RED_BG, color=RED, bold=True)
cf_fill(ws, f'K{P0}:K{LAST}', f'{K1}="Theo dõi"', bg=ORANGE_BG, color=ORANGE, bold=True)
cf_fill(ws, f'K{P0}:K{LAST}', f'{K1}="Ổn"', bg=GREEN_BG, color=GREEN, bold=True)
cf_fill(ws, f'K{P0}:K{LAST}', f'OR({K1}="Chưa dự trù",{K1}="Chưa có lịch sử")', bg=PALE, color=BLUE, bold=True)
M1 = f'$M{P0}'
cf_fill(ws, f'M{P0}:M{LAST}', f'OR(ISNUMBER(SEARCH("cao hơn đề xuất",{M1})),ISNUMBER(SEARCH("2 lần",{M1})),'
        f'ISNUMBER(SEARCH("không có trong",{M1})),ISNUMBER(SEARCH("Chưa chọn",{M1})))', color=RED, stop=False)
cf_fill(ws, f'L{P0}:L{LAST}', 'TRUE', bg=CREAM, stop=False)
cf_fill(ws, f'C{P0}:K{LAST}', f'AND($B{P0}<>"",MOD(ROW(),2)=0)', bg=PALE, stop=False)
widths(ws, {'B': 38, 'C': 13, 'D': 11, 'E': 11, 'F': 11, 'G': 12, 'H': 12, 'I': 12, 'J': 11, 'K': 14, 'L': 13, 'M': 118, 'N': 4})
for cc in 'OPQ':
    ws.column_dimensions[cc].hidden = True
ws.freeze_panes = f'C{P0}'
ws.auto_filter.ref = f'B21:M{LAST}'
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.sheet_view.zoomScale = 90
protect(ws)

# ------------------------------------------------------------------ HƯỚNG DẪN
ws = S_HD
ws['B1'] = 'HƯỚNG DẪN GỬI DỰ TRÙ THÁNG {{T}}'
ws['B1'].font = f(16, True, '000000')
ws['B2'] = 'Nhà phân phối: {{NPP}}   |   Hạn gửi MMH: {{HAN}}'
ws['B2'].font = f(12, True, BLUE)
ws['B3'] = 'Nếu bảng chưa hiển thị số liệu, vui lòng bấm "Enable Editing" (Cho phép chỉnh sửa) trên thanh thông báo màu vàng.'
ws['B3'].font = f(11, True, RED)
ws['B5'] = 'CÁC BƯỚC'; ws['B5'].font = f(12, True, NAVY)
steps = [
    '1.  Mở sheet DỰ TRÙ. Ở cột B, chọn mã cần dự trù (bấm vào ô rồi chọn trong danh sách). Thông tin của mã hiện ra ngay trên dòng đó.',
    '2.  Nhập số lượng dự trù tháng {{T}} ở cột SỐ DỰ TRÙ (nền vàng). Mỗi mã một dòng.',
    '3.  Trước khi nhập, xem cột Cảnh báo và Thông tin cần lưu ý: nếu còn nhiều dự trù chưa đặt hàng, nên giảm hoặc không dự trù thêm.',
    '4.  Muốn xem lịch sử 10 tháng của một mã: chọn mã ở ô "Mã cần xem" phía trên danh sách (để trống thì hiển thị mã vừa chọn gần nhất).',
    '5.  Lưu file (giữ định dạng .xlsx) và gửi lại MMH trước hạn.',
]
for i, s_ in enumerate(steps):
    c = ws.cell(6 + i, 2, s_); c.font = f(11); c.alignment = al('left', 'top', True)
    ws.merge_cells(start_row=6 + i, start_column=2, end_row=6 + i, end_column=9); ws.row_dimensions[6 + i].height = 22
ws['B13'] = 'CẢNH BÁO'; ws['B13'].font = f(12, True, NAVY)
legend = [('Cao', RED_BG, RED, 'Dự trù chưa đặt hàng từ {{THH}} tổng dự trù đã gửi trong kỳ trở lên. Nên đặt hết phần đã dự trù trước, hạn chế dự trù thêm.'),
          ('Theo dõi', ORANGE_BG, ORANGE, 'Dự trù chưa đặt hàng từ {{THW}} đến dưới {{THH}}. Cân nhắc giảm dự trù.'),
          ('Ổn', GREEN_BG, GREEN, 'Dự trù chưa đặt hàng dưới {{THW}}.'),
          ('Chưa dự trù', PALE, BLUE, 'Có đặt hàng nhưng chưa gửi dự trù. Nên dự trù để MMH chuẩn bị hàng.')]
for i, (v, bg, fc, t) in enumerate(legend):
    c = ws.cell(14 + i, 2, v); c.fill = fill(bg); c.font = f(10, True, fc); c.alignment = al('center'); c.border = BORDER_LIGHT
    ws.cell(14 + i, 3, t).font = f(10)
    ws.merge_cells(start_row=14 + i, start_column=3, end_row=14 + i, end_column=9)
ws['B19'] = 'GIẢI THÍCH'; ws['B19'].font = f(12, True, NAVY)
expl = [
    ('Hai việc hằng tháng', 'Gửi dự trù cho tháng thứ 3 kế tiếp để MMH chuẩn bị hàng (tháng {{CM}} gửi dự trù tháng {{T}}); đặt hàng trong tháng, có thể nhiều lần.'),
    ('Dự trù đã gửi / đã đặt hàng', 'Tổng dự trù đã gửi và tổng đơn đã đặt từ đầu kỳ ({{K0}}) đến hết {{LASTP}}.'),
    ('Dự trù chưa đặt hàng', 'Phần đã dự trù nhưng chưa đặt hàng, cộng dồn từ đầu kỳ. MMH đã chuẩn bị hàng cho phần này.'),
    ('Chưa đặt hàng dự kiến', 'Phần dự trù chưa đặt hàng còn lại đến hết {{TM1}}, sau khi tính dự trù đã gửi và đặt hàng dự kiến của {{EDITR}}.'),
    ('Dự trù đề xuất', 'Đặt hàng dự kiến tháng {{T}} trừ phần dự trù chưa đặt hàng dự kiến. Ví dụ: dự kiến đặt 100, còn 30 chưa đặt → đề xuất 70.'),
]
header_row(ws, 20, 2, ['Mục', 'Ý nghĩa'], 22)
for cc in range(4, 10):
    ws.cell(20, cc).fill = fill(LIGHT)
ws.merge_cells('C20:I20')
for i, (k, v) in enumerate(expl):
    r = 21 + i
    a = ws.cell(r, 2, k); a.font = f(10, True); a.border = BORDER_LIGHT; a.alignment = al('left', 'center', True)
    b = ws.cell(r, 3, v); b.font = f(10); b.alignment = al('left', 'center', True)
    for cc in range(3, 10):
        ws.cell(r, cc).border = BORDER_LIGHT
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9); ws.row_dimensions[r].height = 30
c = ws.cell(27, 2, 'Lưu ý: không thay đổi bố cục bảng (không chèn, xoá hoặc di chuyển cột). Đổi tên sheet hoặc tên file vẫn nhập được bình thường.')
c.font = f(11, True, INK); ws.merge_cells('B27:I27')
widths(ws, {'B': 28, 'C': 14, 'D': 12, 'E': 10, 'F': 14, 'G': 12, 'H': 20, 'I': 40})
protect(ws)
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.calculation.fullCalcOnLoad = True
wb.active = 1
wb.save(OUT)
print('template saved', OUT, os.path.getsize(OUT), 'bytes')
