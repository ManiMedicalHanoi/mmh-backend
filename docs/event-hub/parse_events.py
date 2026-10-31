# -*- coding: utf-8 -*-
"""Đọc file kế hoạch sự kiện (mẫu MKT) ⇒ dữ liệu chuẩn hoá cho file tổng.
Bản Python để thử trên file .xlsx; sẽ chép sang Apps Script (đọc Google Sheet trong thư mục Event FY68)."""
import sys, json, re, glob, os, datetime as dt, unicodedata
from openpyxl import load_workbook

def norm(s):
    s = unicodedata.normalize('NFD', str(s or '')).lower()
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').replace('đ', 'd')
    return re.sub(r'[^a-z0-9]+', ' ', s).strip()

def S(v):
    if v is None: return ''
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime('%d/%m/%Y')
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def grid(ws):
    mr, mc = min(ws.max_row, 400), min(ws.max_column, 20)
    g = [[S(c.value) for c in row] for row in ws.iter_rows(min_row=1, max_row=mr, max_col=mc)]
    # ô gộp: điền giá trị ô đầu cho cả vùng (để đọc bảng dễ)
    for rg in ws.merged_cells.ranges:
        v = S(ws.cell(rg.min_row, rg.min_col).value)
        for r in range(rg.min_row, min(rg.max_row, mr) + 1):
            for c in range(rg.min_col, min(rg.max_col, mc) + 1):
                if r - 1 < len(g) and c - 1 < len(g[r - 1]): g[r - 1][c - 1] = v
    return g

def find_header(g, must, start=0):
    must = [norm(m) for m in must]
    for i in range(start, len(g)):
        row = [norm(x) for x in g[i]]
        if all(any(m == x or x.startswith(m) for x in row) for m in must): return i
    return -1

def col_of(hdr, names):
    h = [norm(x) for x in hdr]
    for n in names:
        n = norm(n)
        for j, x in enumerate(h):
            if x == n or x.startswith(n): return j
    return -1

OV_KEYS = {
    'ten su kien': 'name', 'ngay to chuc': 'date_text', 'don vi to chuc': 'organizer', 'dia diem': 'venue', 'thoi gian': 'time',
    'so luong khach': 'guests', 'hinh thuc su kien': 'format', 'muc dich su kien': 'purpose', 'doi tuong huong toi': 'audience',
    'quy mo': 'scale', 'bao cao vien': 'speakers', 'dang ky phi': 'fee', 'npp tham du': 'distributor', 'san pham trong tam': 'products',
    'sale mmh ho tro': 'sales', 'doi tac to chuc': 'partner', 'kpi': 'kpi',
}

def parse_overview(g):
    info, tasks = {}, []
    for i, row in enumerate(g[:30]):
        for j, cell in enumerate(row):
            k = OV_KEYS.get(norm(cell))
            if k and k not in info:
                val = next((row[x] for x in range(j + 1, len(row)) if row[x] and row[x] != cell), '')
                info[k] = val
    h = find_header(g, ['hạng mục', 'chi tiết'])
    if h >= 0:
        hdr = g[h]; cg, cd, cp, cn = col_of(hdr, ['hạng mục']), col_of(hdr, ['chi tiết']), col_of(hdr, ['đơn vị phụ trách']), col_of(hdr, ['ghi chú'])
        grp = ''
        for row in g[h + 1:]:
            if not any(row):
                if tasks: break
                continue
            vals = set(x for x in row if x)
            if len(vals) == 1 and norm(next(iter(vals))).startswith('chuong trinh khuyen mai'): break
            grp = row[cg] or grp
            d = row[cd] if cd >= 0 else ''
            if not d: continue
            notes = [x for x in row[cn:] if x] if cn >= 0 else []
            tasks.append({'group': grp, 'detail': d, 'owner': row[cp] if cp >= 0 else '', 'note': ' / '.join(dict.fromkeys(notes))})
    return info, tasks

def parse_table(g, must, cols, carry=()):
    h = find_header(g, must)
    if h < 0: return []
    hdr = g[h]; idx = {k: col_of(hdr, v) for k, v in cols.items()}
    out, last = [], {}
    for row in g[h + 1:]:
        if not any(row):
            if out: break
            continue
        rec = {k: (row[j] if j >= 0 and j < len(row) else '') for k, j in idx.items()}
        for k in carry:
            if rec[k]: last[k] = rec[k]
            else: rec[k] = last.get(k, '')
        out.append(rec)
    return out

def parse_staff(g):
    h = find_header(g, ['công ty', 'họ tên'])
    if h < 0: return []
    hdr = g[h]; cc, cn = col_of(hdr, ['công ty']), col_of(hdr, ['họ tên']); cnote = col_of(hdr, ['ghi chú'])
    cno = max(col_of(hdr, ['no', 'stt']), 0)
    out, start = [], h + 1
    while start < len(g) and not re.match(r'^\d', g[start][cno] or ''): start += 1
    # nhãn giai đoạn: các dòng giữa tiêu đề và dòng dữ liệu đầu tiên
    phases = {}
    for j in range(cn + 1, cnote if cnote > 0 else len(hdr)):
        lab = ' · '.join(dict.fromkeys(g[r][j] for r in range(h + 1, start) if g[r][j] and g[r][j] != hdr[j]))
        phases[j] = lab.replace('\n', ' ')
    for row in g[start:]:
        if not any(row[cno + 1:]): continue
        if not (row[cc] or row[cn]): continue
        tasks = [{'phase': phases.get(j, ''), 'task': row[j]} for j in phases if row[j]]
        out.append({'company': row[cc], 'name': row[cn], 'tasks': tasks, 'note': row[cnote] if cnote > 0 else ''})
    return out

def parse_activities(g):
    acts = []
    # bảng CTKM
    promo = parse_table(g, ['sản phẩm', 'chương trình khuyến mại'],
                        {'product': ['sản phẩm'], 'offer': ['chương trình khuyến mại'], 'gift': ['quà tặng'], 'qty': ['sl quà'], 'note': ['ghi chú']})
    # bảng hoạt động (No | Hoạt động | Thời lượng | Nội dung | PIC | Ghi chú)
    act = parse_table(g, ['hoạt động', 'thời lượng'],
                      {'title': ['hoạt động'], 'duration': ['thời lượng'], 'content': ['nội dung'], 'pic': ['pic'], 'note': ['ghi chú']})
    # bảng lớp / phiên theo cột (Hạng mục | LỚP SÁNG | LỚP CHIỀU)
    sessions = []
    h = find_header(g, ['hạng mục'])
    if h >= 0 and not promo and not act:
        hdr = g[h]; c0 = col_of(hdr, ['hạng mục'])
        for j in range(c0 + 1, len(hdr)):
            if not hdr[j] or hdr[j] == hdr[c0]: continue
            s = {'session': hdr[j]}
            for row in g[h + 1:]:
                if row[c0] and j < len(row) and row[j]: s[row[c0]] = row[j]
            sessions.append(s)
    return promo, act, sessions

def parse_kv(g):
    out = {}
    for row in g:
        xs = [x for x in row if x]
        if len(xs) >= 2: out[xs[0]] = xs[1]
    return out

def parse_file(path):
    wb = load_workbook(path, data_only=True)
    ev = {'file': os.path.basename(path), 'info': {}, 'tasks': [], 'promo': [], 'activities': [], 'sessions': [], 'staff': [],
          'agenda': [], 'checklist': [], 'booth': {}, 'gifts': [], 'results': [], 'mc': []}
    for ws in wb.worksheets:
        t = norm(ws.title); g = grid(ws)
        if t.startswith('overview'):
            ev['info'], ev['tasks'] = parse_overview(g)
            p, a, s = parse_activities(g); ev['promo'] += p
        elif t.startswith('activities'):
            p, a, s = parse_activities(g); ev['promo'] += p; ev['activities'] += a; ev['sessions'] += s
        elif 'staff' in t or t.startswith('nhan su'):
            ev['staff'] += parse_staff(g)
        elif t.startswith('checklist'):
            ev['checklist'] += [r for r in parse_table(g, ['marketing material'],
                {'no': ['stt'], 'group': ['group'], 'category': ['category'], 'item': ['marketing material'], 'unit': ['unit'],
                 'qty': ['quantity'], 'note': ['note'], 'pic': ['pic'], 'status': ['status']}, carry=('no', 'group', 'category')) if r['item']]
        elif t.startswith('agenda'):
            ev['agenda'] += [r for r in parse_table(g, ['detail', 'pic'],
                {'no': ['no'], 'date': ['date'], 'kind': ['kind'], 'time': ['time'], 'detail': ['detail'], 'pic': ['pic'], 'note': ['note']},
                carry=('no', 'date', 'kind')) if r['detail']]
        elif t.startswith('booth'):
            ev['booth'] = parse_kv(g)
        elif t.startswith('qua ung'):
            ev['gifts'] += [r for r in parse_table(g, ['nhà phân phối'],
                {'npp': ['nhà phân phối'], 'item': ['sản phẩm'], 'advance': ['tổng tạm ứng'], 'used': ['tổng xuất'], 'diff': ['chênh lệch'], 'status': ['tình trạng']},
                carry=('npp',)) if r['item']]
        elif t.startswith('report'):
            ev['results'] += [r for r in parse_table(g, ['chỉ số'], {'no': ['stt'], 'metric': ['chỉ số'], 'unit': ['đơn vị'], 'actual': ['thực đạt'], 'note': ['ghi chú']},
                carry=('no',)) if r['metric']]
        elif 'kich ban' in t:
            ev['mc'] += parse_table(g, ['thời gian', 'hoạt động'], {'time': ['thời gian'], 'activity': ['hoạt động'], 'detail': ['chi tiết']})
    # khoá & ngày từ tên file: yyyymmdd[-dd]_Type_..._City
    m = re.search(r'(\d{8})(?:-(\d{2}))?_([^_]+(?:_[^_]+)*)', os.path.basename(path))
    if m:
        d0 = dt.datetime.strptime(m.group(1), '%Y%m%d').date()
        d1 = d0.replace(day=int(m.group(2))) if m.group(2) else d0
        ev['start'], ev['end'] = d0.isoformat(), d1.isoformat()
        stem = re.sub(r'^[0-9a-f]{8}-', '', os.path.splitext(os.path.basename(path))[0])
        ev['code'] = stem[:60]
        parts = stem.split('_')
        ev['type'] = 'Exhibition' if 'Exhibition' in stem else ('Workshop' if 'Workshop' in stem else ('Seminar' if 'Seminar' in stem else 'Event'))
        CITY = [('Hai_Phong', 'Hải Phòng'), ('Haiphong', 'Hải Phòng'), ('Ninh_Binh', 'Ninh Bình'), ('Hanoi', 'Hà Nội'), ('Ha_Noi', 'Hà Nội'),
                ('HCM', 'TP.HCM'), ('Ho_Chi_Minh', 'TP.HCM'), ('Da_Nang', 'Đà Nẵng')]
        ev['city'] = next((v for k, v in CITY if stem.endswith(k)), parts[-1])
    return ev

if __name__ == '__main__':
    evs = [parse_file(f) for f in sorted(glob.glob(sys.argv[1] + '/*.xlsx'))]
    json.dump(evs, open(sys.argv[2], 'w'), ensure_ascii=False, indent=1)
    for e in evs:
        print(e['start'], e['type'], e['city'], '| info', len(e['info']), 'tasks', len(e['tasks']), 'promo', len(e['promo']), 'act', len(e['activities']),
              'sess', len(e['sessions']), 'staff', len(e['staff']), 'agenda', len(e['agenda']), 'check', len(e['checklist']),
              'gifts', len(e['gifts']), 'results', len(e['results']), 'mc', len(e['mc']), 'booth', len(e['booth']))
