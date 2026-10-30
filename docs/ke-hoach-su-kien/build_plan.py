# -*- coding: utf-8 -*-
"""File chung: Kế hoạch & tiến độ sự kiện MMH (mọi sự kiện, Owner giao việc, Dashboard)."""
import datetime as dt, sys, json, hashlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

OUT = sys.argv[1]
TODAY = dt.date(2026, 10, 30)
F = 'Arial'
NAVY, BLUE, GRAY = '003047', '3A5CAA', '5B6B7B'
SOFT, INPUT = 'F0F7FF', 'FFF7D6'
RED, AMBER, GREEN, MUTE = 'FDE4E1', 'FFF4D6', 'EAF6EE', 'B0B8C2'
thin = Side(style='thin', color='D5DCE4')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def font(**k): k.setdefault('name', F); k.setdefault('size', 10); return Font(**k)
def fill(c): return PatternFill('solid', start_color=c, end_color=c)
WR = Alignment(wrap_text=True, vertical='center')
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()
def sheet(n, first=False):
    ws = wb.active if first else wb.create_sheet(); ws.title = n; ws.sheet_view.showGridLines = False; return ws
def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)
def title(ws, t, sub=None):
    ws['A1'] = t; ws['A1'].font = font(size=16, bold=True, color=NAVY); ws.row_dimensions[1].height = 26
    if sub: ws['A2'] = sub; ws['A2'].font = font(size=9, italic=True, color=GRAY)
def header(ws, row, cols, edit=(), start=1):
    for i, (h, w) in enumerate(cols):
        c = ws.cell(row, start + i, h)
        c.font = font(bold=True, color='FFFFFF', size=9); c.fill = fill(BLUE if i in edit else NAVY)
        c.alignment = CEN; c.border = BOX
        if w is not None: ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[row].height = 32
def dv_list(ws, items, rng):
    d = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True); ws.add_data_validation(d); d.add(rng)
def dv_ref(ws, ref, rng):
    d = DataValidation(type='list', formula1=ref, allow_blank=True); ws.add_data_validation(d); d.add(rng)

STATUS = ['Chưa làm', 'Đang làm', 'Chờ đối tác', 'Xong', 'Bỏ qua']
ROLES = ['MKT Leader', 'Digital MKT', 'Trade MKT', 'Sales', 'NPP / đối tác', 'Giám đốc']

wsD = sheet('Dashboard', first=True)
wsG = sheet('Giao việc')
wsE = sheet('Sự kiện')
wsT = sheet('Mẫu việc')
wsC = sheet('Cài đặt')
wsH = sheet('Hướng dẫn')

# ═════════ Cài đặt
C = wsC
title(C, 'CÀI ĐẶT', 'Danh sách người (để chọn ở cột Người phụ trách) · người mặc định theo vai trò (dùng khi tạo việc cho sự kiện mới) · ngày nghỉ lễ.')
header(C, 4, [('Vai trò', 18), ('Người mặc định', 22)], edit={1})
PEOPLE = [('Leader MKT (ví dụ)', 'Marketing'), ('Digital (ví dụ)', 'Marketing'), ('Trade (ví dụ)', 'Marketing'),
          ('Sales Dental (ví dụ)', 'Sales'), ('Sales Surgical (ví dụ)', 'Sales'), ('NPP (ví dụ)', 'Đối tác'), ('Giám đốc (ví dụ)', 'BGĐ')]
DEF = {'MKT Leader': PEOPLE[0][0], 'Digital MKT': PEOPLE[1][0], 'Trade MKT': PEOPLE[2][0], 'Sales': PEOPLE[3][0],
       'NPP / đối tác': PEOPLE[5][0], 'Giám đốc': PEOPLE[6][0]}
for i, r_ in enumerate(ROLES):
    C.cell(5 + i, 1, r_).border = BOX; c = C.cell(5 + i, 2, DEF[r_]); c.fill = fill(INPUT); c.border = BOX
header(C, 4, [('Danh sách người', 24), ('Bộ phận', 12), ('Email (để nhắc việc sau này)', 28)], edit={0, 1, 2}, start=4)
C.column_dimensions['C'].width = 3
for i in range(40):
    r = 5 + i
    for j in range(3): c = C.cell(r, 4 + j); c.border = BOX; c.fill = fill(INPUT)
    if i < len(PEOPLE): C.cell(r, 4, PEOPLE[i][0]); C.cell(r, 5, PEOPLE[i][1])
header(C, 4, [('Ngày nghỉ lễ (T2–T6)', 18)], edit={0}, start=8)
C.column_dimensions['G'].width = 3
HOL = ['2026-01-01', '2026-02-16', '2026-02-17', '2026-02-18', '2026-02-19', '2026-02-20', '2026-04-27', '2026-04-30', '2026-05-01',
       '2026-09-02', '2026-09-03', '2026-09-04', '2027-01-01', '2027-04-30', '2027-09-02']
for i in range(60):
    c = C.cell(5 + i, 8); c.border = BOX; c.fill = fill(INPUT); c.number_format = 'dd/mm/yyyy'
    if i < len(HOL): c.value = dt.date.fromisoformat(HOL[i])
C['H66'] = '2027: thêm Tết, Giỗ Tổ, nghỉ bù khi có lịch MMH.'; C['H66'].font = font(size=8, italic=True, color=GRAY)
dv_ref(C, "'Cài đặt'!$D$5:$D$44", 'B5:B10')
name('PEOPLE', "'Cài đặt'!$D$5:$D$44"); name('HOLI', "'Cài đặt'!$H$5:$H$64")
C.freeze_panes = 'A5'

# ═════════ Mẫu việc
T = wsT
title(T, 'MẪU VIỆC CHO 1 SỰ KIỆN',
      'Theo SOP MMHSOP_0XX. Sự kiện mới được tạo việc theo bảng này (sửa / thêm / xoá dòng tuỳ ý — chỉ áp cho sự kiện tạo SAU khi sửa). D − n = ngày lịch, D + n = ngày làm việc.')
L_, W_ = 'Lịch', 'Làm việc'
TPL = [
 # mã, hạng mục, việc, vai trò, phối hợp, số ngày, tính theo, áp dụng
 ('K01', '1. Kế hoạch', 'Lập kế hoạch tổng quan (KOL, số khách, mục đích, KPI bán hàng, khuyến mại, khảo sát) + xin Giám đốc duyệt ngân sách', 'MKT Leader', 'Sales, NPP', -30, L_, 'Tất cả'),
 ('K02', '1. Kế hoạch', 'Chia việc với NPP, Trường / Viện (tuyển sinh, đồ thực hành, MC, số bàn trưng bày)', 'MKT Leader', 'NPP, Viện', -30, L_, 'Tất cả'),
 ('K03', '2. KOL & đối tác', 'Chốt KOL: hợp đồng, lịch di chuyển, xem trước bài giảng, người chăm sóc KOL', 'MKT Leader', 'KOL', -30, L_, 'Tất cả'),
 ('K04', '3. Tuyển sinh', 'Mở tuyển sinh: form đăng ký (khảo sát trước, đồng ý dữ liệu, nguồn đăng ký) + nhóm Zalo khách', 'Digital MKT', 'Sales, NPP', -30, L_, 'Tất cả'),
 ('K05', '3. Tuyển sinh', 'Báo cáo số đăng ký theo nguồn cho team (thứ Hai hằng tuần, đến D-1)', 'MKT Leader', 'Sales, NPP', -1, L_, 'Tất cả'),
 ('K06', '3. Tuyển sinh', 'Checkpoint D-14: đăng ký < 60% target ⇒ báo Giám đốc phương án', 'MKT Leader', 'Giám đốc', -14, L_, 'Tất cả'),
 ('K07', '3. Tuyển sinh', 'Sales / NPP gửi đủ thông tin khách đăng ký hộ (Họ tên · SĐT · Nơi công tác)', 'Sales', 'NPP', -5, L_, 'Tất cả'),
 ('K08', '4. Truyền thông', 'Đăng chuỗi bài truyền thông sự kiện', 'Digital MKT', 'Sales, NPP chia sẻ', -14, L_, 'Tất cả'),
 ('K09', '5. Địa điểm & hậu cần', 'Đặt địa điểm (ngày, phòng, sức chứa, báo giá)', 'Trade MKT', 'Địa điểm', -30, L_, 'Trực tiếp'),
 ('K10', '5. Địa điểm & hậu cần', 'Chốt setup với địa điểm: sơ đồ, tea break / ăn trưa, kỹ thuật, sự kiện khác cùng giờ', 'Trade MKT', 'Địa điểm', -14, L_, 'Trực tiếp'),
 ('K11', '5. Địa điểm & hậu cần', 'Webinar: tạo phòng Zoom (đăng ký, ghi hình, báo cáo người tham dự)', 'Digital MKT', 'KOL', -30, L_, 'Webinar'),
 ('K12', '6. Nội dung', 'Timeline chi tiết ngày D', 'MKT Leader', 'KOL, Viện', -7, L_, 'Tất cả'),
 ('K13', '6. Nội dung', 'Kịch bản MC, kịch bản phát biểu, video trình chiếu', 'Digital MKT', 'KOL, Viện', -7, L_, 'Tất cả'),
 ('K14', '7. Ấn phẩm & đồ', 'Thiết kế & in: backdrop, standee, thẻ đeo, bảng tên, folder, QR khảo sát, biểu mẫu check-in ký tay', 'Digital MKT', 'Trade MKT', -7, L_, 'Trực tiếp'),
 ('K15', '7. Ấn phẩm & đồ', 'Đồ marketing + đồ thực hành theo phiên (nguồn, báo giá, đồ mượn)', 'Trade MKT', 'KOL, NPP', -10, L_, 'Trực tiếp'),
 ('K16', '7. Ấn phẩm & đồ', 'VIP / đại biểu (nếu có Viện): danh sách, hoa, bảng tên mica, quà, lời giới thiệu', 'Trade MKT', 'Viện', -7, L_, 'Trực tiếp'),
 ('K17', '7. Ấn phẩm & đồ', 'Gửi đồ đi tỉnh + xác nhận địa điểm nhận đủ', 'Trade MKT', 'Địa điểm', -7, L_, 'Ngoài Hà Nội'),
 ('K18', '8. Chốt trước D', 'Họp với NPP: agenda, nhân sự, bán hàng, sản phẩm', 'MKT Leader', 'NPP, Sales', -3, L_, 'Tất cả'),
 ('K19', '8. Chốt trước D', 'Rà soát dữ liệu khách (thiếu thông tin / chưa đồng ý dữ liệu)', 'Digital MKT', 'Sales, NPP', -3, L_, 'Tất cả'),
 ('K20', '8. Chốt trước D', 'Phân công nhân sự trước / trong / sau, giờ có mặt', 'MKT Leader', 'Sales, NPP', -3, L_, 'Tất cả'),
 ('K21', '8. Chốt trước D', 'Họp review D-1 + gọi nhắc khách + set up địa điểm / chạy thử Zoom', 'MKT Leader', 'Toàn team', -1, L_, 'Tất cả'),
 ('K22', '9. Ngày D', 'Check-in có chữ ký + đồng ý dữ liệu (webinar: xuất báo cáo Zoom ngay sau buổi)', 'Digital MKT', 'Trade, NPP', 0, L_, 'Tất cả'),
 ('K23', '9. Ngày D', 'Điều phối chương trình, khảo sát QR trước bế mạc, ảnh / video', 'MKT Leader', 'Digital MKT', 0, L_, 'Tất cả'),
 ('K24', '9. Ngày D', 'Bán hàng tại bàn trưng bày + ghi khách quan tâm', 'Sales', 'NPP', 0, L_, 'Trực tiếp'),
 ('K25', '9. Ngày D', 'Thu dọn, kiểm đếm, ghi POSM hỏng', 'Trade MKT', 'NPP', 0, L_, 'Trực tiếp'),
 ('K26', '10. Sau sự kiện', 'Bài cảm ơn + ảnh recap; video recap (nếu có)', 'Digital MKT', '', 2, W_, 'Tất cả'),
 ('K27', '10. Sau sự kiện', 'Scan biểu mẫu ký tay + upload ảnh / video lên folder sự kiện', 'Trade MKT', 'Digital MKT', 2, W_, 'Trực tiếp'),
 ('K28', '10. Sau sự kiện', 'Đối chiếu dữ liệu khách (loại trùng, đủ trường) ⇒ CUSTOMER LIST; tổng hợp khảo sát', 'Digital MKT', 'Sales, NPP', 3, W_, 'Tất cả'),
 ('K29', '10. Sau sự kiện', 'Sales + NPP gửi số liệu bán hàng tại sự kiện', 'Sales', 'NPP', 3, W_, 'Tất cả'),
 ('K30', '10. Sau sự kiện', 'Bàn giao danh sách khách cần follow-up cho Sales', 'Digital MKT', 'Sales', 3, W_, 'Tất cả'),
 ('K31', '10. Sau sự kiện', 'Báo cáo sự kiện + chứng từ KOL, hoá đơn địa điểm', 'MKT Leader', 'Trade MKT', 5, W_, 'Tất cả'),
 ('K32', '10. Sau sự kiện', 'Email báo cáo toàn MMH (4 file đính kèm) + hoàn tất lưu hồ sơ', 'MKT Leader', '', 7, W_, 'Tất cả'),
 ('K33', '11. Follow-up', 'Sales liên hệ lần đầu khách cần follow-up', 'Sales', 'NPP', 5, W_, 'Tất cả'),
 ('K34', '11. Follow-up', 'Sales báo kết quả chuyển đổi đơn hàng', 'Sales', 'NPP', 20, W_, 'Tất cả'),
]
header(T, 4, [('Mã', 6), ('Hạng mục', 18), ('Việc', 64), ('Vai trò mặc định', 14), ('Phối hợp', 16), ('Số ngày so với D', 9), ('Tính theo', 10), ('Áp dụng cho', 11)],
       edit=set(range(8)))
for i in range(80):
    r = 5 + i
    if i < len(TPL):
        for j, v in enumerate(TPL[i]): T.cell(r, 1 + j, v)
    for j in range(8):
        c = T.cell(r, 1 + j); c.border = BOX; c.alignment = WR if j in (2,) else Alignment(vertical='center'); c.font = font(size=9)
dv_list(T, ROLES, 'D5:D84'); dv_list(T, [L_, W_], 'G5:G84'); dv_list(T, ['Tất cả', 'Trực tiếp', 'Webinar', 'Ngoài Hà Nội'], 'H5:H84')
T.freeze_panes = 'D5'

# ═════════ Sự kiện
E = wsE
ER0, ERN = 5, 304
title(E, 'DANH SÁCH SỰ KIỆN',
      'Cột xám TỰ ĐỒNG BỘ từ file "2. Vietnam - Marketing Offline FY68" (PLAN_EVENT + REPORT_EVENT, 10 phút / lần) — sửa tên / ngày ở file Offline. Cột xanh do Owner nhập.')
ECOLS = [('Mã SK', 11), ('Tên sự kiện', 38), ('Ngày D', 11), ('Ngày kết thúc', 11), ('Loại sự kiện', 14), ('Hình thức', 10), ('Địa điểm', 14),
         ('KOL', 16), ('Nguồn', 8), ('Trạng thái (file Offline)', 13), ('Owner', 18), ('Theo dõi', 8),
         ('Việc', 6), ('Xong', 6), ('Quá hạn', 7), ('Chưa giao', 7), ('%', 6), ('Tiến độ', 14), ('Thời gian', 14), ('_k', 3)]
header(E, 4, ECOLS, edit={10, 11})
SAMPLE_EV = [
    ('Hội thảo nha khoa tháng 10 (ví dụ)', '2026-10-17', '', 'Hội thảo', 'Trực tiếp', 'Hà Nội', 'KOL A', 'REPORT', 'Done'),
    ('Khoá học nội nha thực hành – Hà Nội (ví dụ)', '2026-11-08', '', 'Workshop', 'Trực tiếp', 'Hà Nội', 'KOL B', 'PLAN', 'Confirmed'),
    ('Webinar phục hồi Composite (ví dụ)', '2026-11-20', '', 'Webinar', 'Webinar', 'Zoom', 'KOL C', 'PLAN', 'Confirmed'),
    ('Hội thảo nhãn khoa – Đà Nẵng (ví dụ)', '2026-12-12', '2026-12-13', 'Hội thảo', 'Trực tiếp', 'Đà Nẵng', 'KOL D', 'PLAN', 'Planning'),
    ('Workshop phẫu thuật – TP.HCM (ví dụ)', '2027-01-16', '', 'Workshop', 'Trực tiếp', 'TP.HCM', 'KOL E', 'PLAN', 'Planning'),
]
EVK = []
for i in range(ERN - ER0 + 1):
    r = ER0 + i
    if i < len(SAMPLE_EV):
        nm, s, e, tp, ht, pl, kol, src, st = SAMPLE_EV[i]
        k = 'VD-' + str(i + 1); EVK.append((k, nm, dt.date.fromisoformat(s), dt.date.fromisoformat(e or s), ht, pl))
        vals = [k, nm, dt.date.fromisoformat(s), dt.date.fromisoformat(e) if e else None, tp, ht, pl, kol, src, st, PEOPLE[0][0], 'Có']
        for j, v in enumerate(vals): E.cell(r, 1 + j, v)
    G = "'Giao việc'"
    E[f'M{r}'] = f'=IF(A{r}="","",COUNTIFS({G}!$M$5:$M$3004,A{r},{G}!$I$5:$I$3004,"<>Bỏ qua"))'
    E[f'N{r}'] = f'=IF(A{r}="","",COUNTIFS({G}!$M$5:$M$3004,A{r},{G}!$I$5:$I$3004,"Xong"))'
    E[f'O{r}'] = f'=IF(A{r}="","",COUNTIFS({G}!$M$5:$M$3004,A{r},{G}!$L$5:$L$3004,"*QUÁ HẠN*"))'
    E[f'P{r}'] = f'=IF(A{r}="","",COUNTIFS({G}!$M$5:$M$3004,A{r},{G}!$L$5:$L$3004,"CHƯA GIAO*"))'
    E[f'Q{r}'] = f'=IF(OR(A{r}="",N(M{r})=0),"",N{r}/M{r})'
    E[f'R{r}'] = f'=IF(Q{r}="","",REPT("█",ROUND(Q{r}*10,0))&REPT("░",10-ROUND(Q{r}*10,0)))'
    E[f'S{r}'] = (f'=IF(C{r}="","Chưa có ngày",IF(C{r}>TODAY(),"Còn "&(C{r}-TODAY())&" ngày",'
                  f'IF(IF(D{r}="",C{r},D{r})>=TODAY(),"Đang diễn ra","Đã qua "&(TODAY()-IF(D{r}="",C{r},D{r}))&" ngày")))')
    E[f'T{r}'] = f'=IF(AND(A{r}<>"",L{r}="Có",ISNUMBER(C{r})),IF(IF(D{r}="",C{r},D{r})>=TODAY()-35,C{r}+ROW()/100000,""),"")'
    for j in range(20):
        c = E.cell(r, 1 + j); c.border = BOX; c.alignment = Alignment(vertical='center', wrap_text=(j == 1))
        c.font = font(size=9, color=GRAY) if j in (0, 8, 9) else font(size=9)
    E[f'C{r}'].number_format = 'dd/mm/yy'; E[f'D{r}'].number_format = 'dd/mm/yy'; E[f'Q{r}'].number_format = '0%'
    E[f'R{r}'].font = font(size=9, color=BLUE)
    E[f'K{r}'].fill = fill(INPUT); E[f'L{r}'].fill = fill(INPUT)
E.column_dimensions['T'].hidden = True
dv_ref(E, 'PEOPLE', f'K{ER0}:K{ERN}'); dv_list(E, ['Có', 'Không'], f'L{ER0}:L{ERN}')
E.conditional_formatting.add(f'A{ER0}:S{ERN}', FormulaRule(formula=[f'$L{ER0}="Không"'], font=Font(color=MUTE, name=F)))
E.conditional_formatting.add(f'O{ER0}:O{ERN}', FormulaRule(formula=[f'N(O{ER0})>0'], fill=fill(RED), font=Font(color='C0392B', bold=True, name=F)))
E.conditional_formatting.add(f'P{ER0}:P{ERN}', FormulaRule(formula=[f'N(P{ER0})>0'], fill=fill(AMBER), font=Font(bold=True, name=F)))
E.freeze_panes = 'C5'; E.auto_filter.ref = f'A4:S{ERN}'
name('SK_A', f"'Sự kiện'!$A${ER0}:$A${ERN}"); name('SK_B', f"'Sự kiện'!$B${ER0}:$B${ERN}"); name('SK_C', f"'Sự kiện'!$C${ER0}:$C${ERN}")
name('SK_D', f"'Sự kiện'!$D${ER0}:$D${ERN}"); name('SK_L', f"'Sự kiện'!$L${ER0}:$L${ERN}")

# ═════════ Giao việc
Gv = wsG
GR0, GRN = 5, 3004
title(Gv, 'GIAO VIỆC & THEO DÕI',
      'Mọi việc của mọi sự kiện. OWNER chọn Người phụ trách (cột F, có thể đổi Hạn ở cột H). NGƯỜI PHỤ TRÁCH cập nhật Trạng thái · Ngày xong · Ghi chú. Lọc cột F để xem việc của mình.')
GCOLS = [('Sự kiện', 30), ('Ngày D', 9), ('Mốc', 7), ('Hạng mục', 15), ('Việc cần làm', 50), ('Người phụ trách', 18), ('Phối hợp', 14),
         ('Hạn', 10), ('Trạng thái', 12), ('Ngày xong', 10), ('Ghi chú / link', 26), ('Cảnh báo', 20),
         ('_sk', 4), ('_ma', 4), ('_n', 4), ('_t', 4), ('_sort', 4), ('_td', 4)]
header(Gv, 4, GCOLS, edit={5, 6, 7, 8, 9, 10})

def g_formulas(r):
    Gv[f'A{r}'] = f'=IF(M{r}="","",IFERROR(INDEX(SK_B,MATCH(M{r},SK_A,0))&"",""))'
    Gv[f'B{r}'] = f'=IF(M{r}="","",IFERROR(INDEX(SK_C,MATCH(M{r},SK_A,0)),""))'
    Gv[f'C{r}'] = f'=IF(O{r}="","",IF(O{r}=0,"D",IF(O{r}<0,"D"&O{r},"D+"&O{r})))'
    end = f'IFERROR(IF(INDEX(SK_D,MATCH(M{r},SK_A,0))="",B{r},INDEX(SK_D,MATCH(M{r},SK_A,0))),B{r})'
    Gv[f'H{r}'] = (f'=IF(OR(B{r}="",O{r}=""),"",IF(O{r}>0,IF(P{r}="Làm việc",WORKDAY({end},O{r},HOLI),{end}+O{r}),B{r}+O{r}))')
    Gv[f'L{r}'] = (f'=IF(E{r}="","",IF(I{r}="Bỏ qua","—",IF(I{r}="Xong",IF(AND(ISNUMBER(J{r}),ISNUMBER(H{r}),J{r}>H{r}),"Xong trễ","Xong"),'
                   f'IF(F{r}="","CHƯA GIAO","")&IF(NOT(ISNUMBER(H{r})),"",IF(TODAY()>H{r},IF(F{r}=""," · ","")&"QUÁ HẠN "&(TODAY()-H{r})&" ngày",'
                   f'IF(H{r}-TODAY()<=3,IF(F{r}=""," · ","")&"Đến hạn "&IF(H{r}=TODAY(),"hôm nay","trong "&(H{r}-TODAY())&" ngày"),""))))))')
    Gv[f'R{r}'] = f'=IF(M{r}="","",IFERROR(INDEX(SK_L,MATCH(M{r},SK_A,0)),""))'
    Gv[f'Q{r}'] = f'=IF(AND(E{r}<>"",R{r}="Có",I{r}<>"Xong",I{r}<>"Bỏ qua",ISNUMBER(H{r})),H{r}+ROW()/100000,"")'

# dữ liệu ví dụ: tạo việc như backend sẽ tạo
def wday_add(d, n):
    hol = {dt.date.fromisoformat(x) for x in HOL}; cur = d; k = 0
    while k < n:
        cur += dt.timedelta(days=1)
        if cur.weekday() < 5 and cur not in hol: k += 1
    return cur
rows = []
import random
random.seed(7)
for k, nm, s, e, ht, pl in EVK:
    for (ma, grp, job, role, co, n, unit, app) in TPL:
        if app == 'Ngoài Hà Nội':
            if ht != 'Trực tiếp' or 'Hà Nội' in pl: continue
        elif app != 'Tất cả' and app != ht: continue
        due = s + dt.timedelta(days=n) if n <= 0 else (wday_add(e, n) if unit == W_ else e + dt.timedelta(days=n))
        who = DEF[role]
        if role == 'Sales' and 'Composite' not in nm and 'nhãn' in nm: who = PEOPLE[4][0]
        st, done = 'Chưa làm', None
        if due < TODAY:
            x = random.random()
            st = 'Xong' if x < 0.82 else ('Đang làm' if x < 0.93 else 'Chưa làm')
            if st == 'Xong': done = due - dt.timedelta(days=random.choice([0, 0, 1, 2])) if random.random() < 0.85 else due + dt.timedelta(days=1)
        elif (due - TODAY).days <= 7:
            st = random.choice(['Đang làm', 'Đang làm', 'Chưa làm', 'Chờ đối tác'])
        if k == 'VD-5' and random.random() < 0.25: who = ''
        rows.append((nm, grp, job, who, co, st, done, k, ma, n, unit))
for i in range(GRN - GR0 + 1):
    r = GR0 + i
    if i < len(rows):
        nm, grp, job, who, co, st, done, k, ma, n, unit = rows[i]
        Gv[f'D{r}'] = grp; Gv[f'E{r}'] = job; Gv[f'F{r}'] = who or None; Gv[f'G{r}'] = co
        Gv[f'I{r}'] = st; Gv[f'J{r}'] = done; Gv[f'M{r}'] = k; Gv[f'N{r}'] = ma; Gv[f'O{r}'] = n; Gv[f'P{r}'] = unit
    if i < 700:          # backend tự ghi công thức cho dòng mới; bản mẫu dựng sẵn 700 dòng
        g_formulas(r)
        for j in range(12):
            c = Gv.cell(r, 1 + j); c.border = BOX; c.alignment = Alignment(vertical='center', wrap_text=(j in (0, 4, 10)))
            c.font = font(size=9, color=GRAY) if j in (1, 2, 11) else font(size=9)
        Gv[f'A{r}'].font = font(size=9)
        Gv[f'B{r}'].number_format = 'dd/mm'; Gv[f'H{r}'].number_format = 'dd/mm/yy'; Gv[f'J{r}'].number_format = 'dd/mm/yy'
        Gv[f'L{r}'].font = font(size=9, bold=True, color=GRAY)
        Gv[f'F{r}'].fill = fill(INPUT)
for col in 'MNOPQR': Gv.column_dimensions[col].hidden = True
dv_ref(Gv, 'PEOPLE', f'F{GR0}:F{GRN}'); dv_list(Gv, STATUS, f'I{GR0}:I{GRN}')
d = DataValidation(type='date', operator='greaterThan', formula1='40000', allow_blank=True); Gv.add_data_validation(d); d.add(f'J{GR0}:J{GRN}')
R = f'A{GR0}:L{GRN}'
Gv.conditional_formatting.add(R, FormulaRule(formula=[f'$I{GR0}="Bỏ qua"'], font=Font(color=MUTE, name=F)))
Gv.conditional_formatting.add(R, FormulaRule(formula=[f'ISNUMBER(SEARCH("QUÁ HẠN",$L{GR0}))'], fill=fill(RED)))
Gv.conditional_formatting.add(R, FormulaRule(formula=[f'ISNUMBER(SEARCH("Đến hạn",$L{GR0}))'], fill=fill(AMBER)))
Gv.conditional_formatting.add(R, FormulaRule(formula=[f'LEFT($L{GR0},4)="Xong"'], fill=fill(GREEN)))
Gv.conditional_formatting.add(f'F{GR0}:F{GRN}', FormulaRule(formula=[f'AND($E{GR0}<>"",$F{GR0}="")'], fill=fill('F6C9A8')))
Gv.conditional_formatting.add(f'A{GR0}:A{GRN}', FormulaRule(formula=[f'$A{GR0}<>$A{GR0 - 1}'], font=Font(bold=True, color=NAVY, name=F)))
Gv.conditional_formatting.add(f'L{GR0}:L{GRN}', FormulaRule(formula=[f'ISNUMBER(SEARCH("QUÁ HẠN",$L{GR0}))'], font=Font(bold=True, color='C0392B', name=F)))
Gv.freeze_panes = 'F5'; Gv.auto_filter.ref = f'A4:L{GRN}'
GV = lambda c: f"'Giao việc'!${c}${GR0}:${c}${GRN}"

# ═════════ Dashboard
D = wsD
title(D, 'DASHBOARD TIẾN ĐỘ SỰ KIỆN', None)
D['A2'] = '="Cập nhật đến "&TEXT(TODAY(),"dd/mm/yyyy")&" · chỉ tính sự kiện đang theo dõi (Sự kiện ▸ Theo dõi = Có)"'
D['A2'].font = font(size=9, italic=True, color=GRAY)
for col, w in zip('ABCDEFGHIJKLMN', [30, 10, 13, 18, 7, 7, 8, 8, 7, 13, 3, 3, 3, 3]): D.column_dimensions[col].width = w
cards = [('A', 'Sự kiện đang theo dõi', f'=COUNT(\'Sự kiện\'!$T${ER0}:$T${ERN})'),
         ('B', 'Việc đang mở', f'=COUNTIFS({GV("R")},"Có",{GV("I")},"<>Xong",{GV("I")},"<>Bỏ qua",{GV("E")},"<>")'),
         ('D', 'QUÁ HẠN', f'=COUNTIFS({GV("R")},"Có",{GV("L")},"*QUÁ HẠN*")'),
         ('F', 'Đến hạn ≤ 3 ngày', f'=COUNTIFS({GV("R")},"Có",{GV("L")},"*Đến hạn*")'),
         ('H', 'Chưa giao người', f'=COUNTIFS({GV("R")},"Có",{GV("L")},"CHƯA GIAO*")'),
         ('J', '% việc đã xong', f'=IFERROR(COUNTIFS({GV("R")},"Có",{GV("I")},"Xong")/COUNTIFS({GV("R")},"Có",{GV("I")},"<>Bỏ qua",{GV("E")},"<>"),0)')]
spans = {'A': 'A', 'B': 'C', 'D': 'E', 'F': 'G', 'H': 'I', 'J': 'J'}
colors = {'D': 'C0392B', 'F': 'B7791F', 'H': 'B85C1E'}
for col, lab, f_ in cards:
    end = spans[col]
    if end != col:
        D.merge_cells(f'{col}4:{end}4'); D.merge_cells(f'{col}5:{end}5')
    c = D[f'{col}4']; c.value = lab; c.font = font(size=9, bold=True, color=GRAY); c.alignment = CEN; c.fill = fill(SOFT)
    v = D[f'{col}5']; v.value = f_; v.font = font(size=22, bold=True, color=colors.get(col, NAVY)); v.alignment = CEN; v.fill = fill(SOFT)
D.row_dimensions[5].height = 40
D['J5'].number_format = '0%'

def sec(cell, t): D[cell] = t; D[cell].font = font(size=12, bold=True, color=NAVY)
def thead(row, heads):
    for j, h in enumerate(heads):
        c = D.cell(row, 1 + j, h); c.font = font(bold=True, color='FFFFFF', size=9); c.fill = fill(NAVY); c.alignment = CEN; c.border = BOX
    D.row_dimensions[row].height = 28

# Tiến độ từng sự kiện
NE = 12
r0 = 8; sec(f'A{r0 - 1}', 'TIẾN ĐỘ TỪNG SỰ KIỆN (sắp tới trước)')
thead(r0, ['Sự kiện', 'Ngày D', 'Thời gian', 'Owner', 'Xong', 'Còn lại', 'Quá hạn', 'Chưa giao', '%', 'Tiến độ'])
EK = lambda c: f"'Sự kiện'!${c}${ER0}:${c}${ERN}"
for i in range(NE):
    r = r0 + 1 + i
    pos = f'IFERROR(MATCH(SMALL({EK("T")},{i + 1}),{EK("T")},0),0)'
    g = lambda c: f'=IF({pos}=0,"",INDEX({EK(c)},{pos}))'
    for j, c_ in enumerate(['B', 'C', 'S', 'K', 'N', None, 'O', 'P', 'Q', 'R']):
        v = (g(c_)[:-1] + '&"")' if c_ in ('B', 'S', 'K', 'R') else g(c_)) if c_ else f'=IF(A{r}="","",N(INDEX({EK("M")},{pos}))-N(E{r})-N(G{r}))'
        cell = D.cell(r, 1 + j, v); cell.border = BOX; cell.font = font(size=9); cell.alignment = Alignment(vertical='center', wrap_text=(j == 0))
    D[f'B{r}'].number_format = 'dd/mm/yy'; D[f'I{r}'].number_format = '0%'; D[f'J{r}'].font = font(size=9, color=BLUE)
    # cột phụ cho biểu đồ: Xong · Còn lại · Quá hạn
    D.row_dimensions[r].height = 30
D.conditional_formatting.add(f'G{r0 + 1}:G{r0 + NE}', FormulaRule(formula=[f'N(G{r0 + 1})>0'], fill=fill(RED), font=Font(bold=True, color='C0392B', name=F)))
D.conditional_formatting.add(f'H{r0 + 1}:H{r0 + NE}', FormulaRule(formula=[f'N(H{r0 + 1})>0'], fill=fill(AMBER), font=Font(bold=True, name=F)))
D.conditional_formatting.add(f'C{r0 + 1}:C{r0 + NE}', FormulaRule(formula=[f'LEFT(C{r0 + 1},3)="Còn"'], font=Font(color=BLUE, bold=True, name=F)))

ch = BarChart(); ch.type = 'bar'; ch.grouping = 'stacked'; ch.overlap = 100; ch.gapWidth = 60
ch.title = 'Việc theo sự kiện'; ch.y_axis.title = None; ch.x_axis.title = None
data = Reference(D, min_col=5, max_col=7, min_row=r0, max_row=r0 + NE)
cats = Reference(D, min_col=1, min_row=r0 + 1, max_row=r0 + NE)
ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
for s, col in zip(ch.series, ['3E9B5F', '9DB4D6', 'E0574B']):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ch.x_axis.scaling.orientation = 'maxMin'; ch.legend.position = 'b'
ch.height = 9.5; ch.width = 17; ch.visible_cells_only = False
D.add_chart(ch, 'L7')

# Theo người
NP = 15
r1 = r0 + NE + 3; sec(f'A{r1 - 1}', 'THEO NGƯỜI PHỤ TRÁCH')
thead(r1, ['Người', 'Đang mở', 'Quá hạn', 'Đến hạn ≤ 3 ngày', 'Xong', 'Tổng', '%', '', '', 'Tiến độ'])
D.merge_cells(start_row=r1, start_column=7, end_row=r1, end_column=9)
for i in range(NP):
    r = r1 + 1 + i
    p = f"'Cài đặt'!$D${5 + i}"
    D[f'A{r}'] = f'=IF({p}="","",{p})'
    D[f'B{r}'] = f'=IF(A{r}="","",COUNTIFS({GV("F")},A{r},{GV("R")},"Có",{GV("I")},"<>Xong",{GV("I")},"<>Bỏ qua"))'
    D[f'C{r}'] = f'=IF(A{r}="","",COUNTIFS({GV("F")},A{r},{GV("R")},"Có",{GV("L")},"*QUÁ HẠN*"))'
    D[f'D{r}'] = f'=IF(A{r}="","",COUNTIFS({GV("F")},A{r},{GV("R")},"Có",{GV("L")},"*Đến hạn*"))'
    D[f'E{r}'] = f'=IF(A{r}="","",COUNTIFS({GV("F")},A{r},{GV("R")},"Có",{GV("I")},"Xong"))'
    D[f'F{r}'] = f'=IF(A{r}="","",B{r}+E{r})'
    D[f'G{r}'] = f'=IF(OR(A{r}="",N(F{r})=0),"",E{r}/F{r})'
    D.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9)
    D[f'J{r}'] = f'=IF(G{r}="","",REPT("█",ROUND(G{r}*10,0))&REPT("░",10-ROUND(G{r}*10,0)))'
    for j in range(10):
        c = D.cell(r, 1 + j); c.border = BOX
        if j != 6 or True: c.font = font(size=9)
    D[f'G{r}'].number_format = '0%'; D[f'J{r}'].font = font(size=9, color=BLUE)
D.conditional_formatting.add(f'C{r1 + 1}:C{r1 + NP}', FormulaRule(formula=[f'N(C{r1 + 1})>0'], fill=fill(RED), font=Font(bold=True, color='C0392B', name=F)))
D.conditional_formatting.add(f'D{r1 + 1}:D{r1 + NP}', FormulaRule(formula=[f'N(D{r1 + 1})>0'], fill=fill(AMBER), font=Font(bold=True, name=F)))
r = r1 + NP + 1
D[f'A{r}'] = 'Chưa giao người'; D[f'A{r}'].font = font(size=9, bold=True, color='B85C1E')
D[f'B{r}'] = f'=COUNTIFS({GV("R")},"Có",{GV("L")},"CHƯA GIAO*")'; D[f'B{r}'].font = font(size=9, bold=True, color='B85C1E')

# Việc cần xử lý
NT = 15
r2 = r + 3; sec(f'A{r2 - 1}', 'VIỆC CẦN XỬ LÝ (hạn gần nhất, chưa xong)')
thead(r2, ['Việc', 'Hạn', 'Cảnh báo', 'Người phụ trách', 'Sự kiện', '', '', '', '', 'Trạng thái'])
D.merge_cells(start_row=r2, start_column=5, end_row=r2, end_column=9)
for i in range(NT):
    r = r2 + 1 + i
    pos = f'IFERROR(MATCH(SMALL({GV("Q")},{i + 1}),{GV("Q")},0),0)'
    g = lambda c: f'=IF({pos}=0,"",INDEX({GV(c)},{pos}))'
    gt = lambda c: g(c)[:-1] + '&"")'
    D[f'A{r}'] = gt('E'); D[f'B{r}'] = g('H'); D[f'C{r}'] = gt('L'); D[f'D{r}'] = gt('F'); D[f'E{r}'] = gt('A'); D[f'J{r}'] = gt('I')
    D.merge_cells(start_row=r, start_column=5, end_row=r, end_column=9)
    for j in range(10):
        c = D.cell(r, 1 + j); c.border = BOX; c.font = font(size=9); c.alignment = Alignment(vertical='center', wrap_text=True)
    D[f'B{r}'].number_format = 'dd/mm'; D.row_dimensions[r].height = 30
D.conditional_formatting.add(f'A{r2 + 1}:J{r2 + NT}', FormulaRule(formula=[f'ISNUMBER(SEARCH("QUÁ HẠN",$C{r2 + 1}))'], fill=fill(RED)))
D.conditional_formatting.add(f'A{r2 + 1}:J{r2 + NT}', FormulaRule(formula=[f'ISNUMBER(SEARCH("Đến hạn",$C{r2 + 1}))'], fill=fill(AMBER)))
D.conditional_formatting.add(f'D{r2 + 1}:D{r2 + NT}', FormulaRule(formula=[f'AND($A{r2 + 1}<>"",$D{r2 + 1}="")'], fill=fill('F6C9A8')))
D.freeze_panes = 'A6'

# ═════════ Hướng dẫn
H = wsH
title(H, 'CÁCH DÙNG', 'File chung cho MỌI sự kiện của MMH — đi kèm SOP MMHSOP_0XX Quy trình tổ chức sự kiện.')
H.column_dimensions['A'].width = 4; H.column_dimensions['B'].width = 24; H.column_dimensions['C'].width = 96
steps = [
 ('Sự kiện tự vào file', 'Mỗi 10 phút hệ thống đọc PLAN_EVENT + REPORT_EVENT của file "2. Vietnam - Marketing Offline FY68" ⇒ tab Sự kiện. '
                         'Sự kiện mới có ngày cụ thể ⇒ tự tạo bộ việc ở tab Giao việc theo tab Mẫu việc, giao sẵn cho người mặc định của từng vai trò (tab Cài đặt).'),
 ('Owner giao việc', 'Tab Giao việc: chọn Người phụ trách (cột F, ô cam = chưa có người), đổi Hạn nếu cần (gõ ngày đè lên). '
                     'Việc không cần ⇒ Trạng thái "Bỏ qua". Thêm việc riêng cho 1 sự kiện: sao chép NGUYÊN DÒNG một việc của sự kiện đó, dán vào dòng mới rồi sửa nội dung. Không theo dõi một sự kiện ⇒ tab Sự kiện ▸ Theo dõi = Không.'),
 ('Người phụ trách', 'Lọc cột F theo tên mình ⇒ cập nhật Trạng thái, Ngày xong, Ghi chú / link ngay khi làm xong (trong ngày, trước 17:00).'),
 ('Đọc Dashboard', 'Đỏ = quá hạn · vàng = đến hạn trong 3 ngày · cam = chưa giao người. Bảng 1: từng sự kiện · Bảng 2: từng người · Bảng 3: 15 việc cần xử lý gần nhất.'),
 ('Đổi ngày sự kiện', 'Sửa ở file Offline ⇒ ngày D và mọi hạn tự tính lại (trừ hạn Owner đã gõ tay).'),
 ('Quy tắc thông tin', '① File này là nơi DUY NHẤT ghi trạng thái; Zalo nhóm sự kiện chỉ để báo nhanh. '
                       '② Vướng mắc / nguy cơ trễ ⇒ ghi ở Ghi chú + báo Owner ngay, không đợi quá hạn. '
                       '③ Quá hạn ≥ 2 ngày làm việc hoặc ảnh hưởng ngày D ⇒ Owner báo Giám đốc trong ngày. '
                       '④ Thay đổi ngày / địa điểm / KOL / ngân sách ⇒ sửa file Offline + email mọi bên trong 24 giờ. '
                       '⑤ Không gửi danh sách khách qua Zalo / email cá nhân — chỉ gửi link Sheet.'),
 ('Ngày tính hạn', 'D − n = ngày lịch; D + n = ngày làm việc T2–T6 trừ ngày nghỉ ở tab Cài đặt. Sự kiện nhiều ngày: D + n tính từ ngày kết thúc.'),
 ('Dữ liệu ví dụ', 'Các sự kiện "(ví dụ)" (mã VD-…) và người "(ví dụ)" chỉ để xem trước — tự xoá ở lần đồng bộ đầu tiên; điền tên thật ở tab Cài đặt.'),
]
for i, (a, b) in enumerate(steps):
    r = 4 + i
    H.cell(r, 2, a).font = font(bold=True, color=NAVY); H.cell(r, 3, b).font = font()
    H.cell(r, 2).alignment = WR; H.cell(r, 3).alignment = WR; H.row_dimensions[r].height = 58 if i == 5 else 42

for ws, col in [(wsD, NAVY), (wsG, BLUE), (wsE, '3E9B5F'), (wsT, GRAY), (wsC, MUTE), (wsH, MUTE)]:
    ws.sheet_properties.tabColor = col
wb.save(OUT)
print('ok', len(rows), 'việc mẫu')
